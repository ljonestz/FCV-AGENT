"""
FCV Screener - Word readout design system (python-docx), "Option 1 / Banner".

One visual system for both routes:
  * FCV route      -> build_fcv_readout.py
  * Climate lens   -> build_climate_readout.py  (same blue palette, "CLIMATE LENS" tag in banner + running header)

Palette matches the web app (WBG navy #002244, WBG blue #009FDA).
Every component is a method on `Report`; the build scripts only decide order and content.
"""
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
import os, re

HERE = os.path.dirname(os.path.abspath(__file__))

# --------------------------------------------------------------------------- low-level helpers

# OOXML child-order lists (Word is strict about element order)
_ORDER = {
    "tblPr": ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc",
              "tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption","tblDescription"],
    "tcPr": ["cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd","noWrap","tcMar","textDirection",
             "tcFitText","vAlign","hideMark"],
    "pPr": ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl","numPr","suppressLineNumbers",
            "pBdr","shd","tabs","suppressAutoHyphens","kinsoku","wordWrap","overflowPunct","topLinePunct","autoSpaceDE",
            "autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents",
            "suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle",
            "rPr","sectPr","pPrChange"],
    "rPr": ["rStyle","rFonts","b","bCs","i","iCs","caps","smallCaps","strike","dstrike","outline","shadow","emboss",
            "imprint","noProof","snapToGrid","vanish","webHidden","color","spacing","w","kern","position","sz","szCs",
            "highlight","u","effect","bdr","shd","fitText","vertAlign","rtl","cs","em","lang","eastAsianLayout",
            "specVanish","oMath"],
    "trPr": ["cnfStyle","divId","gridBefore","gridAfter","wBefore","wAfter","cantSplit","trHeight","tblHeader",
             "tblCellSpacing","jc","hidden"],
}


def put(parent, child):
    """Insert child into parent respecting schema order; replaces an existing element of the same tag."""
    ptag = parent.tag.split("}")[1]
    ctag = child.tag.split("}")[1]
    order = _ORDER[ptag]
    for old in parent.findall(child.tag):
        parent.remove(old)
    ci = order.index(ctag)
    for i, existing in enumerate(list(parent)):
        etag = existing.tag.split("}")[1]
        if etag in order and order.index(etag) > ci:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child

def rgb(h): return RGBColor.from_string(h)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def shade_cell(cell, hexcol):
    put(cell._tc.get_or_add_tcPr(), _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": hexcol}))


def cell_borders(cell, **sides):
    """sides: top/bottom/left/right = (size_eighths_pt, hex) or None for no border."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = put(tcPr, _el("w:tcBorders"))
    existing = {e.tag.split("}")[1]: e for e in b}
    for side in ("top", "left", "bottom", "right"):
        if side in existing and side not in sides:
            continue
        if side in existing:
            b.remove(existing[side])
        if side not in sides:
            continue
        v = sides[side]
        if v is None:
            b.append(_el(f"w:{side}", **{"w:val": "nil"}))
        else:
            b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": v[0], "w:space": 0, "w:color": v[1]}))
    kids = sorted(list(b), key=lambda e: ["top", "left", "bottom", "right"].index(e.tag.split("}")[1]))
    for k in kids:
        b.remove(k); b.append(k)


def cell_margins(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    m = tcPr.find(qn("w:tcMar"))
    cur = {}
    if m is not None:
        cur = {e.tag.split("}")[1]: e.get(qn("w:w")) for e in m}
    m = put(tcPr, _el("w:tcMar"))
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if v is not None:
            m.append(_el(f"w:{side}", **{"w:w": int(v * 1440), "w:type": "dxa"}))
        elif side in cur:
            m.append(_el(f"w:{side}", **{"w:w": cur[side], "w:type": "dxa"}))


def table_setup(table, widths_in, indent_in=0.0, borders=None, cell_pad=(0.06, 0.06, 0.08, 0.08)):
    """Fixed layout, exact column widths, table-level borders (dict side->(sz,hex)|None)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblInd", "w:tblBorders", "w:tblCellMar"):
        for old in tblPr.findall(qn(tag)):
            tblPr.remove(old)
    put(tblPr, _el("w:tblW", **{"w:w": int(sum(widths_in) * 1440), "w:type": "dxa"}))
    put(tblPr, _el("w:tblInd", **{"w:w": int(indent_in * 1440), "w:type": "dxa"}))
    bd = _el("w:tblBorders")
    borders = borders or {}
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        v = borders.get(side)
        if v is None:
            bd.append(_el(f"w:{side}", **{"w:val": "nil"}))
        else:
            bd.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": v[0], "w:space": 0, "w:color": v[1]}))
    put(tblPr, bd)
    put(tblPr, _el("w:tblLayout", **{"w:type": "fixed"}))
    mar = _el("w:tblCellMar")
    t, b, l, r = cell_pad
    for side, v in (("top", t), ("left", l), ("bottom", b), ("right", r)):
        mar.append(_el(f"w:{side}", **{"w:w": int(v * 1440), "w:type": "dxa"}))
    put(tblPr, mar)
    # grid
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths_in):
            c.width = Inches(w)
    table.autofit = False


def row_no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    put(trPr, _el("w:cantSplit"))
    if header:
        put(trPr, _el("w:tblHeader"))


def row_height(row, inches, exact=False):
    trPr = row._tr.get_or_add_trPr()
    put(trPr, _el("w:trHeight", **{"w:val": int(inches * 1440), "w:hRule": "exact" if exact else "atLeast"}))


def pfmt(p, before=0, after=0, line=None, align=None, keep_next=False, left=None, right=None):
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    if line:
        f.line_spacing = line
    if align is not None:
        f.alignment = align
    if keep_next:
        f.keep_with_next = True
    if left is not None:
        f.left_indent = Inches(left)
    if right is not None:
        f.right_indent = Inches(right)
    return p


def p_border(p, side, sz, hexcol, space=4):
    pPr = p._p.get_or_add_pPr()
    b = pPr.find(qn("w:pBdr"))
    if b is None:
        b = put(pPr, _el("w:pBdr"))
    b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": space, "w:color": hexcol}))
    order = ["top", "left", "bottom", "right", "between", "bar"]
    for k in sorted(list(b), key=lambda e: order.index(e.tag.split("}")[1])):
        b.remove(k); b.append(k)


def p_shade(p, hexcol):
    put(p._p.get_or_add_pPr(), _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": hexcol}))


def add_run(p, text, size=None, bold=None, italic=None, color=None, font=None, caps=False, spacing=None, shade=None):
    r = p.add_run(text)
    f = r.font
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)
    if font:
        f.name = font
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    if caps: f.all_caps = True
    rPr = r._r.get_or_add_rPr()
    if spacing is not None:
        put(rPr, _el("w:spacing", **{"w:val": int(spacing * 20)}))
    if shade:
        put(rPr, _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": shade}))
    return r


def add_field(p, instr, size, color, font):
    r = p.add_run()
    r.font.size, r.font.color.rgb, r.font.name = Pt(size), rgb(color), font
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = p.add_run(); r2.font.size, r2.font.color.rgb, r2.font.name = Pt(size), rgb(color), font
    it = _el("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
    r2._r.append(it)
    r3 = p.add_run(); r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = p.add_run("1"); r4.font.size, r4.font.color.rgb, r4.font.name = Pt(size), rgb(color), font
    r5 = p.add_run(); r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def first_para(cell):
    return cell.paragraphs[0]



# --------------------------------------------------------------------------- palette (matches web app)
NAVY = "002244"        # --wbg-navy
BLUE = "009FDA"        # --wbg-blue
BLUE_DARK = "0072BB"
TINT = "E8F5FC"        # --color-blue-soft
TINT2 = "F3F6F9"
SURFACE_GREY = "F7F8FA"
BODY_TXT = "111827"
GREY_TXT = "5F6B7A"
GREY_LIGHT = "D9DDE3"
AMBER_BG, AMBER_EDGE, AMBER_TXT = "FFF7E6", "F9AA2A", "4A3000"
VERIFY_BG = "FDEBC8"
METER = {5: ["CFE6F4", "9FD0EA", "5FB8E2", "009FDA", "0072BB"],
         6: ["D7EAF5", "B5DBEF", "8CC8E8", "55B2E0", "009FDA", "0072BB"]}
METER_OFF = "E5E7EB"
PILL = {  # (background, text)
    "High": ("FDE4E1", "B42318"), "Medium": ("FEF0C7", "8A4B08"), "Low": ("E3F4E8", "1E6B3A"),
    "high": ("E3F4E8", "1E6B3A"), "medium": ("FEF0C7", "8A4B08"), "low": ("EEF0F3", "4B5563"),  # confidence
}
FONT = "Arial"

# Standard advisory disclaimer shown under the banner on every readout
BE_AWARE = ['Be aware: LLMs are not magic, and neither is the FCV Project Screener. It gives advice based only on the documents and guidance it has been given, and it can be wrong. Treat everything here as advisory. Do not assume any finding, rating, or suggestion is definitive, including references to WBG policies, instruments, or procedures, which should be verified with OPCS. You are strongly advised to seek corroboration from relevant FCV and OPCS human experts.', 'Assess findings against project objectives, scope and delivery constraints. Responsiveness is not expected for every project; contributing to it often means going above and beyond. Select recommendations that are relevant and feasible.']

BODY_SIZE = 10
PAGE_W, PAGE_H, MARGIN = 8.5, 11.0, 0.9
CONTENT_W = PAGE_W - 2 * MARGIN   # 6.7in
BANNER_H = 1.75


def first_para(cell):
    return cell.paragraphs[0]


def logo_path(white=True):
    return os.path.join(HERE, "logo", f"fcv_screener_mark_{'white' if white else 'color'}.png")


VERIFY_RE = re.compile(r"(\[Verify before sharing:.*?\])")
_SENTENCE_END_RE = re.compile(r"^(.*?[.!?])(\s+)(.*)$", re.DOTALL)


def _split_first_sentence(text):
    """Return (first_sentence_incl_trailing_space, remainder).

    House style bolds the opening sentence of every body paragraph. Falls back to
    (text, '') for a single sentence (bold the whole thing) and extends past a very
    short leading fragment such as an abbreviation ("e.g. ")."""
    if not text:
        return text, ""
    m = _SENTENCE_END_RE.match(text)
    if not m:
        return text, ""
    lead, gap, rest = m.group(1), m.group(2), m.group(3)
    if len(lead) < 12 and rest:
        m2 = _SENTENCE_END_RE.match(rest)
        if m2:
            return lead + gap + m2.group(1) + m2.group(2), m2.group(3)
    return lead + gap, rest


class Report:
    """lens: 'fcv' or 'climate'.  meta keys: title, project_line, running_label, date, classification,
    generator, disclaimer."""

    def __init__(self, lens, meta, document=None):
        self.lens, self.meta = lens, meta
        self.doc = document if document is not None else Document()
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None:
            z.set(qn("w:percent"), "100")
        self._bm = 0
        self._sec = 0
        self._setup_styles()
        self._setup_page()

    # ================================================================ page furniture
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
        s.left_margin = s.right_margin = Inches(MARGIN)
        s.top_margin, s.bottom_margin = Inches(0.85), Inches(0.8)
        s.header_distance, s.footer_distance = Inches(0.4), Inches(0.35)
        s.different_first_page_header_footer = True
        self._running_header(s.header)
        self._footer(s.footer)
        self._footer(s.first_page_footer)

    def _tabbed(self, p):
        p.style = self.doc.styles["Normal"]
        pfmt(p, line=1.0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_W), alignment=2)

    def _running_header(self, header):
        p = header.paragraphs[0]; self._tabbed(p)
        add_run(p, "FCV SCREENER", 7.5, True, color=BLUE, spacing=1)
        label = self.meta["title"]
        if self.meta["running_label"]:
            label += "  \u00b7  " + self.meta["running_label"]
        add_run(p, "   " + label, 7.5, color=GREY_TXT)
        if self.lens == "climate":
            add_run(p, "   ")
            add_run(p, " CLIMATE LENS ", 7, True, color="FFFFFF", spacing=0.8, shade=BLUE)
        add_run(p, "\t" + self.meta["classification"].upper(), 7.5, True, color=GREY_TXT, spacing=0.5)
        p_border(p, "bottom", 4, GREY_LIGHT, space=4)

    def _footer(self, footer):
        p = footer.paragraphs[0]; self._tabbed(p)
        p_border(p, "top", 4, GREY_LIGHT, space=4)
        add_run(p, "AI-assisted output — verify before operational use", 7.5, italic=True, color=GREY_TXT)
        add_run(p, "\tPage ", 7.5, color=GREY_TXT)
        add_field(p, "PAGE", 7.5, GREY_TXT, FONT)
        add_run(p, " of ", 7.5, color=GREY_TXT)
        add_field(p, "NUMPAGES", 7.5, GREY_TXT, FONT)

    def _setup_styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = FONT
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        n.font.size = Pt(BODY_SIZE)
        n.font.color.rgb = rgb(BODY_TXT)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.15
        for lvl, size in ((1, 15), (2, 11.5), (3, 10.5)):
            h = st[f"Heading {lvl}"]
            rf = h.element.rPr.rFonts
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            h.font.size, h.font.bold, h.font.italic = Pt(size), True, False
            h.font.color.rgb = rgb(NAVY)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.space_before = h.paragraph_format.space_after = Pt(0)

    # ================================================================ masthead
    def banner(self):
        """Full-bleed navy banner on page 1 (floating table in the first-page header)."""
        hdr = self.doc.sections[0].first_page_header
        p0 = hdr.paragraphs[0]
        pfmt(p0, before=112, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)   # reserves banner height
        add_run(p0, self.meta["classification"].upper(), 7.5, True, color=GREY_TXT, spacing=0.5)
        tbl = hdr.add_table(rows=1, cols=2, width=Inches(PAGE_W))
        table_setup(tbl, [1.35, PAGE_W - 1.35], cell_pad=(0, 0, 0, 0))
        put(tbl._tbl.tblPr, _el("w:tblpPr", **{"w:leftFromText": 0, "w:rightFromText": 0, "w:topFromText": 0,
                                               "w:bottomFromText": 0, "w:vertAnchor": "page", "w:horzAnchor": "page",
                                               "w:tblpX": 0, "w:tblpY": 0}))
        put(tbl._tbl.tblPr, _el("w:tblOverlap", **{"w:val": "never"}))
        p0._p.addprevious(tbl._tbl)
        a, b = tbl.rows[0].cells
        row_height(tbl.rows[0], BANNER_H, exact=True)
        for c in (a, b):
            shade_cell(c, NAVY)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_borders(c, top=(4, NAVY), left=(4, NAVY), right=(4, NAVY), bottom=(48, BLUE))
        cell_margins(a, right=0.24)
        cell_margins(b, right=MARGIN)
        pa = first_para(a); pfmt(pa, align=WD_ALIGN_PARAGRAPH.RIGHT)
        pa.add_run().add_picture(logo_path(True), width=Inches(0.7))
        pb = first_para(b); pfmt(pb, after=3)
        add_run(pb, "FCV SCREENER", 8.5, True, color=BLUE, spacing=2)
        if self.lens == "climate":
            add_run(pb, "    ")
            add_run(pb, "  CLIMATE LENS  ", 8.5, True, color="FFFFFF", spacing=1.5, shade=BLUE)
        p = b.add_paragraph(); pfmt(p, after=4)
        add_run(p, self.meta["title"], 24, True, color="FFFFFF")
        p = b.add_paragraph(); pfmt(p)
        add_run(p, self.meta["project_line"], 9.5, color="D6E2EC")
        its_url = self.meta.get("its_url")
        if its_url:
            p = b.add_paragraph(); pfmt(p, before=3)
            add_run(p, self.meta.get("its_label", "Hosted version for colleagues: "), 8, color="9FD3EC")
            self.external_link(p, its_url, its_url.replace("https://", "").replace("http://", ""),
                               8, bold=True, color="FFFFFF")

    def disclaimer_line(self):
        p = self.doc.paragraphs[0] if self.doc.paragraphs and not self.doc.paragraphs[0].text else self.doc.add_paragraph()
        pfmt(p, after=10)
        add_run(p, self.meta["generator"] + ". ", 8, True, color=GREY_TXT)
        add_run(p, self.meta["disclaimer"], 8, italic=True, color=GREY_TXT)

    # ================================================================ text primitives
    def _runs_with_verify(self, p, text, size=None, color=None, italic=None, bold=None):
        """Body text; '[Verify before sharing: ...]' notes are highlighted so reviewers can't miss them."""
        for part in VERIFY_RE.split(text):
            if not part:
                continue
            if VERIFY_RE.fullmatch(part):
                add_run(p, part, (size or BODY_SIZE) - 1, True, italic, color=AMBER_TXT, shade=VERIFY_BG)
            else:
                add_run(p, part, size, bold=bold, italic=italic, color=color)

    def body(self, text, after=6, size=None, color=None, italic=None, keep=False, bold_lead=True):
        """Narrative paragraph. Per house style the opening sentence is bold (set
        bold_lead=False to opt out, e.g. for a muted note)."""
        p = self.doc.add_paragraph(); pfmt(p, after=after, keep_next=keep)
        if bold_lead and not italic:
            lead, rest = _split_first_sentence(text)
            self._runs_with_verify(p, lead, size, color, italic, bold=True)
            if rest:
                self._runs_with_verify(p, rest, size, color, italic)
        else:
            self._runs_with_verify(p, text, size, color, italic)
        return p

    def small(self, text, italic=True, after=6, before=0):
        p = self.doc.add_paragraph(); pfmt(p, before=before, after=after)
        add_run(p, text, 7.5, italic=italic, color=GREY_TXT)
        return p

    def spacer(self, pts=4, keep=False):
        p = self.doc.add_paragraph(); pfmt(p, line=0.5, keep_next=keep)
        add_run(p, "", 1)
        p.paragraph_format.space_after = Pt(pts)

    def label(self, text, before=8, color=BLUE):
        p = self.doc.add_paragraph(); pfmt(p, before=before, after=2, keep_next=True)
        add_run(p, text.upper(), 8, True, color=color, spacing=0.8)
        return p

    # ================================================================ headings
    def h1(self, text, numbered=True):
        self._sec += 1
        p = self.doc.add_paragraph(style="Heading 1")
        pfmt(p, before=18, after=8, keep_next=True)
        if numbered:
            add_run(p, f"{self._sec:02d}", 10, True, color=BLUE)
            add_run(p, "   ")
        add_run(p, text, 15, True, color=NAVY)
        p_border(p, "bottom", 12, BLUE, space=3)
        return p

    def h2_caps(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        pfmt(p, before=10, after=4, keep_next=True)
        add_run(p, text.upper(), 9, True, color=BLUE, spacing=0.8)
        return p

    def h2_question(self, text):
        """Longer headings, e.g. assessment questions."""
        p = self.doc.add_paragraph(style="Heading 2")
        pfmt(p, before=12, after=5, keep_next=True)
        p_border(p, "left", 24, BLUE, space=6)
        add_run(p, text, 11, True, color=NAVY)
        return p

    # ================================================================ boxes
    def _box(self, fill, edge=None, edge_sz=36, pad=(0.1, 0.1, 0.18, 0.18), indent=0.0, outline=None):
        t = self.doc.add_table(rows=1, cols=1)
        borders = {}
        if edge:
            borders["left"] = (edge_sz, edge)
        if outline:
            for s in ("top", "bottom", "right"):
                borders[s] = (4, outline)
            borders.setdefault("left", (4, outline))
        table_setup(t, [CONTENT_W - indent], indent_in=indent, cell_pad=pad, borders=borders)
        c = t.rows[0].cells[0]
        if fill:
            shade_cell(c, fill)
        return t, c

    def status_line(self, text):
        t, c = self._box(SURFACE_GREY, edge=GREY_TXT, edge_sz=18, pad=(0.06, 0.06, 0.14, 0.14))
        p = first_para(c); pfmt(p)
        add_run(p, text, 8.5, True, color=GREY_TXT)
        self.spacer(4)

    def caution(self, title, body=()):
        t, c = self._box(AMBER_BG, edge=AMBER_EDGE, pad=(0.08, 0.08, 0.16, 0.16))
        p = first_para(c); pfmt(p, after=2 if body else 0)
        add_run(p, "⚠  ", 9, True, color=AMBER_EDGE)
        add_run(p, title, 9, True, color=AMBER_TXT)
        for b in body:
            p = c.add_paragraph(); pfmt(p)
            add_run(p, b, 8.5, color=AMBER_TXT)
        row_no_split(t.rows[0])
        self.spacer(4)

    def info_box(self, text, label=None, bold=False, size=None, keep=False):
        t, c = self._box(TINT, edge=BLUE, pad=(0.1, 0.1, 0.18, 0.18))
        p = first_para(c); pfmt(p, after=3 if label else 0)
        if label:
            add_run(p, label.upper(), 7.5, True, color=BLUE, spacing=1)
            p = c.add_paragraph(); pfmt(p)
        add_run(p, text, size or (BODY_SIZE + (1 if bold else 0)), bold, color=NAVY)
        p.paragraph_format.line_spacing = 1.2
        row_no_split(t.rows[0])
        if keep:
            for q in c.paragraphs:
                q.paragraph_format.keep_with_next = True
        self.spacer(6, keep=keep)

    def drafting_box(self, text, for_text="", label="Suggested drafting", note="ready to adapt into the project document"):
        t, c = self._box(TINT, edge=BLUE, edge_sz=24, pad=(0.09, 0.11, 0.16, 0.16), indent=0.25)
        p = first_para(c); pfmt(p, after=3)
        add_run(p, label.upper(), 7, True, color=BLUE, spacing=1)
        if for_text:
            add_run(p, "   FOR: ", 7, True, color=GREY_TXT, spacing=0.5)
            add_run(p, for_text, 7.5, True, color=NAVY)
        elif note:
            add_run(p, "   " + note, 7, italic=True, color=GREY_TXT)
        p = c.add_paragraph(); pfmt(p)
        add_run(p, text, BODY_SIZE - 0.5, italic=True, color="2B3A48")
        row_no_split(t.rows[0])
        self.spacer(6)

    def grey_box(self, label, paras):
        """paras: list of str or (lead, text) tuples."""
        t, c = self._box(SURFACE_GREY, outline=GREY_LIGHT, pad=(0.09, 0.09, 0.16, 0.16))
        p = first_para(c); pfmt(p, after=3)
        add_run(p, label.upper(), 7.5, True, color=GREY_TXT, spacing=1)
        for x in paras:
            p = c.add_paragraph(); pfmt(p, after=3)
            if isinstance(x, tuple):
                add_run(p, x[0] + " ", BODY_SIZE - 0.5, True, color=NAVY)
                add_run(p, x[1], BODY_SIZE - 0.5)
            else:
                add_run(p, x, BODY_SIZE - 0.5)
        self.spacer(6)

    # ================================================================ ratings
    def meter(self, cell, n, idx, width, left_label, right_label):
        """n-step scale bar; idx is 1-based position of the rating."""
        inner = cell.add_table(rows=1, cols=n)
        seg = round(width / n, 3)
        table_setup(inner, [seg] * n, cell_pad=(0, 0, 0, 0))
        for i, c in enumerate(inner.rows[0].cells):
            shade_cell(c, METER[n][i] if i < idx else METER_OFF)
            cell_borders(c, left=(24, "FFFFFF") if i else None, right=(24, "FFFFFF") if i < n - 1 else None)
            pfmt(first_para(c), line=0.1)
        row_height(inner.rows[0], 0.11, exact=True)
        p = cell.paragraphs[-1]; pfmt(p, before=2)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(seg * n), alignment=2)
        add_run(p, left_label, 6.5, color=GREY_TXT)
        add_run(p, "\t" + right_label, 6.5, color=GREY_TXT)

    def scorecard(self, tiles, count_tile, note=None):
        """tiles: [{label, value, n, idx, lo, hi}] ; count_tile: {label, number, lines:[(text, pill_key|None)]}"""
        w = [2.45, 2.45, CONTENT_W - 4.9]
        t = self.doc.add_table(rows=1, cols=3)
        bd = (4, GREY_LIGHT)
        table_setup(t, w, cell_pad=(0.1, 0.12, 0.14, 0.14),
                    borders={"top": bd, "bottom": bd, "left": bd, "right": bd, "insideV": bd})
        for i, tl in enumerate(tiles):
            c = t.rows[0].cells[i]
            p = first_para(c); pfmt(p, after=1)
            add_run(p, tl["label"].upper(), 7.5, True, color=GREY_TXT, spacing=0.8)
            p = c.add_paragraph(); pfmt(p, after=5)
            add_run(p, tl["value"], 13.5, True, color=NAVY)
            self.meter(c, tl["n"], tl["idx"], w[i] - 0.3, tl["lo"], tl["hi"])
        c = t.rows[0].cells[2]; shade_cell(c, NAVY)
        p = first_para(c); pfmt(p, after=1)
        add_run(p, count_tile["label"].upper(), 7.5, True, color="FFFFFF", spacing=0.8)
        p = c.add_paragraph(); pfmt(p, after=2)
        add_run(p, str(count_tile["number"]), 22, True, color="FFFFFF")
        p = c.add_paragraph(); pfmt(p)
        for text, key in count_tile["lines"]:
            if key:
                bg, fg = PILL[key]; add_run(p, f" {text} ", 8, True, color=fg, shade=bg)
            else:
                add_run(p, text, 8, color="D6E2EC")
            add_run(p, "  ")
        row_no_split(t.rows[0])
        if note:
            self.small(note, before=3, after=4)

    def be_aware(self, paras):
        """Prominent advisory disclaimer under the banner (both lenses). First words up to ':' are bolded."""
        t, c = self._box(AMBER_BG, edge=AMBER_EDGE, edge_sz=48, pad=(0.12, 0.12, 0.2, 0.2))
        for i, text in enumerate(paras):
            p = first_para(c) if i == 0 else c.add_paragraph()
            pfmt(p, after=4 if i < len(paras) - 1 else 0)
            p.paragraph_format.line_spacing = 1.2
            lead, sep, rest = text.partition(":")
            if sep and len(lead) < 20:
                add_run(p, lead + ": ", 9.5, True, color=AMBER_TXT)
                add_run(p, rest.strip(), 9, color=AMBER_TXT)
            else:
                add_run(p, text, 9, color=AMBER_TXT)
        row_no_split(t.rows[0])
        self.spacer(6)

    def rating_segment(self, label, value, meaning=(), why=(), notes=(), question="", scale="", first=False):
        """Compact rating write-up for the body (no scale graphic - that lives in the page-1 scorecard)."""
        p = self.doc.add_paragraph(style="Heading 2")
        pfmt(p, before=4 if first else 12, after=4, keep_next=True)
        if not first:
            p_border(p, "top", 4, GREY_LIGHT, space=10)
        add_run(p, label.upper(), 8.5, True, color=GREY_TXT, spacing=0.8)
        add_run(p, "   ")
        add_run(p, f" {value} ", 9.5, True, color="FFFFFF", shade=NAVY)
        for m in meaning:
            q = self.doc.add_paragraph(); pfmt(q, after=3, keep_next=True)
            add_run(q, m, 9.5, True, color=NAVY)
        for w_ in why:
            self.body(w_, after=4)
        for nt in notes:
            self.body(nt, after=4, size=8.5, color=GREY_TXT, italic=True)
        if question or scale:
            q = self.doc.add_paragraph(); pfmt(q, after=6)
            if question:
                add_run(q, "Question assessed: ", 7.5, True, color=GREY_TXT)
                add_run(q, question + "  ", 7.5, color=GREY_TXT)
            if scale:
                add_run(q, scale, 7.5, italic=True, color=GREY_TXT)

    def rating_block(self, label, value, n, idx, lo, hi, meaning=(), why=(), notes=(), question="", scale=""):
        t = self.doc.add_table(rows=1, cols=2)
        table_setup(t, [2.0, CONTENT_W - 2.0], cell_pad=(0.08, 0.08, 0, 0.14))
        a, b = t.rows[0].cells
        p = first_para(a); pfmt(p, after=1)
        add_run(p, label.upper(), 7.5, True, color=GREY_TXT, spacing=0.8)
        p = a.add_paragraph(); pfmt(p, after=4)
        add_run(p, value, 13, True, color=NAVY)
        self.meter(a, n, idx, 1.75, lo, hi)
        for m in meaning:
            p = a.add_paragraph(); pfmt(p, before=4, after=0)
            add_run(p, m, 8, color=BODY_TXT)
        cell_borders(a, right=(12, BLUE))
        cell_margins(b, left=0.22)
        p = first_para(b); pfmt(p, after=2)
        if why or notes:
            add_run(p, "WHY THIS RATING", 7.5, True, color=BLUE, spacing=0.8)
        for w_ in why:
            p = b.add_paragraph(); pfmt(p, after=4); add_run(p, w_)
        for nt in notes:
            p = b.add_paragraph(); pfmt(p, after=4); add_run(p, nt, 8.5, italic=True, color=GREY_TXT)
        if question:
            p = b.add_paragraph(); pfmt(p, after=1)
            add_run(p, "Question assessed: ", 7.5, True, color=GREY_TXT); add_run(p, question, 7.5, color=GREY_TXT)
        if scale:
            p = b.add_paragraph(); pfmt(p)
            add_run(p, scale, 7.5, italic=True, color=GREY_TXT)
        row_no_split(t.rows[0])
        self.spacer(8)

    # ================================================================ tables
    def data_table(self, headers, widths, rows, center=(0,), pill_cols=(), keep_together=True, link_col=None):
        """rows: list of lists of str. Column 0 is rendered as a blue number."""
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table_setup(t, widths, cell_pad=(0.07, 0.07, 0.08, 0.08),
                    borders={"bottom": (8, NAVY), "insideH": (4, GREY_LIGHT)})
        hr = t.rows[0]; row_no_split(hr, header=True)
        for i, h in enumerate(headers):
            c = hr.cells[i]; shade_cell(c, NAVY); cell_borders(c, bottom=(18, BLUE))
            p = first_para(c); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if i in center else None)
            add_run(p, h, 8, True, color="FFFFFF")
        for r, vals in enumerate(rows, start=1):
            row = t.rows[r]; row_no_split(row)
            for i, v in enumerate(vals):
                c = row.cells[i]; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if r % 2 == 0:
                    shade_cell(c, TINT2)
                p = first_para(c); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if i in center else None)
                if i == 0:
                    add_run(p, v, 11, True, color=BLUE)
                elif i in pill_cols and v in PILL:
                    bg, fg = PILL[v]; add_run(p, f" {v[:1].upper() + v[1:]} ", 8, True, color=fg, shade=bg)
                elif i == link_col:
                    self.internal_link(p, v[1], v[0], 9)
                else:
                    add_run(p, v, 9 if i == 1 else 8.5, color=BODY_TXT if i < 3 else GREY_TXT)
        if keep_together:
            for row in t.rows[:-1]:
                for c in row.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
        self.spacer(8)

    # ================================================================ links & bookmarks
    def bookmark(self, p, name):
        self._bm += 1
        start = _el("w:bookmarkStart", **{"w:id": self._bm, "w:name": name})
        end = _el("w:bookmarkEnd", **{"w:id": self._bm})
        p._p.insert(1 if p._p.pPr is not None else 0, start)
        p._p.append(end)

    def internal_link(self, p, anchor, text, size=8, color=BLUE):
        h = _el("w:hyperlink", **{"w:anchor": anchor, "w:history": 1})
        r = add_run(p, text, size, color=color)
        r.font.underline = True
        h.append(r._r)
        p._p.append(h)

    def external_link(self, p, url, text, size=10, bold=True, color=NAVY):
        # Relate against the part that OWNS this paragraph (document body or a
        # header part). Relating to self.doc.part would create an r:id that does
        # not resolve when the hyperlink sits in a header, which Word flags as a
        # repair. p.part resolves correctly for both locations.
        rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
        r = add_run(p, text, size, bold, color=color)
        r.font.underline = True
        h.append(r._r)
        p._p.append(h)

    def related_line(self, label, items):
        """items: [(text, anchor)]"""
        p = self.doc.add_paragraph(); pfmt(p, before=2, after=8)
        add_run(p, label + " ", 8, color=GREY_TXT)
        for i, (text, anchor) in enumerate(items):
            if i:
                add_run(p, "  ·  ", 8, color=GREY_TXT)
            self.internal_link(p, anchor, text, 8)

    # ================================================================ priority cards
    def priority_header(self, num, title, meta, anchor=None, widths=None):
        t = self.doc.add_table(rows=1, cols=2)
        table_setup(t, [0.55, CONTENT_W - 0.55], cell_pad=(0.08, 0.08, 0.1, 0.12))
        a, b = t.rows[0].cells
        shade_cell(a, BLUE); shade_cell(b, TINT)
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = first_para(a); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        add_run(p, str(num), 20, True, color="FFFFFF")
        p = first_para(b); p.style = self.doc.styles["Heading 2"]; pfmt(p, keep_next=True)
        add_run(p, f"PRIORITY {num}", 7.5, True, color=NAVY, spacing=1)
        if anchor:
            self.bookmark(p, anchor)
        p = b.add_paragraph(); pfmt(p, keep_next=True)
        add_run(p, title, 11.5, True, color=NAVY)
        row_no_split(t.rows[0])
        if meta:
            w = widths or [CONTENT_W / len(meta)] * len(meta)
            mt = self.doc.add_table(rows=2, cols=len(meta))
            table_setup(mt, w, cell_pad=(0.05, 0.03, 0.08, 0.08),
                        borders={"bottom": (4, GREY_LIGHT), "insideV": (4, GREY_LIGHT)})
            for i, (k, v) in enumerate(meta):
                pk = first_para(mt.rows[0].cells[i]); pfmt(pk, keep_next=True)
                add_run(pk, k.upper(), 6.5, True, color=GREY_TXT, spacing=0.5)
                pv = first_para(mt.rows[1].cells[i]); pfmt(pv, after=3, keep_next=True)
                if v in PILL:
                    bg, fg = PILL[v]; add_run(pv, f" {v[:1].upper() + v[1:]} ", 8, True, color=fg, shade=bg)
                else:
                    add_run(pv, v, 8.5, True, color=NAVY)
            for r in mt.rows:
                row_no_split(r)
        self.spacer(2, keep=True)

    def who_strip(self, pairs):
        t = self.doc.add_table(rows=1, cols=len(pairs))
        table_setup(t, [CONTENT_W / len(pairs)] * len(pairs), cell_pad=(0.06, 0.06, 0.1, 0.1),
                    borders={"insideV": (24, "FFFFFF")})
        for i, (k, v) in enumerate(pairs):
            c = t.rows[0].cells[i]; shade_cell(c, TINT2)
            p = first_para(c); pfmt(p)
            add_run(p, k.upper() + "   ", 6.5, True, color=GREY_TXT, spacing=0.5)
            add_run(p, v, 8.5, True, color=NAVY)
        self.spacer(8)

    def notes_grid(self, notes):
        t = self.doc.add_table(rows=len(notes), cols=2)
        table_setup(t, [1.55, CONTENT_W - 1.55], cell_pad=(0.06, 0.06, 0, 0.12),
                    borders={"top": (4, GREY_LIGHT), "insideH": (4, GREY_LIGHT), "bottom": (4, GREY_LIGHT)})
        for r, (k, v) in enumerate(notes):
            a, b = t.rows[r].cells
            p = first_para(a); pfmt(p); add_run(p, k, 8.5, True, color=NAVY)
            p = first_para(b); pfmt(p); add_run(p, v, BODY_SIZE - 0.5)
            row_no_split(t.rows[r])
        self.spacer(6)

    def action_heading(self, num, text):
        p = self.doc.add_paragraph(); pfmt(p, before=8, after=3, keep_next=True)
        add_run(p, f"{num}  ", 9.5, True, color=BLUE)
        add_run(p, text, 9.5, True, color=NAVY)

    # ================================================================ lists
    def numbered_items(self, items):
        """items: [{num, title, why, how}] - checklist style."""
        t = self.doc.add_table(rows=len(items), cols=2)
        table_setup(t, [0.5, CONTENT_W - 0.5], cell_pad=(0.08, 0.08, 0, 0.1),
                    borders={"insideH": (4, GREY_LIGHT), "bottom": (4, GREY_LIGHT), "top": (4, GREY_LIGHT)})
        for r, it in enumerate(items):
            a, b = t.rows[r].cells
            p = first_para(a); pfmt(p); add_run(p, it["num"], 12, True, color=BLUE)
            p = first_para(b); pfmt(p, after=3); add_run(p, it["title"], 9.5, True, color=NAVY)
            if it.get("why"):
                p = b.add_paragraph(); pfmt(p, after=3); add_run(p, it["why"], 9)
            if it.get("how"):
                p = b.add_paragraph(); pfmt(p)
                add_run(p, "HOW TO ADDRESS  ", 7, True, color=BLUE, spacing=0.8)
                add_run(p, it["how"], 9, color=BODY_TXT)
            row_no_split(t.rows[r])
        self.spacer(6)

    def link_items(self, items):
        """items: [{title, url, use}] - reference list with live hyperlinks."""
        t = self.doc.add_table(rows=len(items), cols=2)
        table_setup(t, [0.35, CONTENT_W - 0.35], cell_pad=(0.07, 0.07, 0, 0.1),
                    borders={"insideH": (4, GREY_LIGHT), "bottom": (4, GREY_LIGHT), "top": (4, GREY_LIGHT)})
        for r, it in enumerate(items):
            a, b = t.rows[r].cells
            p = first_para(a); pfmt(p); add_run(p, "▸", 11, True, color=BLUE)
            p = first_para(b); pfmt(p, after=2)
            if it.get("url"):
                self.external_link(p, it["url"], it["title"], 9.5)
            else:
                add_run(p, it["title"], 9.5, True, color=NAVY)
            p = b.add_paragraph(); pfmt(p); add_run(p, it["use"], 9)
            row_no_split(t.rows[r])
        self.spacer(6)

    def diamond_item(self, title, flag, text, last=False):
        t = self.doc.add_table(rows=1, cols=2)
        table_setup(t, [0.4, CONTENT_W - 0.4], cell_pad=(0.04, 0.06, 0, 0.1))
        a, b = t.rows[0].cells
        p = first_para(a); pfmt(p); add_run(p, "◆", 11, color=BLUE)
        p = first_para(b); pfmt(p, after=2); add_run(p, title, 10, True, color=NAVY)
        if flag:
            p = b.add_paragraph(); pfmt(p, after=4); add_run(p, flag, 7.5, True, color=BLUE, spacing=0.5)
        p = b.add_paragraph(); pfmt(p); add_run(p, text)
        row_no_split(t.rows[0])
        if not last:
            self.spacer(4)

    def closing(self, text):
        p = self.doc.add_paragraph(); pfmt(p, before=10)
        p_border(p, "top", 4, GREY_LIGHT, space=6)
        add_run(p, text, 8, italic=True, color=GREY_TXT)

    def save(self, path):
        self.doc.save(path)
