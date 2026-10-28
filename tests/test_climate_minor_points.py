"""WS3: smaller climate & fragility points-to-check from non-admitted gaps."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from sector_lenses.climate_verified_pipeline import _admit_minor_climate_points
from sector_lenses.climate_verified_render import (
    POINTS_TO_CHECK_INTRO,
    build_reader_model,
    render_reader_html,
    write_reader_docx,
)


def test_admit_minor_points_gates_on_non_admitted_gaps_and_dedup():
    known = {"RG-001", "RG-009", "RG-011"}
    admitted = {"RG-001"}          # covered by a priority
    reserved = {"placeholder value in the results framework"}  # a readiness-flag text
    payload = {
        "minor_climate_points": [
            # kept: references a known, non-admitted gap
            {"point": "Heat exposure for site workers is not addressed",
             "why": "Rising temperatures noted but no worker-safety measures.",
             "how_to_check": "Consider a heat-safety line in the ESMP.",
             "residual_gap_ids": ["RG-009"]},
            # dropped: gap is admitted (covered by a priority)
            {"point": "Site criteria unclear", "why": "x", "how_to_check": "y",
             "residual_gap_ids": ["RG-001"]},
            # dropped: unknown gap id
            {"point": "Something else", "why": "x", "how_to_check": "y",
             "residual_gap_ids": ["RG-777"]},
            # dropped: duplicates a reserved (readiness-flag) text
            {"point": "Placeholder value in the results framework",
             "why": "x", "how_to_check": "y", "residual_gap_ids": ["RG-011"]},
        ]
    }
    result = _admit_minor_climate_points(payload, known, admitted, reserved)
    assert [p["point"] for p in result] == [
        "Heat exposure for site workers is not addressed"
    ]
    assert result[0]["residual_gap_ids"] == ["RG-009"]


def _docx_body_paragraphs(document: Document) -> list[str]:
    """Return body paragraphs in OOXML order, including paragraphs in tables."""
    return [Paragraph(element, document).text for element in document.element.body.iter(qn("w:p"))]


def _assessment(minor=True) -> dict[str, object]:
    assessment = {
        "schema_version": "climate-verified-v2.1",
        "run_id": "run-mp",
        "evidence_status": "approved",
        "executive_readout": "Executive readout paragraph. " * 20,
        "judgments": {
            "relevance": {"value": "high", "rationale": "Material."},
            "sensitivity": {"value": "strong", "rationale": "Recognized."},
            "responsiveness": {"value": "emerging", "rationale": "Developing."},
            "operationalization": {"value": "partial", "rationale": "Incomplete."},
        },
        "priorities": [],
        "review_readiness_flags": [
            {"flag": "Placeholder value in the results framework",
             "why_it_matters": "Target cannot be read.",
             "document_basis_ids": ["B-1"],
             "suggested_verification": "Insert the intended figure."},
        ],
    }
    if minor:
        assessment["minor_climate_points"] = [
            {"point": "Heat exposure for site workers is not addressed",
             "why": "Rising temperatures are noted but not worker heat-safety.",
             "how_to_check": "Consider a heat-safety line in the ESMP or POM.",
             "residual_gap_ids": ["RG-009"]},
        ]
    return assessment


def test_reader_and_render_show_two_points_to_check_groups():
    model = build_reader_model(_assessment(minor=True))
    assert model["minor_climate_points"][0]["point"].startswith("Heat exposure")
    html = render_reader_html(model)
    assert "Document points to confirm" in html
    assert "Smaller climate &amp; fragility points to consider" in html
    assert "Heat exposure for site workers is not addressed" in html
    assert html.index("Smaller climate &amp; fragility points") < html.index(
        "Document points to confirm"
    )
    assert html.count('class="climate-item-number"') == 2
    stream = BytesIO()
    write_reader_docx(model, stream)
    stream.seek(0)
    paragraphs = _docx_body_paragraphs(Document(stream))
    text = "\n".join(paragraphs)
    normalized = [paragraph.strip().casefold() for paragraph in paragraphs]
    smaller_heading = "smaller climate & fragility points to consider"
    confirm_heading = "document points to confirm"
    heat_point = "heat exposure for site workers is not addressed"
    placeholder_point = "placeholder value in the results framework"
    assert smaller_heading in text.casefold()
    assert confirm_heading in text.casefold()
    assert heat_point in text.casefold()
    assert normalized.index(smaller_heading) < normalized.index(confirm_heading)
    heat_index = normalized.index(heat_point)
    placeholder_index = normalized.index(placeholder_point)
    assert normalized[heat_index - 1] == "01"
    assert normalized[placeholder_index - 1] == "01"


def test_reader_suppresses_points_section_when_both_groups_are_empty():
    model = build_reader_model(_assessment(minor=False))
    model["review_readiness_flags"] = []

    html = render_reader_html(model)
    assert "Points to check before the decision meeting" not in html
    assert POINTS_TO_CHECK_INTRO not in html
    stream = BytesIO()
    write_reader_docx(model, stream)
    stream.seek(0)
    text = "\n".join(p.text for p in Document(stream).paragraphs)
    assert "Points to check before the decision meeting" not in text
    assert POINTS_TO_CHECK_INTRO not in text


def test_render_without_minor_points_omits_second_group():
    model = build_reader_model(_assessment(minor=False))
    assert model["minor_climate_points"] == []
    html = render_reader_html(model)
    assert "Smaller climate &amp; fragility points to consider" not in html
    # The document flags group still renders.
    assert "Placeholder value in the results framework" in html
    stream = BytesIO()
    write_reader_docx(model, stream)  # must not raise
