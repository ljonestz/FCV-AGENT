"""Complete supporting gap findings in detailed exports, without Summary rendering."""

from io import BytesIO
import html
from docx import Document
from sector_lenses.climate_verified_render import build_reader_model, attach_provenance, render_reader_html, write_reader_docx, HEADINGS
from tests.test_climate_summary_completeness import _run_multiple, GAP_TEXTS


def test_every_summary_gap_has_complete_supporting_prose_in_detailed_exports():
    assessment,_,_=_run_multiple(5)
    reader=attach_provenance(build_reader_model(assessment),assessment)
    for priority in reader["priorities"]:
        priority["limitation"]="Site-level conditions remain unverified."
    rendered=html.unescape(render_reader_html(reader))
    stream=BytesIO();write_reader_docx(reader,stream);stream.seek(0)
    doc=Document(stream)
    word="\n".join(p.text for p in doc.paragraphs)
    for text in GAP_TEXTS:
        assert text in rendered
        assert text in word
    assert "Supporting gap findings" in rendered and "Supporting gap findings" in word
    assert "Five-minute readout" not in rendered and "Five-minute readout" not in word
    assert rendered.index("Supporting gap findings")<rendered.index(HEADINGS[2])


def test_legacy_reader_without_gap_projection_keeps_export_behavior():
    reader={"priorities":[]}
    rendered=render_reader_html(reader)
    stream=BytesIO();write_reader_docx(reader,stream);stream.seek(0)
    assert "Supporting gap findings" not in rendered
    assert "Supporting gap findings" not in " ".join(p.text for p in Document(stream).paragraphs)



def test_browser_detailed_renderer_retains_the_same_gap_findings_for_html_sharing():
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function
    helper=_extract_js_function(INDEX.read_text(encoding="utf-8"),"renderClimateVerifiedAssessment")
    assert "r.summary_gaps" in helper
    assert "Supporting gap findings" in helper
    assert "renderClimateVerifiedSummary" not in helper
