"""Supervised admission preserves useful sections, never a false clean pass."""
from copy import deepcopy
import pytest
from docx import Document
from tests.test_climate_verified_pipeline import (
    _arguments, _responses, _final_reader, _native_pass_review, FakeClient,
    PipelineClients, run_verified_climate_pipeline,
)
from sector_lenses.climate_final_reader import reviewed_final_reader
from sector_lenses.climate_verified_render import build_reader_model, render_reader_html, write_reader_docx


@pytest.fixture
def assessment():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    return run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [_final_reader()], []), FakeClient([_native_pass_review] * 2, [])))


def review_for(final, identifier='NARR-inclusion', severity='minor'):
    record = next(r for r in final['summary_paragraphs'] + final['assessment_narratives'] if r['object_id'] == identifier)
    text = record.get('text') or record['paragraphs'][0]
    issue = {'object_id': identifier, 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
             'explanation': 'The wording needs a scope qualification.', 'quoted_text': text,
             'severity': severity, 'decision_impact': 'No change to the proposed action.' if severity == 'minor' else 'Could change which activities receive funding.'}
    return {'verdict': 'revise', 'object_ids': [identifier], 'reason_codes': list(issue['reason_codes']), 'issues': [issue]}


def test_minor_retains_prose_and_visible_review_across_exports(assessment, tmp_path):
    final = assessment['final_reader']
    assessment['reader_review'] = review_for(final)
    summary, narratives = reviewed_final_reader(final, assessment['reader_review'], assessment)
    assert summary and len(narratives) == 3
    model = build_reader_model(assessment)
    assert model['reader_review_notice']
    assert 'minor' in model['reader_review_notice'].lower()
    assert all(t['narrative_status'] == 'generated' for t in model['assessment_themes'])
    assert model['reader_review_notice'] in render_reader_html(model)
    path = write_reader_docx(model, tmp_path/'review.docx')
    assert model['reader_review_notice'] in '\n'.join(p.text for p in Document(path).paragraphs)
    assert assessment['reader_review']['verdict'] == 'revise'


@pytest.mark.parametrize('identifier,count,has_summary', [('NARR-inclusion',2,True),('SUM-001',3,False)])
def test_consequential_withholds_only_affected_section(assessment, identifier, count, has_summary):
    final = assessment['final_reader']
    review = review_for(final, identifier, 'consequential')
    summary, narratives = reviewed_final_reader(final, review, assessment)
    assert bool(summary) == has_summary
    assert len(narratives) == count
    assert identifier not in [r['object_id'] for r in narratives]


def test_legacy_unrated_issue_is_consequential(assessment):
    review = review_for(assessment['final_reader'])
    del review['issues'][0]['severity']
    del review['issues'][0]['decision_impact']
    assert len(reviewed_final_reader(assessment['final_reader'], review, assessment)[1]) == 2


@pytest.mark.parametrize('defect', ['block','unknown','quote','severity','impact','targets','codes'])
def test_invalid_or_block_review_never_unlocks_prose(assessment, defect):
    final = assessment['final_reader']
    review = review_for(final)
    if defect == 'block': review['verdict'] = 'block'
    if defect == 'unknown': review['issues'][0]['object_id'] = 'NARR-missing'
    if defect == 'quote': review['issues'][0]['quoted_text'] = 'Absent assertion'
    if defect == 'severity': review['issues'][0]['severity'] = 'safe'
    if defect == 'impact': review['issues'][0]['decision_impact'] = ''
    if defect == 'targets': review['object_ids'] = []
    if defect == 'codes': review['reason_codes'] = []
    assert reviewed_final_reader(final, review, assessment) == ([], [])


def test_minor_cannot_override_invalid_evidence(assessment):
    final = deepcopy(assessment['final_reader'])
    review = review_for(final)
    final['assessment_narratives'][0]['evidence_ids'] = ['PF-MISSING']
    assert reviewed_final_reader(final, review, assessment) == ([], [])


def test_minor_runtime_stops_without_paid_correction_and_keeps_attention():
    final = _final_reader()
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    generation = FakeClient(responses + [final], [])
    reviewer = FakeClient([_native_pass_review, review_for(final)], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(generation, reviewer))
    assert not any(call['stage'] == 'reader_correction' for call in generation.calls)
    assert result['validation']['status'] == 'attention'
    assert 'FINAL_READER_REVIEW_REQUIRED' in result['validation']['reason_codes']
    assert result['reader_review']['verdict'] == 'revise'
    assert len(result['assessment_narratives']) == 3


def test_mixed_review_preserves_minor_and_unaffected_answers(assessment):
    final = assessment['final_reader']
    review = review_for(final, 'NARR-dividends', 'minor')
    other = review_for(final, 'NARR-inclusion', 'consequential')
    review['issues'] += other['issues']
    review['object_ids'] += other['object_ids']
    assessment['reader_review'] = review
    summary, narratives = reviewed_final_reader(final, review, assessment)
    assert summary and [r['theme'] for r in narratives] == ['design', 'dividends']
    model = build_reader_model(assessment)
    assert 'Minor uncertainties' in model['reader_review_notice']
    assert 'affected narrative is withheld' in model['reader_review_notice']


def test_native_renderer_uses_shared_hard_ceiling_not_old_editorial_target(assessment):
    final = assessment['final_reader']
    record = final['assessment_narratives'][1]
    record['paragraphs'][0] += ' ' + ' '.join(['supported'] * 110) + '.'
    assert 180 < len(record['paragraphs'][0].split()) < 400
    assert reviewed_final_reader(final, assessment['reader_review'], assessment)[1]
    model = build_reader_model(assessment)
    assert model['assessment_themes'][1]['paragraphs'] == record['paragraphs']


def test_review_notice_is_bound_to_export_snapshot(assessment):
    import app as app_module
    from sector_lenses.climate_verified_render import attach_provenance
    assessment['reader_review'] = review_for(assessment['final_reader'])
    import json
    from sector_lenses.pipeline import normalize_climate_assessment
    assessment = normalize_climate_assessment(json.loads(json.dumps(assessment)))
    reader = attach_provenance(build_reader_model(assessment), assessment)
    pair = {'climate_assessment': assessment, 'climate_reader': reader}
    client = app_module.app.test_client()
    result = client.post('/api/climate-report-snapshot', json=pair)
    assert result.status_code == 200
    assert result.get_json()['publication_approved'] is False
    reader['reader_review_notice'] = ''
    assert client.post('/api/climate-report-snapshot', json=pair).status_code == 409


def test_browser_escapes_and_shows_review_notice():
    import subprocess
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function, _js_escape_helper
    source = INDEX.read_text(encoding='utf-8')
    renderer = '\n'.join(_extract_js_function(source, name) for name in (
        'renderPriorityProjectCycle', 'climateAccessibleRouteValue', 'climateDraftingRouteGate',
        'climateSummaryWatchItems', 'renderClimateVerifiedAssessment'))
    script = _js_escape_helper() + "const renderClimateRelevantGuidance=()=>'';\n" + renderer + """
const html=renderClimateVerifiedAssessment({priorities:[],reader_review_notice:'Review <script>unsafe</script>'});
if(!html.includes('Review &lt;script&gt;unsafe&lt;/script&gt;') || html.includes('<script>unsafe')) throw new Error('Notice missing or unescaped');
const native=renderClimateVerifiedAssessment({priorities:[],reader_contract:'final-reader-v1',executive_readout:'UNREVIEWED_EXECUTIVE_SENTINEL'});
if(native.includes('UNREVIEWED_EXECUTIVE_SENTINEL')) throw new Error('Unreviewed native precursor leaked');
"""
    result = subprocess.run(['node', '-'], input=script, text=True, capture_output=True)
    assert result.returncode == 0, result.stderr


def test_native_reports_do_not_reintroduce_unreviewed_executive(assessment, tmp_path):
    assessment['executive_readout'] = 'UNREVIEWED_EXECUTIVE_SENTINEL'
    model = build_reader_model(assessment)
    assert 'UNREVIEWED_EXECUTIVE_SENTINEL' not in render_reader_html(model)
    assert 'UNREVIEWED_EXECUTIVE_SENTINEL' not in '\n'.join(p.text for p in Document(write_reader_docx(model,tmp_path/'native.docx')).paragraphs)
    assert assessment['executive_readout'] == 'UNREVIEWED_EXECUTIVE_SENTINEL'


def test_summary_view_shows_same_escaped_review_notice():
    import subprocess
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function, _js_escape_helper
    renderer = _extract_js_function(INDEX.read_text(encoding='utf-8'), 'renderClimateVerifiedSummary')
    script = _js_escape_helper() + """
const climateSummaryStrengths=()=>[], climateSummaryPriorityItems=()=>[], renderClimateSummaryGaps=()=>'', renderClimateWatchDisclosure=()=>'', renderClimateRelevantGuidance=()=>'', renderStage3AdvisoryTransition=()=>'';
""" + renderer + """
const html=renderClimateVerifiedSummary({reader_review_notice:'Minor <uncertainty> remains.'});
if(!html.includes('Minor &lt;uncertainty&gt; remains.'))throw new Error('Summary review notice missing');
"""
    result = subprocess.run(['node','-'],input=script,text=True,capture_output=True)
    assert result.returncode == 0, result.stderr


def test_redundant_aggregate_code_does_not_discard_complete_minor_findings(assessment):
    final = assessment['final_reader']
    review = review_for(final)
    review['reason_codes'].append('RATING_INCOHERENT')
    summary, narratives = reviewed_final_reader(final, review, assessment)
    assert summary and len(narratives) == 3
    assessment['reader_review'] = review
    assert 'metadata' in build_reader_model(assessment)['reader_review_notice']
    assert review['verdict'] == 'revise' and 'RATING_INCOHERENT' in review['reason_codes']


def test_new_analysis_rotates_id_after_clearing_old_storage():
    import subprocess
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function
    reset = _extract_js_function(INDEX.read_text(encoding='utf-8'), 'reset')
    script = """
const events=[];
const element={value:'',innerHTML:'',style:{},classList:{remove:()=>{}}};
const document={getElementById:()=>element};
const localStorage={removeItem:key=>events.push('remove:'+key)};
const window={startNewAssessmentSession:()=>events.push('new-id'),scrollTo:()=>{}};
const renderLensSelector=()=>{}, clearUI=()=>{}, updateSessionBar=()=>events.push('bar');
""" + reset + """
reset();
if(events.filter(x=>x==='new-id').length!==1)throw new Error('Assessment ID not rotated once');
if(events.lastIndexOf('remove:fcv_express_lensState')>events.indexOf('new-id'))throw new Error('New session storage cleared');
if(events.indexOf('new-id')>events.indexOf('bar'))throw new Error('Session bar still uses old ID');
"""
    result = subprocess.run(['node','-'],input=script,text=True,capture_output=True)
    assert result.returncode == 0, result.stderr
