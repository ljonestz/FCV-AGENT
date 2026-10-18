"""Targeted human verification preserves prose without manufacturing a pass."""
from copy import deepcopy
import pytest
from docx import Document
from tests.test_climate_proportionate_review import assessment, review_for
from sector_lenses.climate_final_reader import reviewed_final_reader, normalize_reader_paragraphs
from sector_lenses.climate_verified_render import build_reader_model, render_reader_html, write_reader_docx


@pytest.mark.parametrize('identifier', ['SUM-001', 'NARR-design', 'NARR-dividends', 'NARR-inclusion'])
def test_uncertainty_retained_with_local_note(assessment, identifier, tmp_path):
    final = assessment['final_reader']
    review = review_for(final, identifier, 'consequential')
    issue = review['issues'][0]
    issue.update(disposition='verify', verification_note='Confirm whether the activity is planned or underway.')
    assessment['reader_review'] = review
    summary, narratives = reviewed_final_reader(final, review, assessment)
    assert summary and len(narratives) == 3
    model = build_reader_model(assessment)
    note = 'Verify before sharing: Confirm whether the activity is planned or underway.'
    assert note in render_reader_html(model)
    assert note in '\n'.join(p.text for p in Document(write_reader_docx(model, tmp_path/'notes.docx')).paragraphs)
    assert review['verdict'] == 'revise'
    assert 'withheld' not in model['reader_review_notice']


def test_sentence_removal_keeps_unaffected_prose(assessment):
    final = assessment['final_reader']
    record = final['assessment_narratives'][0]
    sentence = record['paragraphs'][0].split('. ')[0] + '.'
    review = review_for(final, 'NARR-design', 'consequential')
    review['issues'][0].update(quoted_text=sentence, disposition='omit_sentence',
                              verification_note='The source contradicts this sentence; the remainder stands alone.')
    summary, narratives = reviewed_final_reader(final, review, assessment)
    retained = next(r for r in narratives if r['theme'] == 'design')['paragraphs'][0]
    assert sentence not in retained
    assert record['paragraphs'][0].split('. ', 1)[1] in retained
    assert 'Sentence withheld' in retained
    assert sentence in final['assessment_narratives'][0]['paragraphs'][0]


@pytest.mark.parametrize('defect', ['bad_disposition', 'blank_note', 'partial_sentence'])
def test_invalid_disposition_cannot_unlock_text(assessment, defect):
    final = assessment['final_reader']
    review = review_for(final, 'NARR-design', 'consequential')
    review['issues'][0].update(disposition='verify', verification_note='Confirm status.')
    if defect == 'bad_disposition': review['issues'][0]['disposition'] = 'publish'
    if defect == 'blank_note': review['issues'][0]['verification_note'] = ''
    if defect == 'partial_sentence':
        review['issues'][0].update(disposition='omit_sentence', quoted_text='The project')
    summary, narratives = reviewed_final_reader(final, review, assessment)
    assert not any(r['theme'] == 'design' for r in narratives)


def test_three_paragraph_answers_are_preserved(assessment):
    final = deepcopy(assessment['final_reader'])
    record = final['assessment_narratives'][0]
    sentences = record['paragraphs'][0].split('. ')
    record['paragraphs'] = [sentences[0]+'.', sentences[1]+'.', '. '.join(sentences[2:])]
    assert len(normalize_reader_paragraphs(final)['assessment_narratives'][0]['paragraphs']) == 3
    assert len(reviewed_final_reader(final, assessment['reader_review'], assessment)[1]) == 3


@pytest.mark.parametrize('disposition,retained', [('verify', True), ('withhold', False), ('bad', False)])
def test_recommendation_rereview_preserves_only_explicit_uncertainty(disposition, retained, tmp_path):
    from tests.test_climate_verified_pipeline import _responses, _arguments, _core_review, FakeClient, PipelineClients, run_verified_climate_pipeline
    from sector_lenses.climate_verified_schemas import CANDIDATE_SCHEMA
    responses = _responses()
    corrected = deepcopy(responses[-1])
    corrected['recommendation_candidates'] = [{k:v for k,v in corrected['recommendation_candidates'][0].items() if k in CANDIDATE_SCHEMA['properties']}]
    quote = corrected['recommendation_candidates'][0]['minimum_action']
    def review(**kwargs):
        return {'verdict':'revise', 'passed_core_ids':[], 'passed_readiness_ids':[],
                'core_reviews':[_core_review('REC-001','unsupported')],
                'reason_codes':['PROJECT_FACT_UNSUPPORTED'], 'object_ids':['REC-001'],
                'issues':[{'object_id':'REC-001','reason_codes':['PROJECT_FACT_UNSUPPORTED'],
                           'explanation':'Confirm the scope against the source.', 'quoted_text':quote,
                           'disposition':disposition, 'verification_note':'Confirm the scope before using this action.'}]}
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses+[corrected],[]), FakeClient([review,review],[])))
    assert bool(result['priorities']) == retained
    if retained:
        assert result['validation']['status'] == 'attention'
        model = build_reader_model(result)
        assert model['recommendation_status'] == 'partial'
        note = 'Verify before sharing: Confirm the scope before using this action.'
        assert note in render_reader_html(model)
        assert note in model['priorities'][0]['minimum_action']
        assert note in '\n'.join(p.text for p in Document(write_reader_docx(model,tmp_path/'rec.docx')).paragraphs)


def test_second_reader_review_marks_without_another_paid_retry():
    from tests.test_climate_verified_pipeline import _responses, _arguments, _final_reader, _native_pass_review, FakeClient, PipelineClients, run_verified_climate_pipeline
    final = _final_reader()
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    review = review_for(final, 'NARR-design', 'consequential')
    review['issues'][0].update(disposition='verify', verification_note='Confirm implementation status against the PID.')
    generation = FakeClient(responses+[final, {'corrections':[{'object_id':'NARR-design','edits':[]}]}],[])
    reviewer = FakeClient([_native_pass_review,review,review],[])
    result = run_verified_climate_pipeline(**_arguments(),clients=PipelineClients(generation,reviewer))
    assert len(result['assessment_narratives']) == 3
    assert len([c for c in generation.calls if c['stage']=='reader_correction']) == 1
    assert len([c for c in reviewer.calls if c['stage']=='reader_review']) == 2
    assert result['reader_review']['verdict'] == 'revise'
    assert 'Verify before sharing' in result['assessment_narratives'][0]['paragraphs'][0]


def test_verification_marker_is_bound_to_packet(assessment):
    import app as app_module
    from sector_lenses.climate_verified_render import attach_provenance
    assessment['reader_review'] = review_for(assessment['final_reader'], 'NARR-design', 'consequential')
    assessment['reader_review']['issues'][0].update(disposition='verify', verification_note='Confirm source scope.')
    import json
    from sector_lenses.pipeline import normalize_climate_assessment
    assessment = normalize_climate_assessment(json.loads(json.dumps(assessment)))
    reader = attach_provenance(build_reader_model(assessment), assessment)
    pair = {'climate_assessment':assessment,'climate_reader':reader}
    client = app_module.app.test_client()
    response = client.post('/api/climate-report-snapshot',json=pair)
    assert response.status_code == 200
    assert response.get_json()['publication_approved'] is False
    reader['assessment_themes'][0]['paragraphs'][0] = reader['assessment_themes'][0]['paragraphs'][0].split(' [Verify before sharing:')[0]
    assert client.post('/api/climate-report-snapshot',json=pair).status_code == 409


def test_recommendation_notes_cannot_become_reader_support(assessment):
    from sector_lenses.climate_final_reader import reader_support_records
    identifier = assessment['priorities'][0]['recommendation_id']
    assessment['recommendation_diagnostics']['verification_issues'] = {identifier:{'disposition':'verify'}}
    assert identifier not in reader_support_records(assessment)


@pytest.mark.parametrize('disposition', ['verify', 'omit_sentence'])
def test_optional_drafting_retains_valid_parent_and_marks_second_review(disposition):
    from tests.test_climate_verified_pipeline import _responses, _arguments, _core_review, FakeClient, PipelineClients, run_verified_climate_pipeline
    responses = _responses()
    original = responses[-1]['recommendation_candidates'][0]
    block = {**original['current_document_drafting'], 'drafting_role':'current_document'}
    block['text'] += ' The project supports traders.'
    quote = 'The project supports traders.'
    repaired = {'drafting_sets':[{'recommendation_id':'REC-001','drafting_blocks':[block]}]}
    initial = {'verdict':'revise','reason_codes':['DRAFTING_SCOPE_UNSUPPORTED'],
               'object_ids':['REC-001::current_document'], 'passed_core_ids':['REC-001'],
               'core_reviews':[_core_review()], 'issues':[{'object_id':'REC-001::current_document',
               'reason_codes':['DRAFTING_SCOPE_UNSUPPORTED'],'explanation':'Check wording.'}]}
    second = deepcopy(initial)
    second['issues'][0].update(disposition=disposition,quoted_text=quote,verification_note='Confirm this drafting against the document.')
    result = run_verified_climate_pipeline(**_arguments(),clients=PipelineClients(
        FakeClient(responses+[repaired],[]),FakeClient([initial,second],[])))
    assert result['priorities'][0]['current_document_drafting'] is not None
    model = build_reader_model(result)
    assert 'Verify before sharing:' in model['priorities'][0]['minimum_action']
    assert 'Verify before sharing:' in model['priorities'][0]['current_document_drafting']['text']


@pytest.mark.parametrize('defect', ['list_disposition', 'missing_severity', 'missing_impact'])
def test_new_disposition_requires_complete_typed_issue(assessment, defect):
    review = review_for(assessment['final_reader'], 'NARR-design', 'consequential')
    issue = review['issues'][0]
    issue.update(disposition='verify', verification_note='Confirm status.')
    if defect == 'list_disposition': issue['disposition'] = ['verify']
    if defect == 'missing_severity': del issue['severity']
    if defect == 'missing_impact': del issue['decision_impact']
    assert reviewed_final_reader(assessment['final_reader'],review,assessment) == ([], [])
