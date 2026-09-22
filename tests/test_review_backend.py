"""Regressions for core research integration and Word watch output."""
import io
from types import SimpleNamespace
import pytest
from docx import Document
import app


def _plan(excerpt="Water supply in Conakry"):
    return app.build_stage1_research_plan([], "Guinea", "Water", [
        {"label": "PROJECT DOCUMENT", "name": "PAD.pdf", "raw_text": excerpt}])


def test_existing_research_call_preserves_provider_sources_and_project_focus():
    calls = []
    citation = {"type": "web_search_result_location", "url": "https://reporter.example/guinea",
                "title": "Guinea water access", "cited_text": "Guinea faces water access disruptions."}
    response = SimpleNamespace(content=[{"type": "text", "text": "Water delivery requires careful planning.",
                                         "citations": [citation]}])
    def create(**kwargs):
        calls.append(kwargs)
        return response
    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    result = app.run_fcv_web_research("Guinea", "Water", client, max_uses=4,
        project_profile={"document_excerpt": "Maneah network extension"})
    assert len(calls) == 1
    assert calls[0]["tools"][0]["max_uses"] == 4
    prompt = calls[0]["messages"][0]["content"]
    assert "Maneah network extension" in prompt
    assert "8–9" not in prompt
    assert citation["url"] in result["brief"]
    assert result["sources"][0]["url"] == citation["url"]


def test_cache_separates_project_focus_and_expires(monkeypatch):
    calls, clock = [], [1000.0]
    def research(*args, **kwargs):
        calls.append(kwargs.get("project_profile"))
        return {"brief": "Source briefing", "sources": [{"url": "https://example.org"}], "status": "sourced"}
    monkeypatch.setattr(app, "run_fcv_web_research", research)
    monkeypatch.setattr(app, "get_research_client", lambda: object())
    monkeypatch.setattr(app.time, "time", lambda: clock[0])
    app._research_cache.clear()
    for _ in range(2):
        list(app._iter_stage1_research(_plan()))
    assert len(calls) == 1
    assert calls[0]["document_excerpt"] == "Water supply in Conakry"
    list(app._iter_stage1_research(_plan("Rural water committees")))
    assert len(calls) == 2
    clock[0] += 6 * 60 * 60 + 1
    list(app._iter_stage1_research(_plan()))
    assert len(calls) == 3


@pytest.mark.parametrize("result", [
    {"brief": "No source evidence", "sources": [], "status": "unavailable"},
    {"brief": "Only model interpretation", "sources": [], "status": "limited"}])
def test_research_without_sources_is_not_cached(monkeypatch, result):
    calls = []
    monkeypatch.setattr(app, "run_fcv_web_research", lambda *a, **k: calls.append(1) or result)
    monkeypatch.setattr(app, "get_research_client", lambda: object())
    app._research_cache.clear()
    for _ in range(2):
        events = list(app._iter_stage1_research(_plan()))
        assert events[-1]["result"]["core_brief"] == result["brief"]
    assert len(calls) == 2


def test_standard_rra_prompt_distinguishes_document_reference_from_context_use():
    for stage in (1, 2):
        prompt = app.append_standard_fcv_stage_context(app.DEFAULT_PROMPTS[str(stage)], stage, [])
        assert "primary document references an RRA" in prompt
        assert "uploaded RRA was used" in prompt
        assert "only when" in prompt


@pytest.mark.parametrize("heading", ["### Watch List for Supervision", "**Watch List for Supervision**", "Watch List for Supervision"])
def test_word_keeps_watch_content_with_single_heading(heading):
    result = app.app.test_client().post("/api/download-report", json={
        "summary": "## Assessment\n\nSynthetic project review.", "priorities": [],
        "horizon_considerations": heading + "\n\n### Access risk\n\nMonitor access in Conakry."})
    assert result.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(result.data)).paragraphs)
    assert text.count("Watch List for Supervision") == 1
    assert "Access risk" in text
    assert "Monitor access in Conakry." in text


def test_word_preserves_all_actions_and_drafting():
    result = app.app.test_client().post("/api/download-report", json={
        "summary": "## Assessment\n\nSynthetic project review.",
        "risk_exposure": {"risks_to": "Supply interruption", "risks_from": "Exclusion risk"},
        "sensitivity_summary": "Sensitivity explanation", "responsiveness_summary": "Responsiveness explanation",
        "fcv_rating": "Adequate", "fcv_responsiveness_rating": "Low",
        "priorities": [{"title": "Protect water access", "the_gap": "Access gap", "why_it_matters": "Delivery risk",
            "actions": [{"document_element": "POM", "guidance": "First complete action", "suggested_language": "First proposed wording"},
                        {"document_element": "RF", "guidance": "Second complete action", "suggested_language": "Second proposed wording"}]}]})
    assert result.status_code == 200
    doc = Document(io.BytesIO(result.data))
    text = "\n".join(node.text or "" for node in doc.element.iter() if node.tag.endswith("}t"))
    for expected in ("First complete action", "Second complete action", "First proposed wording", "Second proposed wording", "Supply interruption", "Exclusion risk", "Sensitivity explanation", "Responsiveness explanation"):
        assert expected in text



def test_watch_heading_rejects_large_near_match_in_bounded_time():
    """Untrusted report text must not trigger quadratic regex backtracking."""
    import subprocess
    import sys
    code = (
        "from fcv_presentation import strip_watch_heading; "
        "text='Watch List for Supervision'+'\\t'*100000+'not a heading\\nBody'; "
        "assert strip_watch_heading(text)==text"
    )
    subprocess.run([sys.executable, "-c", code], check=True, timeout=8)
