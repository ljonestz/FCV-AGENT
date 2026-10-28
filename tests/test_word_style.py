"""Focused presentation checks for editable FCV Word exports."""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
import pytest

from fcv_docx_style import AMBER_TXT, BE_AWARE, CONTENT_W, METER_OFF, Report, logo_path
from fcv_management_brief import render_management_brief_docx


def _brief():
    readout = {
        "headline": "Clearer access decisions would improve delivery.",
        "overview": "The project identifies access constraints and proposes practical follow-up.",
        "strengths": [{"title": "Context", "text": "The design recognizes local constraints."}],
    }
    priorities = [
        {
            "number": 1,
            "title": "Clarify access",
            "concise": {
                "title": "Clarify access ownership",
                "why": "Unclear ownership may delay delivery.",
                "how": ["Assign a named decision owner."],
            },
        }
    ]
    return readout, priorities


def _fill(element, tag):
    child = element.find(qn(tag))
    if child is None:
        return None
    return child.get(qn("w:fill"))


def test_management_brief_uses_editable_navy_word_treatment():
    readout, priorities = _brief()
    document = Document(BytesIO(render_management_brief_docx(readout, priorities)))

    assert document.styles["Normal"].font.name == "Arial"
    assert document.styles["Title"].font.color.rgb == RGBColor(0, 0, 0)
    assert document.styles["Heading 1"].font.color.rgb == RGBColor(0x15, 0x39, 0x56)
    assert document.styles["Heading 1"].paragraph_format.keep_with_next is True
    assert all(
        paragraph.alignment is None
        or paragraph.alignment.value == 0
        for paragraph in document.paragraphs
    )

    header = document.sections[0].header.paragraphs[0]
    assert "FCV MANAGEMENT BRIEF" in header.text
    assert _fill(header._p.pPr, "w:shd") == "153956"

    footer = document.sections[0].footer.paragraphs[0]
    assert "Page" in footer.text
    assert any(
        instruction.text.strip() == "PAGE"
        for instruction in footer._p.xpath(".//w:instrText")
    )

    action = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Suggested action:")
    )
    assert _fill(action._p.pPr, "w:shd") == "E7F3F1"
    assert action.paragraph_format.keep_together is not True


def test_brief_keeps_a4_geometry_and_native_priority_flow():
    readout, priorities = _brief()
    document = Document(BytesIO(render_management_brief_docx(readout, priorities)))
    section = document.sections[0]

    assert section.page_width.mm == pytest.approx(210, abs=0.1)
    assert section.page_height.mm == pytest.approx(297, abs=0.1)
    assert section.left_margin.mm == pytest.approx(17, abs=0.1)
    assert section.right_margin.mm == pytest.approx(17, abs=0.1)
    assert "Clarify access ownership" in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )


def _xml_text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def test_report_components_keep_brand_advisory_scorecard_links_and_qualifications():
    meta = {
        "title": "FCV Recommendations Note",
        "project_line": "Supplied project | Somalia | Exported: 2026-10-27",
        "running_label": "Somalia",
        "classification": "",
        "generator": "",
        "disclaimer": "",
        "its_url": "",
    }
    report = Report("climate", meta)
    report.banner()
    report.be_aware(BE_AWARE)
    report.scorecard(
        [
            {
                "label": "FCV Sensitivity",
                "value": "Unclear",
                "n": 6,
                "idx": 0,
                "lo": "Very low",
                "hi": "Very high",
            },
            {
                "label": "FCV Responsiveness",
                "value": "Not expected for this project",
                "n": 6,
                "idx": 0,
                "lo": "Very low",
                "hi": "Very high",
            },
        ],
        {"label": "Operational priorities", "number": 0, "lines": []},
    )
    report.h1("Assessment findings")
    report.body("The design records the main access constraints.")
    report.h1("Climate and FCV Sensitivity and Responsiveness")
    report.rating_segment("FCV Sensitivity", "Unclear", first=True)
    report.rating_segment("FCV Responsiveness", "Not expected for this project")
    report.h1("Ranked operational priorities")
    report.data_table(
        ["#", "Operational priority"],
        [0.35, CONTENT_W - 0.35],
        [["1", ("Clarify access", "priority_1")]],
        link_col=1,
    )
    report.priority_header(
        1,
        "Clarify access",
        [("Suggested lead", "Task team")],
        anchor="priority_1",
    )
    grey_note = "[Verify before sharing: confirm this note against the source.]"
    draft_note = "[Verify before sharing: confirm the drafting detail with the team.]"
    report.grey_box("Practical considerations", [f"Source qualification. {grey_note}"])
    report.drafting_box(f"Suggested wording. {draft_note}")

    document = report.doc
    section = document.sections[0]
    first_header = section.first_page_header
    header_text = _xml_text(first_header._element)
    body_elements = list(document.element.body.iterchildren())

    def body_index(marker, after=-1):
        return next(
            index
            for index, element in enumerate(body_elements)
            if index > after and marker in _xml_text(element)
        )

    assert section.header_distance.inches == pytest.approx(0)
    banner_position = first_header.tables[0]._tbl.tblPr.find(qn("w:tblpPr"))
    assert banner_position.get(qn("w:tblpX")) == "0"
    assert banner_position.get(qn("w:tblpY")) == "0"
    assert "FCV Recommendations Note" in header_text
    assert "Supplied project | Somalia | Exported: 2026-10-27" in header_text
    logo_bytes = Path(logo_path(True)).read_bytes()
    assert any(
        relationship.target_part.blob == logo_bytes
        for relationship in first_header.part.rels.values()
        if relationship.reltype.endswith("/image")
    )
    assert "FCV SCREENER" in _xml_text(section.header._element)
    assert "AI-assisted output" in _xml_text(section.footer._element)

    scorecard = next(
        table
        for table in document.tables
        if len(table.columns) == 3
        and "FCV SENSITIVITY" in _xml_text(table._tbl)
        and "FCV RESPONSIVENESS" in _xml_text(table._tbl)
    )
    assert "Unclear" in _xml_text(scorecard._tbl)
    assert "Not expected for this project" in _xml_text(scorecard._tbl)
    assert "0" in _xml_text(scorecard._tbl)
    meter = scorecard.cell(0, 0).tables[0]
    meter_fills = [
        cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill"))
        for cell in meter.rows[0].cells
    ]
    assert meter_fills == [METER_OFF] * 6

    advisory_index = body_index(BE_AWARE[0])
    scorecard_index = body_index("FCV SENSITIVITY")
    findings_index = body_index("Assessment findings")
    rating_section_index = body_index("Climate and FCV Sensitivity and Responsiveness")
    sensitivity_detail_index = body_index("FCV SENSITIVITY", after=rating_section_index)
    responsiveness_detail_index = body_index(
        "FCV RESPONSIVENESS", after=rating_section_index
    )
    priorities_index = body_index("Ranked operational priorities")
    assert advisory_index < scorecard_index < findings_index
    assert (
        findings_index
        < rating_section_index
        < sensitivity_detail_index
        < responsiveness_detail_index
        < priorities_index
        < body_index("Operational priority")
        < body_index("PRIORITY 1")
    )
    bookmark_names = {
        item.get(qn("w:name"))
        for item in document.element.body.iter(qn("w:bookmarkStart"))
    }
    assert "priority_1" in bookmark_names
    assert any(
        item.get(qn("w:anchor")) == "priority_1"
        for item in document.element.body.iter(qn("w:hyperlink"))
    )

    for marker in (grey_note, draft_note):
        highlighted_runs = []
        for run in document.element.body.iter(qn("w:r")):
            text = "".join(node.text or "" for node in run.iter(qn("w:t")))
            if marker in text:
                run_properties = run.find(qn("w:rPr"))
                color = run_properties.find(qn("w:color")) if run_properties is not None else None
                highlighted_runs.append(color.get(qn("w:val")) if color is not None else None)
        assert highlighted_runs == [AMBER_TXT]
