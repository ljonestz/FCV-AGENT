from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, replace
import json
import pytest

from sector_lenses.climate_analysis import ContextEvidenceRef
from sector_lenses.climate_source_blocks import (
    DocumentApplicability,
    SourceBlock,
    SourceDocument,
)
from sector_lenses.climate_verified_pipeline import (
    _candidate,
    _valid_reader_revision,
    PipelineClients,
    run_verified_climate_pipeline,
)


@dataclass
class FakeClient:
    responses: list[dict[str, object]]
    calls: list[dict[str, object]]

    @property
    def last_output_tokens(self):
        return 0  # Fake clients incur no provider output-token charge.

    def complete_json(self, **kwargs):
        self.calls.append(kwargs)
        response = self.responses.pop(0)
        return response(**kwargs) if callable(response) else response



def _native_pass_review(**kwargs):
    """Fake verifier explicitly accepting each supplied core ID in native JSON.

    FakeClient returns malformed dictionaries unchanged, so missing required
    fields still fail the real acceptance gate and remain testable.
    """
    result = {'verdict': 'pass', 'reason_codes': [], 'object_ids': [], 'issues': []}
    if kwargs['stage'] == 'conditional_review':
        recommendations = kwargs['payload']['recommendations']
        block_id = kwargs['payload']['source_blocks'][0]['block_id']
        result['passed_core_ids'] = [item['recommendation_id'] for item in recommendations]
        result['core_reviews'] = [{
            'recommendation_id': item['recommendation_id'],
            'premise_verdict': 'supported',
            'source_block_ids': [block_id],
            'reason_codes': [],
            'explanation': 'The cited source supports the bounded premise.',
        } for item in recommendations]
    return result

def _pass_review_client() -> FakeClient:
    return FakeClient(
        [_native_pass_review],
        [],
    )


def _core_review(identifier="REC-001", verdict="supported", *, codes=None):
    return {
        "recommendation_id": identifier,
        "premise_verdict": verdict,
        "source_block_ids": ["DOC-1-B-1"],
        "reason_codes": (
            ["PROJECT_FACT_UNSUPPORTED"] if verdict == "unsupported" else []
        ) if codes is None else codes,
        "explanation": "The source comparison supports this premise decision.",
    }


READER_QUOTE = "Verdict and foundation: The project recognizes"


def _source() -> tuple[list[SourceDocument], list[SourceBlock]]:
    document = SourceDocument(
        document_id="DOC-1",
        filename="pcn.docx",
        sha256="abc123",
        applicability=DocumentApplicability.VERIFIED,
        relationship="primary",
        version_status="latest",
        operation_match="verified",
    )
    block = SourceBlock(
        block_id="DOC-1-B-1",
        document_id="DOC-1",
        text="The Project Operations Manual will define site selection.",
        normalized_hash="block123",
        heading_path=("Components",),
        paragraph_index=1,
    )
    return [document], [block]


def _responses(*, unresolved_routing: bool = False):
    routing = "team_to_confirm" if unresolved_routing else "verified_existing"
    instrument_ids = [] if unresolved_routing else ["PF-001"]
    return [
        {
            "facts": [
                {
                    "claim_id": "PF-001",
                    "claim_type": "named_instrument",
                    "subject": "Project Operations Manual",
                    "predicate": "will define",
                    "object": "site selection",
                    "epistemic_status": "explicit",
                    "source_block_ids": ["DOC-1-B-1"],
                    "supporting_excerpt": (
                        "The Project Operations Manual will define site selection."
                    ),
                    "confidence": "high",
                }
            ],
            "derived_assertions": [],
        },
        {
            "existing_responses": [],
            "pathways": [
                {
                    "pathway_id": "PW-001",
                    "direction": "climate_to_fcv",
                    "chain": ["flood", "access disruption", "exclusion"],
                    "project_anchor_ids": ["PF-001"],
                    "evidence_ids": ["CE-001"],
                    "confidence": "medium",
                }
            ],
            "residual_gaps": [
                {
                    "gap_id": "RG-001",
                    "gap_type": "not_yet_specified",
                    "statement": "Site criteria are not yet specified.",
                    "pathway_ids": ["PW-001"],
                    "project_anchor_ids": ["PF-001"],
                    "existing_response_ids": [],
                    "evidence_ids": ["CE-001"],
                    "confidence": "medium",
                }
            ],
        },
        {
            "executive_readout": (
                "Verified project facts support a bounded Climate-FCV readout. " * 45
            ).strip(),
            "relevance": {
                "value": "high",
                "evidence_ids": ["PW-001"],
                "rationale": "A material pathway is connected to project design.",
            },
            "sensitivity": {
                "value": "limited",
                "evidence_ids": ["RG-001"],
                "rationale": "The residual gap remains material.",
            },
            "responsiveness": {
                "value": "not_expected",
                "evidence_ids": ["PF-001"],
                "rationale": "Responsiveness is not required for this design choice.",
            },
            "operationalization": {
                "value": "early",
                "evidence_ids": ["PF-001"],
                "rationale": "An instrument is named but its requirements are open.",
            },
        },
        {
            "recommendation_candidates": [
                {
                    "recommendation_id": "REC-001",
                    "title": "Define climate-FCV site-selection criteria",
                    "pathway_ids": ["PW-001"],
                    "existing_response_ids": [],
                    "residual_gap_ids": ["RG-001"],
                    "project_anchor_ids": ["PF-001"],
                    "decision": "Decide how site selection will address the pathway.",
                    "minimum_action": "Document exposure and access considerations.",
                    "enhanced_action": None,
                    "enhanced_activation": None,
                    "routing_status": routing,
                    "instrument_claim_ids": instrument_ids,
                    "responsible_function": "Task team",
                    "authority_basis": "project_commitment",
                    "recommendation_basis": "project_evidence",
                    "completion_evidence": "Updated site-selection method",
                    "completion_evidence_status": "updated_section",
                    "confidence": "medium",
                    "limitation": "Site-level conditions remain to be verified.",
                    "caution": "Avoid restricting necessary seasonal mobility.",
                    "narrative": "The project needs an evidenced site-selection method.",
                    "supported_numeric_tokens": [],
                    "current_document_drafting": {
                        "target_document": "PCN",
                        "target_section": "Project Description",
                        "drafting_status": "existing_commitment",
                        "text": ("Add bounded project design language here. " * 18).strip(),
                        "project_basis_ids": ["PF-001"],
                        "gap_basis_ids": ["RG-001"],
                        "guidance_ids": ["GUIDE-PCN-DESIGN"],
                    },
                    "operational_instrument_drafting": None,
                    "score": {
                        "materiality": 3,
                        "gap_strength": 2,
                        "leverage_urgency": 2,
                        "evidence": 2,
                        "feasibility": 1,
                    },
                    "gate_results": {
                        "connection": True,
                        "residuality": True,
                        "materiality": True,
                        "actionability": True,
                        "timing": True,
                        "distinctiveness": True,
                    },
                }
            ],
            "readiness_flags": [],
        },
    ]


def _arguments():
    documents, blocks = _source()
    context = [
        ContextEvidenceRef(
            evidence_id="CE-001",
            evidence_class="country",
            scope="national",
            statement="Flooding can disrupt access.",
            source_ref="bank:ssd-2026.08:E-001",
            confidence="medium",
        )
    ]
    return {
        "source_documents": documents,
        "source_blocks": blocks,
        "context_evidence": context,
        "bank_release_id": "2026.08",
        "run_id": "run-test",
        "doc_type": "PCN",
        "instrument_type": "IPF",
    }


def test_four_calls_run_when_semantic_review_is_not_required():
    assessment = FakeClient(_responses(), [])
    reviewer = FakeClient(
        [_native_pass_review],
        [],
    )
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, reviewer),
    )

    assert [call["stage"] for call in assessment.calls] == [
        "fact_extraction",
        "bounded_analysis",
        "judgment_review",
        "recommendation_compiler",
    ]
    assert [call["timeout_seconds"] for call in assessment.calls] == [
        300,
        180,
        240,
        240,
    ]
    assert all(call["max_transient_retries"] == 1 for call in assessment.calls)
    assert [call["stage"] for call in reviewer.calls] == ["conditional_review"]
    assert result["schema_version"] == "climate-verified-v2.1"
    assert result["validation"]["status"] == "attention"
    assert "SUMMARY_OVERVIEW_FALLBACK_REQUIRED" in result["validation"]["reason_codes"]
    assert len(result["priorities"]) == 1
    compiler_payload = assessment.calls[-1]["payload"]
    assert compiler_payload["guidance_registry_version"] == "climate-guidance-v4"
    assert "GUIDE-PCN-DESIGN" in {
        item["guidance_id"] for item in compiler_payload["operational_guidance"]
    }
    assert result["executive_readout"].startswith("Verified project facts")


def test_operation_context_is_passed_to_analysis_and_recommendation_prompts():
    assessment = FakeClient(_responses(), [])
    context = {
        "document_type": "Program Paper",
        "instrument_type": "PforR",
        "country_scope": "single",
        "is_mpa": False,
        "has_ipf_component": False,
        "preparation_regime": "new_model",
        "processing_model": "two_step",
        "es_regime": "INSTRUMENT_SPECIFIC",
        "warning_codes": [],
    }

    arguments = _arguments()
    arguments.update(doc_type="Program Paper", instrument_type="PforR")
    result = run_verified_climate_pipeline(
        **arguments,
        clients=PipelineClients(assessment, _pass_review_client()),
        operation_context=context,
    )

    prompt_contexts = [
        call["payload"].get("operation_context")
        for call in assessment.calls
        if call["stage"] in {
            "bounded_analysis", "judgment_review", "recommendation_compiler"
        }
    ]
    assert prompt_contexts == [context, context, context]
    assert result["operation_context"] == context


def test_guidance_backed_section_variants_do_not_suppress_all_candidates():
    responses = _responses()
    template = responses[3]["recommendation_candidates"][0]
    variants = (
        (
            "Project Description and Components",
            "GUIDE-PCN-DESIGN",
        ),
        (
            "Project Design Narrative",
            "GUIDE-PCN-DESIGN",
        ),
        (
            "Implementation and Institutional Arrangements",
            "GUIDE-FCV-CONTINUITY",
        ),
        (
            "Environmental and Social Risk Management",
            "GUIDE-ES-INSTRUMENT-ROUTING",
        ),
    )
    candidates = []
    for index, (section, guidance_id) in enumerate(variants, start=1):
        candidate = deepcopy(template)
        candidate["recommendation_id"] = f"REC-{index:03d}"
        candidate["title"] = f"Bounded recommendation {index}"
        candidate["current_document_drafting"]["target_section"] = section
        candidate["current_document_drafting"]["guidance_ids"] = [guidance_id]
        candidates.append(candidate)
    responses[3]["recommendation_candidates"] = candidates

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(
            FakeClient(responses, []),
            _pass_review_client(),
        ),
    )

    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["raw_candidate_count"] == 4
    assert diagnostics["valid_candidate_count"] == 4
    assert diagnostics["final_priority_count"] == 4
    assert diagnostics["reason_codes"] == []
    assert "DRAFTING_CURRENT_SECTION_CANONICALIZED" in (
        result["manifest"]["repair_actions"]
    )


def test_dpf_pid_candidates_use_current_concept_document_targets():
    responses = _responses()
    template = responses[3]["recommendation_candidates"][0]
    target_sections = (
        "Concept Description",
        "Poverty and Social Impacts and Environmental Aspects",
        "Proposed Development Objectives",
        "Concept Description",
    )
    candidates = []
    for index, section in enumerate(target_sections, start=1):
        candidate = deepcopy(template)
        candidate["recommendation_id"] = f"REC-{index:03d}"
        candidate["title"] = f"Strengthen a bounded DPF concept decision {index}"
        candidate["current_document_drafting"]["target_document"] = "PID"
        candidate["current_document_drafting"]["target_section"] = section
        candidate["current_document_drafting"]["guidance_ids"] = [
            "GUIDE-DPF-POLICY-PROGRAM"
        ]
        candidates.append(candidate)
    responses[3]["recommendation_candidates"] = candidates
    arguments = _arguments()
    arguments.update(doc_type="PID", instrument_type="DPF")

    result = run_verified_climate_pipeline(
        **arguments,
        clients=PipelineClients(
            FakeClient(responses, []),
            _pass_review_client(),
        ),
    )

    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["raw_candidate_count"] == 4
    assert diagnostics["valid_candidate_count"] == 4
    assert diagnostics["final_priority_count"] == 4
    assert "DRAFTING_CURRENT_TARGET_INVALID" not in diagnostics["reason_codes"]


def test_unresolved_routing_triggers_one_source_first_review():
    assessment = FakeClient(_responses(), [])
    reviewer = FakeClient(
        [
            {
                "verdict": "revise",
                "reason_codes": ["ROUTING_SCOPE_UNVERIFIED"],
                "object_ids": ["REC-001"],
                "passed_core_ids": [],
                "core_reviews": [_core_review("REC-001", "unsupported",
                                               codes=["ROUTING_SCOPE_UNVERIFIED"])]}
        ],
        [],
    )
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, reviewer),
    )

    assert len(reviewer.calls) == 1
    assert reviewer.calls[0]["stage"] == "conditional_review"
    assert reviewer.calls[0]["timeout_seconds"] == 120
    assert "source_blocks" in reviewer.calls[0]["payload"]
    review_payload = reviewer.calls[0]["payload"]
    compiler_payload = next(call["payload"] for call in assessment.calls
                            if call["stage"] == "recommendation_compiler")
    assert review_payload["operational_guidance"] == compiler_payload["operational_guidance"]
    assert review_payload["guidance_registry_version"] == compiler_payload["guidance_registry_version"]
    assert review_payload["current_document"] == compiler_payload["operation_context"]["document_type"]
    assert result["validation"]["status"] == "attention"
    assert result["priorities"] == []
    assert result["manifest"]["suppressed_counts"]["recommendations"] == 1
    assert result["recommendation_diagnostics"] == {
        "raw_candidate_count": 1,
        "parsed_candidate_count": 1,
        "valid_candidate_count": 1,
        "admitted_count": 1,
        "final_priority_count": 0,
        "reviewer_invoked": True,
        "reviewer_verdict": "revise",
        "reason_codes": [
            "ROUTING_SCOPE_UNVERIFIED",
            "RECOMMENDATIONS_ALL_SUPPRESSED",
        ],
        "unsupported_numeric_tokens": [],
        "semantic_review_object_ids": ["REC-001"],
        "candidate_suppressions": [
            {
                "recommendation_id": "REC-001",
                "stage": "semantic_review",
                "reason_codes": ["ROUTING_SCOPE_UNVERIFIED"],
                "unsupported_numeric_fields": [],
            }
        ],
    }


def test_semantic_review_targets_only_affected_recommendations():
    responses = _responses()
    second = deepcopy(responses[3]["recommendation_candidates"][0])
    second["recommendation_id"] = "REC-002"
    second["title"] = "Confirm a second bounded site decision"
    responses[3]["recommendation_candidates"].append(second)
    reviewer = FakeClient(
        [
            {
                "verdict": "revise",
                "reason_codes": ["ROUTING_SCOPE_UNVERIFIED"],
                "object_ids": ["REC-001"],
                "passed_core_ids": ["REC-002"],
                "core_reviews": [
                    _core_review("REC-001", "unsupported", codes=["ROUTING_SCOPE_UNVERIFIED"]),
                    _core_review("REC-002"),
                ]}
        ],
        [],
    )

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(responses, []), reviewer),
    )

    assert [item["recommendation_id"] for item in result["priorities"]] == [
        "REC-002"
    ]
    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["semantic_review_object_ids"] == ["REC-001"]
    assert diagnostics["candidate_suppressions"] == [
        {
            "recommendation_id": "REC-001",
            "stage": "semantic_review",
            "reason_codes": ["ROUTING_SCOPE_UNVERIFIED"],
            "unsupported_numeric_fields": [],
        }
    ]


def test_semantic_review_reason_codes_are_bounded():
    reviewer = FakeClient(
        [
            {
                "verdict": "revise",
                "reason_codes": ["routing scope needs review"],
                "object_ids": ["REC-001"],
                "passed_core_ids": [],
                "core_reviews": [_core_review("REC-001", "unsupported",
                                               codes=["ROUTING_SCOPE_UNVERIFIED"])]}
        ],
        [],
    )
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(
            FakeClient(_responses(), []),
            reviewer,
        ),
    )

    assert result["recommendation_diagnostics"]["reason_codes"] == [
        "SEMANTIC_REVIEW_REASON_INVALID",
        "RECOMMENDATIONS_ALL_SUPPRESSED",
    ]


def test_semantic_review_with_unresolved_targets_fails_safe():
    reviewer = FakeClient(
        [
            {
                "verdict": "revise",
                "reason_codes": ["ROUTING_SCOPE_UNVERIFIED"],
                "object_ids": ["RG-001"],
                "passed_core_ids": [],
                "core_reviews": [_core_review("REC-001", "unsupported",
                                               codes=["ROUTING_SCOPE_UNVERIFIED"])]}
        ],
        [],
    )

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(
            FakeClient(_responses(), []),
            reviewer,
        ),
    )

    assert result["priorities"] == []
    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["semantic_review_object_ids"] == ["REC-001"]
    assert "SEMANTIC_REVIEW_TARGET_UNRESOLVED" in diagnostics["reason_codes"]


@pytest.mark.parametrize("verdict", ["revise", "block"])
def test_negative_review_without_any_target_cannot_certify_core(verdict):
    review = {
        "verdict": verdict,
        "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
        "object_ids": [],
        "issues": [],
        "passed_core_ids": ["REC-001"],
        "core_reviews": [_core_review()],
    }

    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(
            FakeClient(_responses(), []), FakeClient([review], []),
        ),
    )

    assert result["priorities"] == []
    assert "SEMANTIC_REVIEW_TARGET_UNRESOLVED" in result["validation"]["reason_codes"]


def test_admission_suppression_exposes_bounded_reason_codes():
    responses = _responses()
    recommendation = responses[3]["recommendation_candidates"][0]
    recommendation["score"]["materiality"] = 1
    recommendation["gate_results"]["timing"] = False

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(responses, []), FakeClient([], [])),
    )

    assert result["priorities"] == []
    assert result["recommendation_diagnostics"] == {
        "raw_candidate_count": 1,
        "parsed_candidate_count": 1,
        "valid_candidate_count": 1,
        "admitted_count": 0,
        "final_priority_count": 0,
        "reviewer_invoked": False,
        "reviewer_verdict": "not_invoked",
            "reason_codes": [
                "ADMISSION_MATERIALITY_BELOW_MIN",
                "ADMISSION_GATE_FAILED_TIMING",
                "RECOMMENDATIONS_ALL_SUPPRESSED",
            ],
        "unsupported_numeric_tokens": [],
        "semantic_review_object_ids": [],
        "candidate_suppressions": [
            {
                "recommendation_id": "REC-001",
                "stage": "admission",
                "reason_codes": [
                    "ADMISSION_MATERIALITY_BELOW_MIN",
                    "ADMISSION_GATE_FAILED_TIMING",
                ],
                "unsupported_numeric_fields": [],
            }
        ],
    }


def test_invalid_optional_target_preserves_independently_reviewed_decision():
    responses = _responses()
    responses[3]['recommendation_candidates'][0]['current_document_drafting']['target_section'] = 'Procurement Strategy'
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )
    assert len(result['priorities']) == 1
    assert result['priorities'][0]['current_document_drafting'] is None
    assert result['diagnostic_trace']['drafting_suppressions'][0]['reason_codes'] == ['DRAFTING_CURRENT_TARGET_INVALID']


def test_no_generated_candidates_does_not_claim_all_were_suppressed():
    responses = _responses()
    responses[3]["recommendation_candidates"] = []

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(
            FakeClient(responses, []),
            FakeClient([], []),
        ),
    )

    assert result["recommendation_diagnostics"]["raw_candidate_count"] == 0
    assert "RECOMMENDATIONS_ALL_SUPPRESSED" not in (
        result["validation"]["reason_codes"]
    )
    assert "RECOMMENDATIONS_ALL_SUPPRESSED" not in (
        result["recommendation_diagnostics"]["reason_codes"]
    )


def test_source_linked_numeric_label_is_supported_without_model_echo():
    responses = _responses()
    fact = responses[0]["facts"][0]
    excerpt = (
        "The Project Operations Manual will define Component 1.4 site selection."
    )
    fact["object"] = "Component 1.4 site selection"
    fact["supporting_excerpt"] = excerpt
    recommendation = responses[3]["recommendation_candidates"][0]
    recommendation["decision"] = "Confirm the Component 1.4 site-selection method."

    arguments = _arguments()
    arguments["source_blocks"][0] = replace(
        arguments["source_blocks"][0],
        text=excerpt,
    )
    result = run_verified_climate_pipeline(
        **arguments,
        clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )

    assert len(result["priorities"]) == 1
    assert (
        result["recommendation_diagnostics"]["unsupported_numeric_tokens"] == []
    )


def test_model_cannot_self_attest_an_unsourced_numeric_token():
    responses = _responses()
    recommendation = responses[3]["recommendation_candidates"][0]
    recommendation["decision"] = "Complete the review in 2027."
    recommendation["supported_numeric_tokens"] = ["2027"]

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )

    assert result["priorities"][0]["decision"] == "Complete the review."
    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["unsupported_numeric_tokens"] == []
    assert "RECOMMENDATION_UNSUPPORTED_PRECISION_REMOVED" in (
        result["manifest"]["repair_actions"]
    )


def test_numeric_repair_removes_only_bounded_unsupported_tokens():
    responses = _responses()
    recommendation = responses[3]["recommendation_candidates"][0]
    recommendation["minimum_action"] = (
        "Update Components 1 and 2 before the 2027 review."
    )

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )

    assert result["priorities"][0]["minimum_action"] == (
        "Update Components before the review."
    )
    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["reason_codes"] == []
    assert diagnostics["unsupported_numeric_tokens"] == []
    assert diagnostics["candidate_suppressions"] == []


def test_bad_fact_suppresses_dependent_analysis_and_recommendation():
    responses = _responses()
    responses[0]["facts"][0]["supporting_excerpt"] = "Invented excerpt"
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(responses, []), FakeClient([], [])),
    )

    assert result["facts"] == []
    assert result["analysis"]["pathways"] == []
    assert result["priorities"] == []
    assert "FACT_SOURCE_UNRESOLVED" in result["validation"]["reason_codes"]


def test_unsupported_drafting_system_is_withheld_without_removing_decision():
    responses = _responses()
    responses[3]['recommendation_candidates'][0]['current_document_drafting']['text'] = (
        'A hydrometeorological system will establish continuity alerts and procedures for access disruption. ' * 8
    ).strip()
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )
    assert len(result['priorities']) == 1
    assert result['priorities'][0]['current_document_drafting'] is None
    assert result['recommendation_diagnostics']['candidate_suppressions'] == []
    assert result['diagnostic_trace']['drafting_suppressions'][0]['reason_codes'] == ['DRAFTING_SYSTEM_UNVERIFIED']


def test_manifest_is_privacy_safe_and_scoped_to_the_run():
    first = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(FakeClient(_responses(), []), _pass_review_client()),
    )
    second_args = {**_arguments(), "run_id": "run-two"}
    second = run_verified_climate_pipeline(
        **second_args,
        clients=PipelineClients(FakeClient(_responses(), []), _pass_review_client()),
    )

    assert first["manifest"]["run_id"] == "run-test"
    assert second["manifest"]["run_id"] == "run-two"
    assert first["manifest"]["source_count"] == 1
    assert first["manifest"]["renderer_version"] == "climate-reader-v2.20"
    assert first["manifest"]["extraction_version"] == "source-blocks-v3"
    assert (
        first["manifest"]["prompt_versions"]["fact_extraction"]
        == "climate-facts-v2.9"
    )
    assert (
        first["manifest"]["prompt_versions"]["bounded_analysis"]
        == "climate-analysis-v2.12"
    )
    assert (
        first["manifest"]["prompt_versions"]["judgment_review"]
        == "climate-judgments-v2.16"
    )
    assert (
        first["manifest"]["prompt_versions"]["recommendation_compiler"]
        == "climate-recommendations-v2.22"
    )
    assert (
        first["manifest"]["prompt_versions"]["conditional_review"]
        == "climate-review-v2.21"
    )
    assert (
        first["manifest"]["prompt_versions"]["drafting_compiler"]
        == "climate-drafting-v1.7"
    )
    assert set(first["manifest"]["prompt_versions"]) == {
        "fact_extraction",
        "bounded_analysis",
        "judgment_review",
        "recommendation_compiler",
        "drafting_compiler",
        "conditional_review",
        "reader_compiler",
        "reader_correction",
        "reader_review",
    }
    manifest_text = str(first["manifest"])
    assert "Project Operations Manual will define" not in manifest_text
    assert "prompt_text" not in manifest_text


def test_preview_label_survives_into_reader_payload():
    args = _arguments()
    args["context_evidence"] = [
        ContextEvidenceRef(
            evidence_id="CE-001",
            evidence_class="country",
            scope="national",
            statement="Flooding can disrupt access.",
            source_ref="bank-preview:2026.08:E-001",
            confidence="medium",
            preview_status="preview; not approved",
        )
    ]
    result = run_verified_climate_pipeline(
        **args,
        clients=PipelineClients(FakeClient(_responses(), []), _pass_review_client()),
    )

    assert result["evidence_status"] == "preview; not approved"


def test_candidate_maps_compact_drafting_blocks_to_domain_fields():
    record = deepcopy(_responses()[3]["recommendation_candidates"][0])
    current = record.pop("current_document_drafting")
    record.pop("operational_instrument_drafting")
    record["drafting_blocks"] = [
        {"drafting_role": "current_document", **current}
    ]

    candidate = _candidate(record)

    assert candidate.current_document_drafting is not None
    assert (
        candidate.current_document_drafting.target_section
        == "Project Description"
    )
    assert candidate.operational_instrument_drafting is None


def test_missing_transport_drafting_uses_bounded_drafting_compiler():
    responses = _responses()
    recommendation = responses[3]["recommendation_candidates"][0]
    current = recommendation.pop("current_document_drafting")
    recommendation.pop("operational_instrument_drafting")
    responses.append({
        "drafting_sets": [{
            "recommendation_id": "REC-001",
            "drafting_blocks": [{
                "drafting_role": "current_document",
                **current,
            }],
        }],
    })
    assessment = FakeClient(responses, [])

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )

    assert [call["stage"] for call in assessment.calls] == [
        "fact_extraction",
        "bounded_analysis",
        "judgment_review",
        "recommendation_compiler",
        "drafting_compiler",
    ]
    assert result["priorities"][0]["current_document_drafting"]["text"] == (
        current["text"]
    )
    drafting_payload = assessment.calls[-1]["payload"]
    assert drafting_payload["recommendation_candidates"][0][
        "recommendation_id"
    ] == "REC-001"


def test_unresolved_instrument_route_preserves_actions_and_withholds_drafting():
    arguments = _arguments()
    arguments["instrument_type"] = "Unknown"

    result = run_verified_climate_pipeline(
        **arguments,
        clients=PipelineClients(
            FakeClient(_responses(), []),
            _pass_review_client(),
        ),
    )

    diagnostics = result["recommendation_diagnostics"]
    assert diagnostics["parsed_candidate_count"] == 1
    assert diagnostics["valid_candidate_count"] == 1
    assert diagnostics["final_priority_count"] == 1
    assert "DRAFTING_CURRENT_TARGET_INVALID" not in diagnostics["reason_codes"]
    assert result["priorities"][0]["current_document_drafting"] is None
    assert "DRAFTING_CURRENT_UNRESOLVED_ROUTE_DROPPED" in (
        result["manifest"]["repair_actions"]
    )

def _summary_payload():
    return {
        "paragraphs": [
            (
                "Verdict and foundation: The project recognizes the climate-FCV "
                "interaction and the overall takeaway is credible. "
                + "The Project Operations Manual will define site selection. " * 12
            ).strip(),
            (
                "Four-dimensional assessment: Relevance is material, sensitivity "
                "follows do no harm, responsiveness supports resilience, and "
                "operationalization sets roles and indicators. Practical "
                "implication: confirm site criteria and bridge them to ranked "
                "priorities. "
                + "The residual gap remains material. " * 12
            ).strip(),
        ]
    }


def test_summary_overview_normalization_admits_two_or_three_and_rejects_other_shapes():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()

    assert helper(valid) == valid["paragraphs"]
    assert helper({"paragraphs": ["Relevance only."]}) == []
    assert helper({"paragraphs": ["Relevance.", "Sensitivity.", "Responsiveness.", "Operationalization."]}) == []
    assert helper({"paragraphs": ["Relevance.", "   "]}) == []


def test_summary_accepts_a_plain_language_verdict_without_magic_heading_words():
    from sector_lenses.climate_verified_schemas import validate_summary_overview
    summary = _summary_payload()
    summary["paragraphs"][0] = (
        "The project already acts on climate-FCV design. "
        + "The Project Operations Manual will define site selection. " * 12).strip()
    canonical = "The Project Operations Manual will define site selection. The residual gap remains material."
    assert validate_summary_overview(summary, canonical_text=canonical) == summary["paragraphs"]


def test_pipeline_admits_summary_overview_without_a_second_model_call():
    assessment = FakeClient(_responses(), [])
    assessment.responses[2]["summary_overview"] = _summary_payload()
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )

    assert result["summary_overview"] == _summary_payload()
    assert result["summary_overview_status"] == "generated"
    assert [call["stage"] for call in assessment.calls] == [
        "fact_extraction", "bounded_analysis", "judgment_review",
        "recommendation_compiler",
    ]


def test_invalid_summary_overview_is_empty_and_has_no_repair_call():
    assessment = FakeClient(_responses(), [])
    assessment.responses[2]["summary_overview"] = {"paragraphs": ["Only one paragraph."]}
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )

    assert result["summary_overview"] == {"paragraphs": []}
    assert "SUMMARY_OVERVIEW_INVALID" in result["validation"]["reason_codes"]
    import json
    trace = result["diagnostic_trace"]
    assert json.loads(trace["summary_overview"]["json"]) == {"paragraphs": ["Only one paragraph."]}
    assert trace["summary_validation"]["reason"] == "SUMMARY_PARAGRAPH_COUNT"
    assert [call["stage"] for call in assessment.calls] == [
        "fact_extraction", "bounded_analysis", "judgment_review",
        "recommendation_compiler",
    ]


def test_rejected_candidates_and_review_context_survive_in_diagnostics_only():
    import json
    from sector_lenses.climate_verified_render import build_reader_model
    review = {"verdict": "block", "reason_codes": ["AUTHORITY_UNSUPPORTED"],
              "object_ids": ["REC-001"], "issues": [{
                  "object_id": "REC-001", "reason_codes": ["AUTHORITY_UNSUPPORTED"],
                  "explanation": "The recommendation asserts unverified authority."}], "passed_core_ids": [],
              "core_reviews": [_core_review("REC-001", "unsupported",
                                             codes=["AUTHORITY_UNSUPPORTED"])]}
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(_responses(), []), FakeClient([review], [])))
    trace = result["diagnostic_trace"]
    assert json.loads(trace["semantic_review"]["json"]) == review
    reviewed = json.loads(trace["review_input"]["json"])
    assert reviewed["recommendations"][0]["recommendation_id"] == "REC-001"
    assert reviewed["operational_guidance"]
    assert json.loads(trace["reviewed_recommendations"]["json"]) == reviewed["recommendations"]
    raw = json.loads(trace["recommendation_candidates"]["json"])
    assert raw[0]["recommendation_id"] == "REC-001"
    assert result["priorities"] == []
    assert "diagnostic_trace" not in build_reader_model(result)


def test_diagnostic_snapshot_is_bounded_and_marks_truncation():
    from sector_lenses.climate_verified_pipeline import _diagnostic_snapshot
    small = _diagnostic_snapshot({"text": "normal"})
    assert not small["truncated"]
    large = _diagnostic_snapshot({"text": "x" * 100000})
    assert large["truncated"] and len(large["json"]) == 60000

def test_summary_validator_rejects_raw_shape_markup_and_length_failures():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()

    assert helper({"paragraphs": tuple(valid["paragraphs"])}) == []
    assert helper({"paragraphs": [valid["paragraphs"][0], ""]}) == []
    assert helper({"paragraphs": [valid["paragraphs"][0], 42]}) == []
    assert helper({"paragraphs": ["# heading", valid["paragraphs"][1]]}) == []
    assert helper({"paragraphs": [valid["paragraphs"][0] + "\n- bullet item", valid["paragraphs"][1]]}) == []
    assert helper({"paragraphs": ["<b>markup</b>", valid["paragraphs"][1]]}) == []
    assert helper({"paragraphs": ["Short.", "Still short."]}) == []


def test_summary_validator_rejects_exact_prefix_and_slice_copies_both_directions():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    executive = " ".join(valid["paragraphs"]) + " Additional executive context."

    assert helper(valid, executive_readout=executive) == []
    assert helper(
        {"paragraphs": [executive, valid["paragraphs"][1]]},
        executive_readout=valid["paragraphs"][0],
    ) == []


def test_summary_validator_rejects_unsupported_numbers_and_named_entities():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    canonical = "The Project Operations Manual will define site selection. The residual gap remains material."

    with_number = {"paragraphs": [
        valid["paragraphs"][0],
        valid["paragraphs"][1] + " The 2026 allocation is confirmed.",
    ]}
    with_entity = {"paragraphs": [
        valid["paragraphs"][0],
        valid["paragraphs"][1] + " Lake Victoria Authority is confirmed.",
    ]}
    assert helper(with_number, canonical_text=canonical) == []
    assert helper(with_entity, canonical_text=canonical) == []


def test_invalid_or_absent_summary_sets_stable_status_diagnostics():
    absent = FakeClient(_responses(), [])
    absent_result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(absent, _pass_review_client()),
    )
    assert absent_result["summary_overview_status"] == "fallback"
    assert "SUMMARY_OVERVIEW_FALLBACK_REQUIRED" in absent_result["validation"]["reason_codes"]

    invalid = FakeClient(_responses(), [])
    invalid.responses[2]["summary_overview"] = {"paragraphs": ["Only one paragraph."]}
    invalid_result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(invalid, _pass_review_client()),
    )
    assert invalid_result["summary_overview_status"] == "fallback"
    assert "SUMMARY_OVERVIEW_INVALID" in invalid_result["validation"]["reason_codes"]


def test_summary_admission_uses_judgments_after_validation_normalization():
    assessment = FakeClient(_responses(), [])
    assessment.responses[2]["relevance"]["value"] = "not-a-rating"
    assessment.responses[2]["relevance"]["rationale"] = (
        "Normalization marker is supported."
    )
    summary = _summary_payload()
    summary["paragraphs"][0] = summary["paragraphs"][0].replace(
        "The Project Operations Manual will define site selection.",
        "Normalization marker is supported.",
        1,
    )
    assessment.responses[2]["summary_overview"] = summary
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )

    assert result["judgments"]["relevance"]["value"] == "unclear"
    assert result["summary_overview_status"] == "fallback"
    assert "SUMMARY_OVERVIEW_INVALID" in result["validation"]["reason_codes"]
def _grounding_canonical(include_core=False, include_context=False, extras=""):
    text = (
        "The Project Operations Manual will define site selection. "
        "The residual gap remains material. "
        "The project recognizes the climate-FCV interaction and the overall "
        "takeaway is credible. Relevance sensitivity responsiveness "
        "operationalization practical implication confirm priorities. "
        "context relevant route."
    )
    if include_core:
        text += " Distinctive geographic footprint finding."
    if include_context:
        text += " Flooding can disrupt access."
    return text + extras


def test_summary_validator_rejects_unsupported_proper_names_and_acronyms():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    name_summary = {
        "paragraphs": [
            valid["paragraphs"][0] + " Somalia context is relevant.",
            valid["paragraphs"][1],
        ]
    }
    acronym_summary = {
        "paragraphs": [
            valid["paragraphs"][0],
            valid["paragraphs"][1] + " The UNDP route is relevant.",
        ]
    }
    canonical = _grounding_canonical()
    assert helper(name_summary, canonical_text=canonical) == []
    assert helper(acronym_summary, canonical_text=canonical) == []
    assert helper(
        name_summary,
        canonical_text=_grounding_canonical(extras=" Somalia context."),
    ) == name_summary["paragraphs"]
    assert helper(
        acronym_summary,
        canonical_text=_grounding_canonical(extras=" UNDP route."),
    ) == acronym_summary["paragraphs"]


def test_summary_validator_rejects_unsupported_action_but_admits_supported_action():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    action_summary = {
        "paragraphs": [
            valid["paragraphs"][0],
            valid["paragraphs"][1].replace(
                "The residual gap remains material.",
                "The project should establish a new committee.",
                1,
            ),
        ]
    }
    assert helper(
        action_summary,
        canonical_text=_grounding_canonical(extras=" project committee."),
    ) == []
    assert helper(
        action_summary,
        canonical_text=_grounding_canonical(
            extras=" project should establish a new committee."
        ),
    ) == action_summary["paragraphs"]


def test_summary_validator_accepts_paraphrased_narrative_jobs():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    paraphrased = {
        "paragraphs": [
            valid["paragraphs"][0].replace(
                "Verdict and foundation: The project recognizes the climate-FCV "
                "interaction and the overall takeaway is credible.",
                "The project demonstrates a clear finding about the climate-FCV interaction.",
            ),
            valid["paragraphs"][1].replace(
                "Practical implication: confirm site criteria and bridge them to ranked "
                "priorities.",
                "Remaining attention should focus on follow-up and decision points.",
            ),
        ]
    }
    assert helper(paraphrased) == paraphrased["paragraphs"]


def test_pipeline_grounds_summary_in_core_questions_and_context_evidence():
    assessment = FakeClient(_responses(), [])
    summary = _summary_payload()
    summary["paragraphs"][0] = summary["paragraphs"][0].replace(
        "The Project Operations Manual will define site selection.",
        "Distinctive geographic footprint finding.",
        1,
    )
    assessment.responses[2]["summary_overview"] = summary
    assessment.responses[2]["core_questions"] = [{
        "question_id": "dq-geo-overlap",
        "theme": "driver_geo_overlap",
        "question": "Where does the project land?",
        "source": "FCV-Sensitive Climate Action Framework",
        "summary": "Distinctive geographic footprint finding.",
        "evidence_ids": ["PF-001"],
        "watch": "Check the footprint.",
    }]
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )
    assert result["summary_overview_status"] == "generated"

    context_assessment = FakeClient(_responses(), [])
    context_summary = _summary_payload()
    context_summary["paragraphs"][0] = context_summary["paragraphs"][0].replace(
        "The Project Operations Manual will define site selection.",
        "Flooding can disrupt access.",
        1,
    )
    context_assessment.responses[2]["summary_overview"] = context_summary
    context_args = _arguments()
    context_args["context_evidence"] = []
    context_result = run_verified_climate_pipeline(
        **context_args,
        clients=PipelineClients(context_assessment, _pass_review_client()),
    )
    assert context_result["summary_overview_status"] == "fallback"


def test_legacy_fallback_requires_known_non_judgment_evidence_ids():
    from sector_lenses.climate_verified_render import build_reader_model

    assessment = {
        "overview_summary": "Overview foundation. Overview implication.",
        "facts": [{
            "claim_id": "PF-001",
            "statement": "Known rationale. Known response. Known operation.",
        }],
        "judgments": {
            "relevance": {"value": "high", "rationale": "Fake rationale.", "evidence_ids": ["FAKE"]},
            "sensitivity": {"value": "strong", "rationale": "Known rationale.", "evidence_ids": ["PF-001"]},
            "responsiveness": {"value": "moderate", "rationale": "Known response.", "evidence_ids": ["PF-001"]},
            "operationalization": {"value": "early", "rationale": "Known operation.", "evidence_ids": ["PF-001"]},
        },
        "priorities": [],
    }
    paragraphs = build_reader_model(assessment)["summary_overview"]
    assert "Fake rationale." not in " ".join(paragraphs)
    assert "Known rationale." in " ".join(paragraphs)


def test_summary_validator_rejects_unanchored_route_acronyms_but_accepts_canonical_ones():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    canonical = _grounding_canonical()
    for acronym in ("MPA", "IPF", "WBG"):
        summary = {
            "paragraphs": [
                valid["paragraphs"][0].replace(
                    "The Project Operations Manual will define site selection.",
                    f"The {acronym} route is relevant.",
                    1,
                ),
                valid["paragraphs"][1],
            ]
        }
        assert helper(summary, canonical_text=canonical) == []
        assert helper(
            summary,
            canonical_text=canonical + f" {acronym} route.",
        ) == summary["paragraphs"]


def test_summary_validator_rejects_inflected_unsupported_actions_but_admits_grounded_forms():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    cases = (
        ("establishes", "committee"),
        ("approves", "allocation"),
        ("ensures", "inclusion"),
    )
    for verb, noun in cases:
        summary = {
            "paragraphs": [
                valid["paragraphs"][0],
                valid["paragraphs"][1].replace(
                    "The residual gap remains material.",
                    f"The project {verb} {noun}.",
                    1,
                ),
            ]
        }
        assert helper(
            summary,
            canonical_text=_grounding_canonical(extras=f" project {noun}."),
        ) == []
        assert helper(
            summary,
            canonical_text=_grounding_canonical(
                extras=f" project {verb} {noun}."
            ),
        ) == summary["paragraphs"]


def test_pipeline_persists_context_evidence_for_reader_summary_grounding():
    from sector_lenses.climate_verified_render import build_reader_model

    assessment = FakeClient(_responses(), [])
    summary = _summary_payload()
    summary["paragraphs"][0] = summary["paragraphs"][0].replace(
        "The Project Operations Manual will define site selection.",
        "Flooding can disrupt access.",
        1,
    )
    assessment.responses[2]["summary_overview"] = summary
    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, _pass_review_client()),
    )

    assert result["summary_overview_status"] == "generated"
    assert result["context_evidence"][0]["evidence_id"] == "CE-001"
    assert result["context_evidence"][0]["statement"] == (
        "Flooding can disrupt access."
    )
    reader = build_reader_model(result)
    assert reader["summary_overview_status"] == "generated"
    assert "Flooding can disrupt access." in " ".join(
        reader["summary_overview"]
    )


def test_summary_validator_allows_connectives_and_canonical_morphology():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    summary = _summary_payload()
    summary["paragraphs"][0] = summary["paragraphs"][0].replace(
        "Verdict and foundation: The project recognizes the climate-FCV "
        "interaction and the overall takeaway is credible.",
        "Verdict and foundation: However, the project remains resilient in the "
        "climate-FCV interaction and the overall takeaway is credible.",
        1,
    )
    canonical = _grounding_canonical(
        extras=" resilience resilience-building."
    )
    assert helper(summary, canonical_text=canonical) == summary["paragraphs"]


def test_summary_validator_rejects_projectile_but_admits_curated_implementation_variant():
    from sector_lenses import climate_verified_pipeline as climate_pipeline

    helper = getattr(climate_pipeline, "_summary_overview_paragraphs")
    valid = _summary_payload()
    projectile = {
        "paragraphs": [
            valid["paragraphs"][0].replace(
                "The Project Operations Manual will define site selection.",
                "The projectile route is relevant.",
                1,
            ),
            valid["paragraphs"][1],
        ]
    }
    assert helper(
        projectile,
        canonical_text=_grounding_canonical(extras=" project route relevant."),
    ) == []
    implemented = {
        "paragraphs": [
            valid["paragraphs"][0],
            valid["paragraphs"][1].replace(
                "The residual gap remains material.",
                "The project implemented requirements.",
                1,
            ),
        ]
    }
    assert helper(
        implemented,
        canonical_text=_grounding_canonical(
            extras=" project implementation requirements."
        ),
    ) == implemented["paragraphs"]


def test_explicit_revise_gets_one_validated_correction_and_independent_review():
    responses = _responses()
    responses[0]["document_integrity_findings"] = [{
        "flag_id": "DIF-1", "category": "document_inconsistency",
        "flag": "Six suspended and four cancelled sites disagree.",
        "why_it_matters": "Verify the document's stated scope.",
        "document_basis_ids": ["DOC-1-B-1"], "residual_gap_ids": [],
        "suggested_verification": "Check whether the populations and statuses match.",
    }]
    corrected = deepcopy(responses[-1])
    corrected["recommendation_candidates"][0]["minimum_action"] = "Document flood exposure and access considerations."
    assessment = FakeClient(responses + [corrected], [])
    reviewer = FakeClient([
        {"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
         "object_ids": ["REC-001::current_document"], "issues": [{"object_id": "REC-001::current_document",
         "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
         "explanation": "Keep the action within the evidenced site-selection scope."}],
         "passed_core_ids": ["REC-001"],
         "core_reviews": [_core_review("REC-001", "unsupported")]},
        _native_pass_review,
    ], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert len(assessment.calls) == 5
    assert len(reviewer.calls) == 2
    assert result["priorities"][0]["minimum_action"] == "Document flood exposure and access considerations."
    assert result["diagnostic_trace"]["recommendation_correction"]["recovered_ids"] == ["REC-001"]
    assert reviewer.calls[0]["payload"]["readiness_flags"]
    assert reviewer.calls[1]["payload"]["readiness_flags"] == []
    assert assessment.calls[-1]["payload"]["readiness_flags"] == []
    assert result["review_readiness_flags"] == []


def test_malformed_core_review_does_not_trigger_speculative_correction():
    reviewer = FakeClient([{
        "verdict": "revise",
        "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
        "object_ids": ["REC-001"],
        "issues": [{
            "object_id": "REC-001",
            "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
            "explanation": "The premise needs evidence.",
        }],
        "passed_core_ids": [],
        "core_reviews": [{**_core_review(), "source_block_ids": ["DOC-unknown"]}],
    }], [])
    assessment = FakeClient(_responses(), [])

    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(assessment, reviewer),
    )

    assert result["priorities"] == []
    assert not any(call["payload"].get("call_scope") == "core_correction"
                   for call in assessment.calls)


def test_mixed_core_reviews_repair_only_precise_finding_and_preserve_valid_core():
    responses = _responses()
    base = responses[-1]["recommendation_candidates"][0]
    for index in (2, 3):
        candidate = deepcopy(base)
        candidate["recommendation_id"] = f"REC-{index:03d}"
        candidate["title"] = f"Bounded action {index}"
        responses[-1]["recommendation_candidates"].append(candidate)
    corrected = deepcopy(base)
    corrected["recommendation_id"] = "REC-002"
    corrected["minimum_action"] = "Document the bounded second action."
    review = {
        "verdict": "revise",
        "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-002::current_document"],
        "issues": [{
            "object_id": "REC-002::current_document",
            "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
            "explanation": "The drafting overstates the evidenced scope.",
        }],
        "passed_core_ids": ["REC-001", "REC-002", "REC-003"],
        "core_reviews": [
            {**_core_review("REC-001"), "source_block_ids": ["DOC-unknown"]},
            _core_review("REC-002", "unsupported"),
            _core_review("REC-003"),
        ],
    }
    assessment = FakeClient(responses + [{"recommendation_candidates": [corrected]}], [])

    result = run_verified_climate_pipeline(
        **_arguments(),
        clients=PipelineClients(assessment, FakeClient([review, _native_pass_review], [])),
    )

    assert [item["recommendation_id"] for item in result["priorities"]] == [
        "REC-002", "REC-003",
    ]
    correction = next(call for call in assessment.calls
                      if call["payload"].get("call_scope") == "core_correction")
    assert [item["recommendation_id"] for item in correction["payload"]["recommendations"]] == ["REC-002"]


def test_native_compiler_correction_needs_no_optional_drafting():
    from sector_lenses.climate_verified_schemas import CANDIDATE_SCHEMA
    responses = _responses()
    corrected = deepcopy(responses[-1])
    corrected["recommendation_candidates"] = [{
        key: value for key, value in corrected["recommendation_candidates"][0].items()
        if key in CANDIDATE_SCHEMA["properties"]
    }]
    assert set(corrected["recommendation_candidates"][0]) == set(CANDIDATE_SCHEMA["required"])
    reviewer = FakeClient([
        {"verdict": "revise", "reason_codes": ["RECOMMENDATION_DISPROPORTIONATE"],
         "object_ids": ["REC-001"], "issues": [{"object_id": "REC-001",
         "explanation": "Keep the decision within site-selection scope."}], "passed_core_ids": [],
         "core_reviews": [_core_review("REC-001", "unsupported",
                                       codes=["RECOMMENDATION_DISPROPORTIONATE"])]},
        _native_pass_review,
    ], [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses + [corrected], []), reviewer),
    )
    assert [item["recommendation_id"] for item in result["priorities"]] == ["REC-001"]
    assert result["priorities"][0]["current_document_drafting"] is None
    assert len(reviewer.calls) == 2


def test_bad_optional_drafting_does_not_suppress_reviewed_core():
    responses = _responses()
    responses[-1]["recommendation_candidates"][0]["current_document_drafting"]["text"] = "Too short."
    reviewer = _pass_review_client()
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses, []), reviewer),
    )
    assert len(result["priorities"]) == 1
    assert result["priorities"][0]["current_document_drafting"] is None
    assert reviewer.calls[0]["payload"]["recommendations"][0]["current_document_drafting"] is None


def test_explicit_draft_block_does_not_restore_or_remove_a_core_action():
    reviewer = FakeClient([{
        "verdict": "block", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-001::current_document"],
        "passed_core_ids": ["REC-001"],
        "core_reviews": [_core_review()],
        "issues": [{"object_id": "REC-001::current_document", "explanation": "Unsupported wording."}],
    }], [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(_responses(), []), reviewer),
    )
    assert len(result["priorities"]) == 1
    assert result["priorities"][0]["current_document_drafting"] is None
    assert reviewer.calls[0]["payload"]["drafting_artifacts"][0]["object_id"] == "REC-001::current_document"


def test_draft_failure_without_independent_core_pass_is_conservative():
    reviewer = FakeClient([{"verdict": "block", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
                           "object_ids": ["REC-001::current_document"], "passed_core_ids": [],
                           "core_reviews": [_core_review("REC-001", "unsupported")] }], [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(_responses(), []), reviewer),
    )
    assert result["priorities"] == []


def test_draft_repair_uses_drafting_schema_and_keeps_reviewed_decision():
    responses = _responses()
    original = responses[-1]["recommendation_candidates"][0]
    block = {**original["current_document_drafting"], "drafting_role": "current_document"}
    repaired = {"drafting_sets": [{"recommendation_id": "REC-001", "drafting_blocks": [block]}]}
    reviewer = FakeClient([
        {"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
         "object_ids": ["REC-001::current_document"], "passed_core_ids": ["REC-001"],
         "core_reviews": [_core_review()],
         "issues": [{"object_id": "REC-001::current_document", "explanation": "Keep proposed wording bounded."}]},
        _native_pass_review,
    ], [])
    assessment = FakeClient(responses + [repaired], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert assessment.calls[-1]["stage"] == "drafting_compiler"
    assert len(reviewer.calls) == 2
    assert result["priorities"][0]["decision"] == original["decision"]
    assert result["priorities"][0]["current_document_drafting"] is not None


@pytest.mark.parametrize("second_review", [
    {"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
     "object_ids": [], "issues": [{"object_id": "REC-001::current_document",
                                     "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
                                     "explanation": "The repaired draft remains too broad."}],
     "passed_core_ids": ["REC-001"], "core_reviews": [_core_review()]},
    {"verdict": "pass", "reason_codes": [], "object_ids": [], "issues": [],
     "passed_core_ids": []},
    {"verdict": "pass", "reason_codes": [], "object_ids": [], "issues": None,
     "passed_core_ids": ["REC-001"], "core_reviews": [_core_review()]},
    {"verdict": "pass", "reason_codes": [], "object_ids": [], "issues": [None],
     "passed_core_ids": ["REC-001"], "core_reviews": [_core_review()]},
])
def test_draft_rereview_withholds_bad_draft_without_removing_sound_parent(second_review):
    responses = _responses()
    original = responses[-1]["recommendation_candidates"][0]
    block = {**original["current_document_drafting"], "drafting_role": "current_document"}
    repaired = {"drafting_sets": [{"recommendation_id": "REC-001", "drafting_blocks": [block]}]}
    initial = {
        "verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-001::current_document"],
        "issues": [{"object_id": "REC-001::current_document",
                    "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
                    "explanation": "Keep the wording bounded."}],
        "passed_core_ids": ["REC-001"], "core_reviews": [_core_review()],
    }

    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(
            FakeClient(responses + [repaired], []),
            FakeClient([initial, second_review], []),
        ),
    )

    assert [item["recommendation_id"] for item in result["priorities"]] == ["REC-001"]
    assert result["priorities"][0]["current_document_drafting"] is None


def test_optional_draft_timeout_preserves_core_and_malformed_draft_is_isolated():
    responses = _responses()
    responses[-1]["recommendation_candidates"][0].pop("current_document_drafting")
    responses[-1]["recommendation_candidates"][0].pop("operational_instrument_drafting")
    class MissingDraft(FakeClient):
        def complete_json(self, **kwargs):
            if kwargs["stage"] == "drafting_compiler":
                raise TimeoutError("Draft timed out")
            return super().complete_json(**kwargs)
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(MissingDraft(responses, []), _pass_review_client()),
    )
    assert len(result["priorities"]) == 1
    assert "DRAFTING_GENERATION_UNAVAILABLE" in result["validation"]["reason_codes"]
    responses = _responses()
    responses[-1]["recommendation_candidates"][0]["drafting_blocks"] = [None]
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses, []), _pass_review_client()),
    )
    assert len(result["priorities"]) == 1
    assert result["priorities"][0]["current_document_drafting"] is None


def test_native_schema_three_decisions_survive_missing_optional_drafts(monkeypatch):
    """Transport benchmark only: it does not establish live decision quality."""
    from sector_lenses.climate_verified_schemas import CANDIDATE_SCHEMA
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    base = responses[-1]["recommendation_candidates"][0]
    candidates = []
    for index, decision in enumerate((
        "Agree climate exposure and fair access considerations for site selection.",
        "Agree a proportionate continuity decision for climate-related access disruption.",
        "Agree how exposed groups can use the climate-related support arrangements.",
    ), 1):
        record = {key: deepcopy(value) for key, value in base.items()
                  if key in CANDIDATE_SCHEMA["properties"]}
        record.update(recommendation_id=f"REC-{index:03d}", title=decision,
                      decision=decision, minimum_action=decision)
        assert set(record) == set(CANDIDATE_SCHEMA["required"])
        candidates.append(record)
    responses[-1]["recommendation_candidates"] = candidates
    assessment = FakeClient(responses + [{"drafting_sets": []}], [])
    reviewer = FakeClient([{"verdict": "pass", "passed_core_ids": [item["recommendation_id"] for item in candidates],
                            "core_reviews": [_core_review(item["recommendation_id"]) for item in candidates],
                            "reason_codes": [], "object_ids": [], "issues": []}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert len(result["priorities"]) == 3
    assert all(item["current_document_drafting"] is None for item in result["priorities"])
    assert any(call["stage"] == "drafting_compiler" for call in assessment.calls)


    # Exercise the same runtime -> normalizer -> reader gate -> export boundary
    # that stopped quality7, after independently reviewed synthetic decisions.
    from io import BytesIO
    from docx import Document
    from sector_lenses.climate_verified_runtime import run_verified_from_doc_parts
    from sector_lenses.climate_verified_render import render_reader_html, write_reader_docx
    monkeypatch.setattr("sector_lenses.climate_verified_runtime.run_verified_climate_pipeline",
                        lambda **kwargs: result)
    bundle = run_verified_from_doc_parts(doc_parts=[{
        "label": "PROJECT DOCUMENT", "name": "synthetic-pcn.txt",
        "raw_text": "The Project Operations Manual will define site selection.",
    }], climate_grounding={}, clients=object(), run_id="synthetic-runtime",
        doc_type="PCN", instrument_type="IPF")
    assert len(bundle["reader"]["priorities"]) == 3
    html = render_reader_html(bundle["reader"])
    stream = BytesIO(); write_reader_docx(bundle["reader"], stream); stream.seek(0)
    word = " ".join(p.text for p in Document(stream).paragraphs)
    for priority in bundle["reader"]["priorities"]:
        assert priority["title"] in html and priority["title"] in word

def test_missing_native_core_pass_list_never_accepts_a_recommendation():
    reviewer = FakeClient([{"verdict": "pass", "reason_codes": [], "object_ids": [], "issues": []}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(FakeClient(_responses(), []), reviewer))
    assert result["priorities"] == []
    assert "SEMANTIC_REVIEW_TARGET_UNRESOLVED" in result["validation"]["reason_codes"]


def test_optional_drafting_reserves_time_for_core_review():
    responses = _responses()
    candidate = responses[-1]["recommendation_candidates"][0]
    candidate.pop("current_document_drafting")
    candidate.pop("operational_instrument_drafting")
    assessment = FakeClient(responses, [])
    result = run_verified_climate_pipeline(
        **_arguments(), wall_clock_seconds=120,
        clients=PipelineClients(assessment, _pass_review_client()),
    )
    assert len(result["priorities"]) == 1
    assert "drafting_compiler" not in [call["stage"] for call in assessment.calls]
    assert "DRAFTING_GENERATION_UNAVAILABLE" in result["validation"]["reason_codes"]


def test_correction_without_usage_consumes_shared_allowance_and_cannot_skip_review():
    responses = _responses()
    class OpaqueCorrection(FakeClient):
        @property
        def last_output_tokens(self):
            return None
    reviewer = FakeClient([{"verdict": "revise", "passed_core_ids": [],
        "reason_codes": ["RECOMMENDATION_DISPROPORTIONATE"], "object_ids": ["REC-001"],
        "issues": [{"object_id": "REC-001", "explanation": "Bound the action."}],
        "core_reviews": [_core_review("REC-001", "unsupported",
                                      codes=["RECOMMENDATION_DISPROPORTIONATE"])]}], [])
    client = OpaqueCorrection(responses + [deepcopy(responses[-1])], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(client, reviewer))
    assert result["priorities"] == []
    assert result["manifest"]["sampling"]["correction_output_token_remaining"] == 0
    assert len(reviewer.calls) == 1
    assert client.calls[-1]["max_transient_retries"] == 0


def test_new_reader_export_rechecks_support_and_never_revives_rating_recap():
    from sector_lenses.climate_verified_render import build_reader_model
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    reviewer = FakeClient([_native_pass_review] * 2, [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses + [_final_reader()], []), reviewer),
    )
    assert build_reader_model(result)["summary_overview_status"] == "generated"
    result["final_reader"]["summary_paragraphs"][0]["evidence_ids"].append("REC-withheld")
    model = build_reader_model(result)
    assert model["summary_overview_status"] == "fallback"
    text = " ".join(model["summary_overview"])
    assert "could not be independently verified" in text
    assert "Sensitivity is" not in text and "Operationalization" not in text
    result.pop("final_reader")
    assert build_reader_model(result)["summary_overview_status"] == "fallback"


def test_correction_still_rejected_is_withheld_without_loop():
    responses = _responses()
    review = {"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
              "object_ids": ["REC-001"], "issues": [{"object_id": "REC-001",
              "explanation": "Remove unsupported scope."}], "passed_core_ids": [],
              "core_reviews": [_core_review("REC-001", "unsupported")]}
    assessment = FakeClient(responses + [deepcopy(responses[-1])], [])
    reviewer = FakeClient([review, review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert len(assessment.calls) == 5
    assert len(reviewer.calls) == 2
    assert result["priorities"] == []


@pytest.mark.parametrize("second_review", [
    {"verdict": "pass", "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
     "object_ids": ["REC-001"], "issues": [], "passed_core_ids": ["REC-001"],
     "core_reviews": [_core_review()]},
    {"verdict": "revise", "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
     "object_ids": ["REC-001"], "issues": [], "passed_core_ids": ["REC-001"],
     "core_reviews": [_core_review()]},
    {"verdict": "pass", "reason_codes": [], "object_ids": ["REC-unknown"],
     "issues": [], "passed_core_ids": ["REC-001"], "core_reviews": [_core_review()]},
    {"verdict": "revise", "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
     "object_ids": [], "issues": [], "passed_core_ids": ["REC-001"],
     "core_reviews": [_core_review()]},
])
def test_core_correction_rereview_rejects_incoherent_envelopes(second_review):
    responses = _responses()
    correction = deepcopy(responses[-1])
    first_review = {
        "verdict": "revise", "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
        "object_ids": ["REC-001"],
        "issues": [{"object_id": "REC-001", "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
                    "explanation": "Bound the premise."}],
        "passed_core_ids": [], "core_reviews": [_core_review("REC-001", "unsupported")],
    }

    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(
            FakeClient(responses + [correction], []),
            FakeClient([first_review, second_review], []),
        ),
    )

    assert result["priorities"] == []


def test_correction_timeout_preserves_unaffected_priority():
    import pytest
    responses = _responses()
    other = deepcopy(responses[-1]["recommendation_candidates"][0])
    other["recommendation_id"] = "REC-002"
    other["title"] = "Maintain delivery access"
    responses[-1]["recommendation_candidates"].append(other)
    class TimeoutCorrection(FakeClient):
        def complete_json(self, **kwargs):
            if kwargs["payload"].get("bounded_correction"):
                raise TimeoutError("Remaining budget exhausted")
            return super().complete_json(**kwargs)
    reviewer = FakeClient([{"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-001"], "issues": [{"object_id": "REC-001", "explanation": "Remove unsupported scope."}],
        "passed_core_ids": ["REC-002"], "core_reviews": [
            _core_review("REC-001", "unsupported"), _core_review("REC-002")]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(TimeoutCorrection(responses, []), reviewer))
    assert [p["recommendation_id"] for p in result["priorities"]] == ["REC-002"]
    assert result["diagnostic_trace"]["recommendation_correction"]["failure"] == "TimeoutError"


def test_correction_cannot_expand_evidence_or_replace_passed_actions():
    responses = _responses()
    correction = deepcopy(responses[-1])
    correction["recommendation_candidates"][0]["project_anchor_ids"].append("PF-invented")
    assessment = FakeClient(responses + [correction], [])
    reviewer = FakeClient([{"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-001"], "issues": [{"object_id": "REC-001", "explanation": "Remove unsupported scope."}],
        "passed_core_ids": [], "core_reviews": [_core_review("REC-001", "unsupported")]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result["priorities"] == []
    assert len(reviewer.calls) == 1
    trace = result["diagnostic_trace"]["recommendation_correction"]
    assert trace["response"]["json"]
    assert trace["rejections"] == [{"recommendation_id": "REC-001",
                                      "reason_codes": ["CORRECTION_EVIDENCE_EXPANDED"]}]


def test_cancellation_during_correction_propagates():
    import pytest
    class CancelCorrection(FakeClient):
        def complete_json(self, **kwargs):
            if kwargs["payload"].get("bounded_correction"):
                raise RuntimeError("Verified Climate-FCV pipeline cancelled")
            return super().complete_json(**kwargs)
    reviewer = FakeClient([{"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["REC-001"], "issues": [{"object_id": "REC-001", "explanation": "Remove unsupported scope."}],
        "passed_core_ids": [], "core_reviews": [_core_review("REC-001", "unsupported")]}], [])
    with pytest.raises(RuntimeError, match="cancelled"):
        run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(CancelCorrection(_responses(), []), reviewer))


def test_unresolved_review_targets_do_not_trigger_correction_even_with_issue_text():
    reviewer = FakeClient([{"verdict": "revise", "reason_codes": ["DRAFTING_SCOPE_UNSUPPORTED"],
        "object_ids": ["unknown"], "issues": [{"object_id": "REC-001", "explanation": "Remove unsupported scope."}],
        "passed_core_ids": [], "core_reviews": [_core_review("REC-001", "unsupported")]}], [])
    assessment = FakeClient(_responses(), [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result["priorities"] == []
    assert len(assessment.calls) == 4


def test_judgment_transports_narrative_findings_without_an_extra_call():
    responses = _responses()
    narrative = [{"theme": "design", "paragraphs": ["The project uses site selection."], "evidence_ids": ["PF-001"]}]
    responses[2]["assessment_narratives"] = narrative
    client = FakeClient(responses, [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(client, _pass_review_client()),
    )
    assert result["assessment_narratives"] == narrative
    assert len(client.calls) == 4


def _final_reader():
    return {
        "summary_paragraphs": [
            {"object_id": f"SUM-{index + 1:03d}", "text": paragraph,
             "evidence_ids": ["PF-001", "PW-001", "RG-001", "CE-001", "REC-001"]}
            for index, paragraph in enumerate(_summary_payload()["paragraphs"])
        ],
        "assessment_narratives": [
            {"object_id": f"NARR-{theme}", "theme": theme,
             "paragraphs": [_summary_payload()["paragraphs"][0]],
             "evidence_ids": ["PF-001", "PW-001", "RG-001", "CE-001", "REC-001"]}
            for theme in ("design", "dividends", "inclusion")
        ],
    }


def test_reader_revision_quote_must_identify_one_span_in_current_record():
    final_reader = _final_reader()
    current_quote = "Verdict and foundation: The project recognizes"
    review = {
        "verdict": "revise",
        "object_ids": ["NARR-dividends"],
        "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
        "issues": [{
            "object_id": "NARR-dividends",
            "reason_codes": ["PROJECT_FACT_UNSUPPORTED"],
            "quoted_text": current_quote,
            "explanation": "Keep the factual premise bounded.",
        }],
    }

    assert _valid_reader_revision(review, final_reader)
    review["issues"][0]["quoted_text"] = "Text removed by an earlier correction."
    assert not _valid_reader_revision(review, final_reader)
    review["issues"][0]["quoted_text"] = "the"
    assert not _valid_reader_revision(review, final_reader)
    review["issues"][0]["quoted_text"] = current_quote
    review["issues"][0]["object_id"] = "NARR-unknown"
    review["object_ids"] = ["NARR-unknown"]
    assert not _valid_reader_revision(review, final_reader)

    overlapping = _final_reader()
    overlapping["assessment_narratives"][1]["paragraphs"] = ["a a a"]
    review["issues"][0]["object_id"] = "NARR-dividends"
    review["object_ids"] = ["NARR-dividends"]
    review["issues"][0]["quoted_text"] = "a a"
    assert not _valid_reader_revision(review, overlapping)


def test_final_reader_compiles_from_accepted_decisions_and_receives_semantic_review():
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    final = _final_reader()
    assessment = FakeClient(responses + [final], [])
    reviewer = FakeClient([_native_pass_review] * 2, [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert assessment.calls[-1]["stage"] == "reader_compiler"
    assert assessment.calls[-1]["payload"]["accepted_recommendation_ids"] == ["REC-001"]
    assert reviewer.calls[-1]["stage"] == "reader_review"
    assert result["summary_overview_status"] == "generated"
    assert result["final_reader"]["summary_paragraphs"] == final["summary_paragraphs"]


def test_semantic_reader_rejection_does_not_reduce_valid_actions():
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    reviewer = FakeClient([
        _native_pass_review,
        {"verdict": "block", "object_ids": ["SUM-001"],
         "reason_codes": ["PROJECT_FACT_UNSUPPORTED"], "issues": [
             {"object_id": "SUM-001", "explanation": "The sentence reverses the supported meaning."}], "passed_core_ids": []},
    ], [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses + [_final_reader()], []), reviewer),
    )
    assert len(result["priorities"]) == 1
    assert result["summary_overview_status"] == "fallback"
    assert result["summary_overview"]["paragraphs"] == []


def test_final_reader_cannot_reference_withheld_action_even_if_words_overlap():
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    final = deepcopy(_final_reader())
    final["summary_paragraphs"][0]["evidence_ids"].append("REC-withheld")
    reviewer = FakeClient([_native_pass_review], [])
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(FakeClient(responses + [final], []), reviewer),
    )
    assert result["summary_overview_status"] == "fallback"
    assert len(result["priorities"]) == 1
    assert len(reviewer.calls) == 1


def test_optional_reader_provider_errors_preserve_accepted_actions():
    for failed_stage in ("reader_compiler", "reader_review"):
        responses = _responses()
        responses[2]["reader_contract"] = "final-reader-v1"
        class FailingProvider(FakeClient):
            def complete_json(self, **kwargs):
                if kwargs["stage"] == failed_stage:
                    raise ConnectionError("Provider unavailable")
                return super().complete_json(**kwargs)
        assessment = FailingProvider(responses + [_final_reader()], [])
        reviewer = FailingProvider([_native_pass_review] * 2, [])
        result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
        assert len(result["priorities"]) == 1
        assert result["summary_overview_status"] == "fallback"


def test_exports_use_reviewed_narratives_instead_of_mutable_top_level_prose():
    from sector_lenses.climate_verified_render import build_reader_model
    responses = _responses()
    responses[2]["reader_contract"] = "final-reader-v1"
    final = _final_reader()
    paragraphs = [("The Project Operations Manual will define site selection. "
                   "Flooding can disrupt access. Site criteria are not yet specified. "
                   "The residual gap remains material. ") * 4]
    final["assessment_narratives"] = [{"object_id": f"NARR-{theme}", "theme": theme,
        "paragraphs": paragraphs, "evidence_ids": ["PF-001", "PW-001", "RG-001", "CE-001"]}
        for theme in ("design", "dividends", "inclusion")]
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [final], []), FakeClient([_native_pass_review] * 2, [])))
    reviewed_copy = deepcopy(final)
    result["assessment_narratives"][0]["paragraphs"] = [
        "The Project Operations Manual will not define site selection. " * 12]
    # Break the alias deliberately, as a modified saved JSON document would.
    result["final_reader"] = reviewed_copy
    model = build_reader_model(result)
    assert model["assessment_themes"][0]["paragraphs"] == [paragraphs[0].strip()]


def test_native_final_reader_requires_all_three_theme_answers():
    from sector_lenses.climate_final_reader import validate_final_reader
    responses = _responses()
    result = run_verified_climate_pipeline(**_arguments(),
        clients=PipelineClients(FakeClient(responses, []), _pass_review_client()))
    final = _final_reader()
    final["assessment_narratives"] = []
    assert validate_final_reader(final, result) == ([], [])


def test_native_reader_preserves_richer_reviewed_benefit_findings(tmp_path):
    """Synthetic prose tests transport/admission, not project analytical quality."""
    from sector_lenses.climate_final_reader import reviewed_final_reader
    from sector_lenses.climate_verified_render import (
        build_reader_model, render_reader_html, write_reader_docx,
    )
    from docx import Document
    from sector_lenses.climate_verified_pipeline import PipelineClients, run_verified_climate_pipeline

    final = _final_reader()
    for record in final['assessment_narratives']:
        if record['theme'] in {'dividends', 'inclusion'}:
            base = record['paragraphs'][0]
            second = (
                'The project could support more reliable livelihoods where people can use '
                'the facilities fairly. This potential depends on access and delivery '
                'arrangements that work for the intended users during disruption. '
                'The existing measures provide a foundation, while the remaining uncertainty '
                'concerns their practical application. Clear access arrangements would '
                'help connect the planned investments to benefits people can use.'
            )
            record['paragraphs'] = [base, second]
            assert 130 < sum(len(p.split()) for p in record['paragraphs']) <= 180
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(
            FakeClient(responses + [final], []), FakeClient([_native_pass_review] * 2, [])))
    assert result['summary_overview_status'] == 'generated'
    reader = build_reader_model(result)
    for theme in reader['assessment_themes']:
        assert theme['narrative_status'] == 'generated'
    assert second in render_reader_html(reader)
    docx_path = write_reader_docx(reader, tmp_path / 'synthetic-benefits.docx')
    assert second in '\n'.join(p.text for p in Document(docx_path).paragraphs)
    assert reviewed_final_reader(final, {'verdict': 'block'}, result) == ([], [])


def test_native_reader_still_withholds_overlong_findings():
    from sector_lenses.climate_final_reader import validate_final_reader

    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    result = run_verified_climate_pipeline(
        **_arguments(), clients=PipelineClients(
            FakeClient(responses + [_final_reader()], []),
            FakeClient([_native_pass_review] * 2, [])))
    final = _final_reader()
    assert validate_final_reader(final, result)[0]
    final['assessment_narratives'][1]['paragraphs'] = ['The project supports traders. ' * 101]
    assert len(final['assessment_narratives'][1]['paragraphs'][0].split()) > 400
    assert validate_final_reader(final, result) == ([], [])


def test_final_reader_repairs_only_targeted_prose_and_rechecks_meaning():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    corrected = deepcopy(original)
    corrected['assessment_narratives'][1]['paragraphs'][0] = corrected['assessment_narratives'][1]['paragraphs'][0].replace('The project', 'This project', 1)
    patch = {'corrections': [{'object_id': 'NARR-dividends',
                            'edits': [{'paragraph_index': 0, 'old_text': 'The project', 'new_text': 'This project'}]}]}
    assessment = FakeClient(responses + [original, patch], [])
    reviewer = FakeClient([_native_pass_review,
        {'verdict':'revise', 'object_ids':['NARR-dividends'],
         'reason_codes':['PROJECT_FACT_UNSUPPORTED'], 'issues':[
             {'object_id':'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
              'quoted_text': READER_QUOTE, 'explanation':'Keep planned benefits conditional.'}]},
        _native_pass_review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    assert result['final_reader'] == corrected
    assert len(result['priorities']) == 1
    assert [call['stage'] for call in reviewer.calls].count('reader_review') == 2
    assert assessment.calls[-1]['payload']['bounded_correction'][0]['object_id'] == 'NARR-dividends'
    assert assessment.calls[-1]['max_transient_retries'] == 0


def test_invalid_reader_references_cannot_be_repaired_by_text_changes():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    invalid = deepcopy(_final_reader())
    invalid['summary_paragraphs'][0]['evidence_ids'].append('PF-unknown')
    assessment = FakeClient(responses + [invalid, _final_reader()], [])
    reviewer = FakeClient([_native_pass_review] * 2, [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'fallback'
    assert assessment.calls[-1]['stage'] == 'reader_compiler'
    assert result['diagnostic_trace']['final_reader']['unrepairable_issues'][0]['reason_codes'] == ['READER_EVIDENCE_INVALID']
    assert len(reviewer.calls) == 1


def test_reader_correction_cannot_rewrite_untargeted_records():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    correction = {'corrections': [{'object_id': 'NARR-design',
                                  'edits': []}]}
    assessment = FakeClient(responses + [original, correction], [])
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Qualify the benefit.'}]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'fallback'
    assert len(result['priorities']) == 1
    assert len(reviewer.calls) == 2
    assert result['diagnostic_trace']['final_reader']['correction']['issues'][0]['reason_codes'] == ['READER_CORRECTION_TARGET_INVALID']


def test_reader_correction_cannot_add_other_admitted_evidence():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    correction = {'corrections': [{'object_id': 'NARR-dividends',
                                  'edits': [],
                                  'evidence_ids': ['PF-001', 'CE-001']}]}
    assessment = FakeClient(responses + [original, correction], [])
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Qualify the benefit.'}]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'fallback'
    assert len(result['priorities']) == 1
    assert 'FINAL_READER_CORRECTION_SCOPE_INVALID' in result['manifest']['validation_reason_codes']
    assert len(reviewer.calls) == 2


def test_reader_repairs_combined_precheck_issues_once():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    invalid = deepcopy(_final_reader())
    excess = ' Additional ' + 'supported access ' * 200 + 'remains.'
    invalid['summary_paragraphs'][0]['text'] += excess
    invalid['assessment_narratives'][2]['paragraphs'][0] += excess
    valid = _final_reader()
    patch = {'corrections': [{'object_id': record['object_id'], 'edits': ([{'paragraph_index': 0, 'old_text': excess, 'new_text': ''}] if index == 0 else [])}
                             for index, record in enumerate(valid['summary_paragraphs'])] + [
        {'object_id': 'NARR-inclusion', 'edits': [{'paragraph_index': 0, 'old_text': excess, 'new_text': ''}]}]}
    assessment = FakeClient(responses + [invalid, patch], [])
    reviewer = FakeClient([_native_pass_review] * 3, [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    assert {issue['object_id'] for issue in assessment.calls[-1]['payload']['bounded_correction']} == {
        record['object_id'] for record in valid['summary_paragraphs']} | {'NARR-inclusion'}
    assert [call['stage'] for call in assessment.calls].count('reader_compiler') == 1
    assert [call['stage'] for call in assessment.calls].count('reader_correction') == 1


def test_native_reader_common_noun_reaches_meaning_review():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    final['assessment_narratives'][2]['paragraphs'][0] += ' Protections include support for traders.'
    assessment = FakeClient(responses + [final], [])
    reviewer = FakeClient([_native_pass_review] * 2, [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    assert reviewer.calls[-1]['stage'] == 'reader_review'


def test_native_reader_uses_text_only_correction_and_retains_original():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    patch = {'corrections': [{'object_id': 'NARR-dividends', 'edits': [{'paragraph_index': 0, 'old_text': 'The project', 'new_text': 'This project'}]}]}
    assessment = FakeClient(responses + [original, patch], [])
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'],
         'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': [
             {'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
              'quoted_text': READER_QUOTE, 'explanation': 'Keep the benefit conditional.'}]},
        _native_pass_review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    assert result['final_reader']['summary_paragraphs'] == original['summary_paragraphs']
    assert result['final_reader']['assessment_narratives'][0] == original['assessment_narratives'][0]
    assert result['final_reader']['assessment_narratives'][1]['evidence_ids'] == original['assessment_narratives'][1]['evidence_ids']
    assert assessment.calls[-1]['stage'] == 'reader_correction'
    assert assessment.calls[-1]['max_transient_retries'] == 0
    trace = result['diagnostic_trace']['final_reader']
    assert json.loads(trace['initial']['json']) == original
    assert trace['correction']['target_ids'] == ['NARR-dividends']
    assert len(trace['meaning_reviews']) == 2


@pytest.mark.parametrize('verdict', ['pass', 'block'])
@pytest.mark.parametrize('word_count', [239, 260, 280, 500])
def test_modest_native_summary_overrun_reaches_meaning_review(verdict, word_count):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    count = sum(len(record['text'].split()) for record in final['summary_paragraphs'])
    assert count < word_count
    final['summary_paragraphs'][0]['text'] += ' ' + ' '.join(['supported'] * (word_count - count - 1) + ['support.'])
    review = {'verdict': verdict, 'object_ids': [], 'reason_codes': [], 'issues': []}
    if verdict == 'block':
        review.update(object_ids=['SUM-001'], reason_codes=['PROJECT_FACT_UNSUPPORTED'],
                      issues=[{'object_id': 'SUM-001', 'explanation': 'Planned support is stated as delivered.'}])
    assessment = FakeClient(responses + [final], [])
    reviewer = FakeClient([_native_pass_review, review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert reviewer.calls[-1]['stage'] == 'reader_review'
    assert result['summary_overview_status'] == ('generated' if verdict == 'pass' else 'fallback')
    assert not any(call['stage'] == 'reader_correction' for call in assessment.calls)
    assert len(result['priorities']) == 1
    from sector_lenses.climate_verified_schemas import validate_summary_overview
    diagnostics = {}
    assert validate_summary_overview({'paragraphs': [r['text'] for r in final['summary_paragraphs']]},
                                     diagnostics=diagnostics) == []
    assert diagnostics['reason'] == 'SUMMARY_WORD_COUNT'


def test_native_reader_reports_actual_summary_word_count():
    from sector_lenses.climate_final_reader import validate_final_reader

    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [_final_reader()], []), FakeClient([_native_pass_review] * 2, [])))
    final = _final_reader()
    final['summary_paragraphs'][0]['text'] += ' The project supports traders.' * 100
    count = sum(len(record['text'].split()) for record in final['summary_paragraphs'])
    issues = []
    assert validate_final_reader(final, result, issues=issues) == ([], [])
    assert any(str(count) in issue['explanation'] and '160-230' in issue['explanation'] for issue in issues)


@pytest.mark.parametrize('words', [20, 500, 501])
def test_native_summary_advisory_target_and_hard_ceiling(words):
    from sector_lenses.climate_verified_schemas import validate_summary_overview

    paragraphs = [' '.join(['Supported'] * (words // 2)) + '.',
                  ' '.join(['Supported'] * (words - words // 2)) + '.']
    diagnostics = {}
    result = validate_summary_overview({'paragraphs': paragraphs},
                                      semantic_review_required=True,
                                      diagnostics=diagnostics)
    assert bool(result) == (words <= 500)
    if words > 500:
        assert diagnostics['reason'] == 'SUMMARY_WORD_COUNT'
    assert not validate_summary_overview({'paragraphs': paragraphs})


def test_native_reader_rejected_patch_preserves_initial_for_diagnosis():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    patch = {'corrections': [{'object_id': 'NARR-dividends', 'edits': [],
                              'evidence_ids': ['PF-invented']} ]}
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'],
         'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': [
             {'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
              'quoted_text': READER_QUOTE, 'explanation': 'Keep the benefit conditional.'}]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [original, patch], []), reviewer))
    assert result['final_reader'] == original
    assert result['summary_overview_status'] == 'fallback'
    assert result['diagnostic_trace']['final_reader']['correction']['issues'][0]['reason_codes'] == ['READER_CORRECTION_FIELDS_INVALID']
    assert len(result['priorities']) == 1


@pytest.mark.parametrize('response', [
    {'corrections': []},
    {'corrections': [{'object_id': 'NARR-unknown', 'edits': []}]},
    {'corrections': [{'object_id': 'NARR-dividends', 'edits': []}] * 2},
    {'corrections': [{'object_id': 'NARR-dividends', 'edits': [], 'theme': 'design'}]},
    {'corrections': [{'object_id': 'NARR-dividends', 'edits': [{'paragraph_index': 0, 'old_text': 'The project', 'new_text': 'The project supports 987654 traders'}]}]},
    {'corrections': [{'object_id': 'NARR-dividends', 'edits': [{'paragraph_index': 0, 'old_text': 'The project', 'new_text': 'The project in January 2099'}]}]},
    {'corrections': [{'object_id': 'NARR-dividends', 'edits': [{'paragraph_index': 0, 'old_text': '', 'new_text': ''}]}]},
])
def test_native_reader_patch_rejects_scope_and_numeric_expansion(response):
    from sector_lenses.climate_final_reader import apply_reader_corrections

    original = _final_reader()
    snapshot = deepcopy(original)
    issues = []
    assert apply_reader_corrections(original, response, {'NARR-dividends'}, issues=issues) is None
    assert issues and issues[0]['reason_codes']
    assert original == snapshot


def test_reader_correction_payload_exposes_exact_existing_precision_and_rejects_new_date():
    from sector_lenses.climate_verified_prompts import build_verified_stage_prompt

    source_text = (
        "The Project Operations Manual will define site selection for 70 percent "
        "of sites by June 2028."
    )
    arguments = _arguments()
    arguments['source_blocks'] = [
        replace(arguments['source_blocks'][0], text=source_text)
    ]
    responses = _responses()
    fact = responses[0]['facts'][0]
    fact['object'] = 'site selection for 70 percent of sites by June 2028'
    fact['supporting_excerpt'] = source_text
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    original['assessment_narratives'][1]['paragraphs'][0] += (
        ' The schedule covers 70 percent of sites by June 2028.'
    )
    patch = {'corrections': [{'object_id': 'NARR-dividends',
                              'edits': [{'paragraph_index': 0, 'old_text': 'June 2028', 'new_text': 'June 2029'}]}]}
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'],
         'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': [
             {'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
              'quoted_text': READER_QUOTE, 'explanation': 'Keep the benefit conditional.'}]}], [])
    assessment = FakeClient(responses + [original, patch], [])

    result = run_verified_climate_pipeline(
        **arguments, clients=PipelineClients(assessment, reviewer))

    correction_payload = next(
        call['payload'] for call in assessment.calls
        if call['stage'] == 'reader_correction'
    )
    assert correction_payload['allowed_numeric_date_tokens_by_id'] == {
        'NARR-dividends': {
            'numeric_tokens': ['2028', '70'],
            'date_phrases': ['june 2028'],
        }
    }
    prompt = build_verified_stage_prompt('reader_correction', correction_payload)
    assert 'allowed_numeric_date_tokens_by_id' in prompt
    assert 'full listed date phrase' in ' '.join(prompt.split())
    assert result['summary_overview_status'] == 'fallback'
    issues = result['diagnostic_trace']['final_reader']['correction']['issues']
    assert issues[0]['reason_codes'] == [
        'READER_CORRECTION_SUMMARY_NUMBER_UNSUPPORTED'
    ]


@pytest.mark.parametrize('review', [
    {}, {'verdict': 'pass'},
    {'verdict': 'pass', 'reason_codes': [], 'object_ids': [], 'issues': None},
    {'verdict': 'pass', 'reason_codes': [], 'object_ids': None, 'issues': []},
])
def test_native_reader_requires_explicit_complete_clean_review(review):
    from sector_lenses.climate_final_reader import reviewed_final_reader

    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [_final_reader()], []), FakeClient([_native_pass_review] * 2, [])))
    assert reviewed_final_reader(_final_reader(), review, result) == ([], [])


@pytest.mark.parametrize('unsupported', [
    'Atlantis manages the markets.',
    'The project has completed all planned works.',
    'Country flood evidence proves every project site is currently flooded.',
    'The project will certainly prevent conflict.',
    'The project does not need any protections.',
])
def test_native_reader_semantic_block_never_publishes_prose(unsupported):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    final['assessment_narratives'][2]['paragraphs'][0] += ' ' + unsupported
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'block', 'object_ids': ['NARR-inclusion'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-inclusion', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Unsupported assertion.'}]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [final], []), reviewer))
    assert reviewer.calls[-1]['stage'] == 'reader_review'
    assert result['summary_overview_status'] == 'fallback'
    assert result['assessment_narratives'] == []
    assert len(result['priorities']) == 1


def test_reader_correction_review_checks_original_claim_scope():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    patch = {'corrections': [{'object_id': 'NARR-dividends', 'edits': []}]}
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Qualify the benefit.'}]}, _native_pass_review], [])
    assessment = FakeClient(responses + [original, patch], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    assert reviewer.calls[-1]['payload']['correction_originals'] == [original['assessment_narratives'][1]]
    correction_input = assessment.calls[-1]['payload']
    assert 'summary_max_words_by_id' not in correction_input
    assert 'final_reader' not in correction_input
    assert set(correction_input['support_records']) == set(original['assessment_narratives'][1]['evidence_ids'])


def test_reader_correction_exhausted_allowance_cannot_skip_fresh_review(monkeypatch):
    from sector_lenses.climate_verified_contracts import CALL_BUDGETS

    monkeypatch.setitem(CALL_BUDGETS, 'recommendation_compiler',
                        replace(CALL_BUDGETS['recommendation_compiler'], output_tokens=1))
    class OpaqueClient(FakeClient):
        @property
        def last_output_tokens(self):
            return None

    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    patch = {'corrections': [{'object_id': 'NARR-dividends', 'edits': []}]}
    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Qualify the benefit.'}]}], [])
    assessment = OpaqueClient(responses + [final, patch], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert assessment.calls[-1]['max_output_tokens'] == 1
    assert assessment.calls[-1]['max_transient_retries'] == 0
    assert result['manifest']['sampling']['correction_output_token_remaining'] == 0
    assert result['summary_overview_status'] == 'fallback'
    assert len(result['priorities']) == 1
    assert len(reviewer.calls) == 2


def test_reader_correction_provider_failure_preserves_initial_and_actions():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    class FailedCorrection(FakeClient):
        def complete_json(self, **kwargs):
            if kwargs['stage'] == 'reader_correction':
                self.calls.append(kwargs)
                raise ConnectionError('Provider unavailable')
            return super().complete_json(**kwargs)

    reviewer = FakeClient([_native_pass_review,
        {'verdict': 'revise', 'object_ids': ['NARR-dividends'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
         'issues': [{'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                     'quoted_text': READER_QUOTE, 'explanation': 'Qualify the benefit.'}]}], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FailedCorrection(responses + [final], []), reviewer))
    assert result['final_reader'] == final
    assert result['summary_overview_status'] == 'fallback'
    assert result['diagnostic_trace']['final_reader']['failure'] == 'ConnectionError'
    assert len(result['priorities']) == 1


def test_reader_diagnostics_never_enter_reader_or_exports(tmp_path):
    from docx import Document
    from sector_lenses.climate_verified_render import build_reader_model, render_reader_html, write_reader_docx

    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [_final_reader()], []), FakeClient([_native_pass_review] * 2, [])))
    result['diagnostic_trace']['final_reader']['initial']['json'] = 'PRIVATE_DIAGNOSTIC_SENTINEL'
    reader = build_reader_model(result)
    assert 'PRIVATE_DIAGNOSTIC_SENTINEL' not in json.dumps(reader)
    assert 'PRIVATE_DIAGNOSTIC_SENTINEL' not in render_reader_html(reader)
    document = Document(write_reader_docx(reader, tmp_path / 'reader.docx'))
    assert 'PRIVATE_DIAGNOSTIC_SENTINEL' not in '\n'.join(p.text for p in document.paragraphs)


def test_reader_combines_mechanical_and_meaning_issues_before_single_repair():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    excess = ' Additional ' + 'supported access ' * 200 + 'planned for July 2099.'
    original['summary_paragraphs'][0]['text'] += excess
    valid = _final_reader()
    targets = {r['object_id'] for r in valid['summary_paragraphs']} | {'NARR-dividends'}
    patch = {'corrections': [
        {'object_id': r['object_id'], 'edits': ([{'paragraph_index': 0, 'old_text': excess, 'new_text': ''}] if r['object_id'] == 'SUM-001' else [])}
        for r in valid['summary_paragraphs'] + valid['assessment_narratives']
        if r['object_id'] in targets]}
    revise = {'verdict': 'revise', 'object_ids': ['NARR-dividends'],
              'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': [
                  {'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                       'quoted_text': READER_QUOTE,
                   'explanation': 'Keep proposed benefits conditional.'}]}
    assessment = FakeClient(responses + [original, patch], [])
    reviewer = FakeClient([_native_pass_review, revise, _native_pass_review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'generated'
    correction = next(c['payload'] for c in assessment.calls if c['stage'] == 'reader_correction')
    assert {i['object_id'] for i in correction['bounded_correction']} == targets
    assert sum(correction['summary_max_words_by_id'].values()) == 220
    assert set(correction['summary_max_words_by_id']) == {r['object_id'] for r in valid['summary_paragraphs']}
    assert len(result['diagnostic_trace']['final_reader']['meaning_reviews']) == 2
    assert sum(c['stage'] == 'reader_correction' for c in assessment.calls) == 1


@pytest.mark.parametrize('review', [
    {'verdict': 'block', 'object_ids': ['SUM-001'], 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
     'issues': [{'object_id': 'SUM-001', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
                 'quoted_text': READER_QUOTE, 'explanation': 'False scope.'}]},
    {'verdict': 'pass'},
])
def test_reader_length_does_not_hide_block_or_malformed_meaning_review(review):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    original['summary_paragraphs'][0]['text'] += ' Supported project access.' * 60
    assessment = FakeClient(responses + [original], [])
    reviewer = FakeClient([_native_pass_review, review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert reviewer.calls[-1]['stage'] == 'reader_review'
    assert not any(c['stage'] == 'reader_correction' for c in assessment.calls)
    assert result['summary_overview_status'] == 'fallback'
    assert len(result['priorities']) == 1


def test_fact_and_review_prompts_require_comparable_contradictions_and_gap_premises():
    from sector_lenses.climate_verified_prompts import build_verified_stage_prompt
    facts = build_verified_stage_prompt('fact_extraction', {})
    review = build_verified_stage_prompt('conditional_review', {})
    for prompt in (facts, review):
        assert 'same subject or population' in prompt
        assert 'status' in prompt and 'period' in prompt and 'denominator' in prompt
        assert 'can both be true' in prompt
    assert 'Verify the residual-gap premise' in review
    assert 'not the residual gap' not in review


@pytest.mark.parametrize('defect', ['overlong_patch', 'unknown_evidence'])
def test_reader_repair_never_admits_overlong_patch_or_unknown_evidence(defect):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    original['summary_paragraphs'][0]['text'] += ' Supported project access.' * 120
    if defect == 'unknown_evidence':
        original['assessment_narratives'][0]['evidence_ids'] = ['PF-not-admitted']
    patch = {'corrections': [{'object_id': r['object_id'], 'edits': []}
                             for r in original['summary_paragraphs']]}
    assessment = FakeClient(responses + [original, patch], [])
    reviewer = FakeClient([_native_pass_review] * 3, [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(assessment, reviewer))
    assert result['summary_overview_status'] == 'fallback'
    assert len(result['priorities']) == 1
    assert sum(c['stage'] == 'reader_correction' for c in assessment.calls) == (defect == 'overlong_patch')
    assert sum(c['stage'] == 'reader_review' for c in reviewer.calls) == (defect == 'overlong_patch')


@pytest.mark.parametrize('defect', ['missing_codes', 'empty_codes', 'missing_explanation', 'duplicate_ids', 'mismatched_codes'])
def test_malformed_reader_revision_does_not_spend_correction(defect):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    issue = {'object_id': 'NARR-dividends', 'reason_codes': ['PROJECT_FACT_UNSUPPORTED'],
             'quoted_text': READER_QUOTE,
             'explanation': 'Keep the proposed benefit conditional.'}
    review = {'verdict': 'revise', 'object_ids': ['NARR-dividends'],
              'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': [issue]}
    if defect == 'missing_codes':
        del review['reason_codes']
    elif defect == 'empty_codes':
        issue['reason_codes'] = []
    elif defect == 'missing_explanation':
        del issue['explanation']
    elif defect == 'duplicate_ids':
        review['object_ids'] *= 2
    else:
        review['reason_codes'] = ['RESIDUAL_GAP_UNSUPPORTED']
    assessment = FakeClient(responses + [_final_reader()], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        assessment, FakeClient([_native_pass_review, review], [])))
    assert not any(c['stage'] == 'reader_correction' for c in assessment.calls)
    assert result['summary_overview_status'] == 'fallback'


def test_reader_markup_is_repaired_after_meaning_review():
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    original = _final_reader()
    original['summary_paragraphs'][0]['text'] = '**' + original['summary_paragraphs'][0]['text'] + '**'
    valid = _final_reader()
    patch = {'corrections': [{'object_id': r['object_id'], 'edits': ([
        {'paragraph_index': 0, 'old_text': '**' + r['text'].split()[0], 'new_text': r['text'].split()[0]},
        {'paragraph_index': 0, 'old_text': r['text'].split()[-1] + '**', 'new_text': r['text'].split()[-1]},
    ] if r['object_id'] == 'SUM-001' else [])}
                             for r in valid['summary_paragraphs']]}
    assessment = FakeClient(responses + [original, patch], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        assessment, FakeClient([_native_pass_review] * 3, [])))
    assert result['summary_overview_status'] == 'generated'
    assert len(result['diagnostic_trace']['final_reader']['meaning_reviews']) == 2
    assert 'summary_max_words_by_id' not in assessment.calls[-1]['payload']


@pytest.mark.parametrize('words', [202, 400, 401])
def test_reader_reports_unknown_evidence_even_with_length_defects(words):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    final['assessment_narratives'][0]['paragraphs'] = [' '.join(['Supported'] * words) + '.']
    final['assessment_narratives'][0]['evidence_ids'] = ['PF-invented']
    reviewer = FakeClient([_native_pass_review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [final], []), reviewer))
    assert result['summary_overview_status'] == 'fallback'
    assert 'READER_EVIDENCE_INVALID' in result['validation']['reason_codes']
    codes = {code for issue in result['diagnostic_trace']['final_reader']['prechecks'][0]['issues']
             for code in issue['reason_codes']}
    assert 'READER_EVIDENCE_INVALID' in codes
    assert ('READER_WORD_COUNT' in codes) == (words > 400)
    assert not any(call['stage'] == 'reader_review' for call in reviewer.calls)


@pytest.mark.parametrize('verdict', ['pass', 'block'])
@pytest.mark.parametrize('words', [20, 185, 202, 400])
def test_minor_reader_length_overrun_still_requires_meaning_review(verdict, words):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    text = final['assessment_narratives'][0]['paragraphs'][0]
    text = ' '.join(text.split()[:words])
    text += ' ' + ' '.join(['Supported'] * (words - len(text.split()))) + '.'
    final['assessment_narratives'][0]['paragraphs'] = [text]
    review = _native_pass_review if verdict == 'pass' else {
        'verdict': 'block', 'object_ids': ['NARR-design'],
        'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': []}
    reviewer = FakeClient([_native_pass_review, review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [final], []), reviewer))
    assert reviewer.calls[-1]['stage'] == 'reader_review'
    assert result['summary_overview_status'] == ('generated' if verdict == 'pass' else 'fallback')


@pytest.mark.parametrize('verdict', ['pass', 'block'])
def test_reader_paragraph_cleanup_preserves_text_and_requires_meaning_review(verdict):
    responses = _responses()
    responses[2]['reader_contract'] = 'final-reader-v1'
    final = _final_reader()
    original = final['assessment_narratives'][0]['paragraphs'][0]
    words = original.split()
    final['assessment_narratives'][0]['paragraphs'] = [
        ' '.join(words[:20]), '.', ' '.join(words[20:40]), '', ' '.join(words[40:])]
    review = _native_pass_review if verdict == 'pass' else {
        'verdict': 'block', 'object_ids': ['NARR-design'],
        'reason_codes': ['PROJECT_FACT_UNSUPPORTED'], 'issues': []}
    reviewer = FakeClient([_native_pass_review, review], [])
    result = run_verified_climate_pipeline(**_arguments(), clients=PipelineClients(
        FakeClient(responses + [final], []), reviewer))
    assert reviewer.calls[-1]['stage'] == 'reader_review'
    reviewed = reviewer.calls[-1]['payload']['final_reader']['assessment_narratives'][0]
    assert ' '.join(reviewed['paragraphs']) == original
    assert len(reviewed['paragraphs']) <= 2
    assert result['summary_overview_status'] == ('generated' if verdict == 'pass' else 'fallback')
