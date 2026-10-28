"""Contracts for the normal-FCV detailed DOCX advisory."""

from io import BytesIO

import pytest
from docx import Document

from app import app


FCV_ADVISORY = (
    "The following priorities are suggestions to strengthen the project's design "
    "in its FCV context. They are not mandatory requirements, nor may they all be "
    "relevant. The task team may wish to focus on those they believe to be more "
    "pertinent based on their own sectoral expertise and knowledge of the context "
    "within which the project will be applied. For more complex or difficult issues "
    "raised, the team are encouraged to reach out and discuss them with the FCV "
    "Country Coordinator or relevant Global Practice experts as needed."
)
IMPLEMENTATION_ADVISORY = FCV_ADVISORY.replace(
    "project's design in its FCV context",
    "project's design and implementation arrangements in its FCV context",
)


def _docx_text(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


@pytest.mark.parametrize(
    ("metadata", "expected_advisory"),
    [
        ({}, FCV_ADVISORY),
        ({"review_mode": "implementation"}, IMPLEMENTATION_ADVISORY),
    ],
)
def test_normal_docx_export_advisory_uses_review_mode_metadata(
    metadata, expected_advisory
):
    response = app.test_client().post(
        "/api/download-report",
        json={
            "summary": "",
            "priorities": [
                {
                    "title": "Priority 1 · Keep the alternate delivery route available",
                    "the_gap": "The route is not yet recorded in the project document.",
                    "why_it_matters": "Access can change during implementation.",
                    "actions": [
                        {
                            "document_element": "Operations Manual",
                            "guidance": "Record the alternate route and responsible unit.",
                            "suggested_language": (
                                "The project will document an alternate delivery route."
                            ),
                        }
                    ],
                    "when": "Implementation",
                    "who_acts": "Project implementation unit",
                }
            ],
            "metadata": metadata,
        },
    )

    assert response.status_code == 200, response.get_data(as_text=True)
    document = Document(BytesIO(response.data))
    body = _docx_text(document)
    assert body.count(expected_advisory) == 1
    assert "The project will document an alternate delivery route." in body
    assert "Draft - human review required" not in body
