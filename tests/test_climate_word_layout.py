"""Word layout contracts for live Climate-FCV and offline reader exports."""

import json
import re
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

import build_climate_readout
import build_fcv_readout
from fcv_docx_style import AMBER_TXT, BE_AWARE, BLUE, METER_OFF, NAVY, logo_path
from sector_lenses.climate_verified_render import (
    CANDIDATE_PREVIEW_WARNING,
    DIMENSIONS,
    HEADINGS,
    SMOKE_RUNTIME_WARNING,
    attach_provenance,
    build_reader_model,
    write_reader_docx,
)
from sector_lenses.pipeline import normalize_climate_assessment
from tests.test_climate_summary_completeness import _run_multiple
from tests.test_climate_verified_render import _assessment


FIXTURES = Path(__file__).resolve().parent / "fixtures"
SNAPSHOT = "a" * 64


def _body_blocks(document):
    """Return body paragraphs and tables, including nested callout content."""
    def visit(element):
        for child in element.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                yield table
                seen_cells = set()
                for row in table.rows:
                    for cell in row.cells:
                        if cell._tc in seen_cells:
                            continue
                        seen_cells.add(cell._tc)
                        yield from visit(cell._tc)

    yield from visit(document.element.body)


def _block_text(block):
    if isinstance(block, Paragraph):
        return block.text
    return "\n".join(
        cell.text for row in block.rows for cell in row.cells
    )


def _all_body_text(blocks):
    return "\n".join(
        block.text for block in blocks if isinstance(block, Paragraph)
    )


def _contains_index(blocks, text):
    return next(
        index for index, block in enumerate(blocks) if text in _block_text(block)
    )


def _heading_index(blocks, text):
    return next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Paragraph)
        and block.style.name.startswith("Heading")
        and text.casefold() in block.text.casefold()
    )


def _priority_overview_table_index(blocks, titles):
    return next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Table)
        and len(block.columns) >= 2
        and "Operational priority" in _block_text(block)
        and all(title in _block_text(block) for title in titles)
    )


def _scorecard_table_index(blocks):
    return next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Table)
        and len(block.columns) == 3
        and "SENSITIVITY" in _block_text(block).upper()
        and "RESPONSIVENESS" in _block_text(block).upper()
    )


def _rating_heading_index(blocks, label, after):
    return next(
        index
        for index, block in enumerate(blocks)
        if index > after
        and isinstance(block, Paragraph)
        and block.style.name.startswith("Heading")
        and label.casefold() in block.text.casefold()
    )


def _first_sentence(text):
    match = re.search(r".*?[.!?](?:\s|$)", text.strip())
    return match.group(0).strip() if match else text.strip()


def _verified_download_payload():
    assessment, _, _ = _run_multiple(1)
    assessment = normalize_climate_assessment(json.loads(json.dumps(assessment)))
    return {
        "climate_assessment": assessment,
        "climate_reader": attach_provenance(
            build_reader_model(assessment), assessment
        ),
    }


def _part_text(part):
    return "\n".join(paragraph.text for paragraph in part.paragraphs) + "\n" + "\n".join(
        cell.text for table in part.tables for row in table.rows for cell in row.cells
    )


def _assert_approved_word_brand(document, report_title, expected_project_line=None):
    section = document.sections[0]
    first_page_header = section.first_page_header
    assert section.header_distance.inches == pytest.approx(0)
    logo_bytes = Path(logo_path(True)).read_bytes()
    assert any(
        relationship.reltype == RT.IMAGE
        and relationship.target_part.blob == logo_bytes
        for relationship in first_page_header.part.rels.values()
    )
    assert any(
        cell._tc.get_or_add_tcPr().find(qn("w:shd")) is not None
        and cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")) == NAVY
        for table in first_page_header.tables
        for row in table.rows
        for cell in row.cells
    )
    banner_position = first_page_header.tables[0]._tbl.tblPr.find(
        qn("w:tblpPr")
    )
    assert banner_position is not None
    assert banner_position.get(qn("w:vertAnchor")) == "page"
    assert banner_position.get(qn("w:horzAnchor")) == "page"
    assert banner_position.get(qn("w:tblpY")) == "0"
    assert banner_position.get(qn("w:tblpX")) == "0"
    banner_text = _part_text(first_page_header)
    assert report_title in banner_text
    if expected_project_line is not None:
        assert expected_project_line in banner_text
    assert "FCV SCREENER" in _part_text(section.header)
    footer_text = _part_text(section.footer)
    assert "AI-assisted output" in footer_text
    assert "verify before operational use" in footer_text
    assert not any(paragraph.style.name == "Title" for paragraph in document.paragraphs)

    heading_colors = set()
    for paragraph in (
        block for block in _body_blocks(document)
        if isinstance(block, Paragraph)
    ):
        if not paragraph.style.name.startswith("Heading"):
            continue
        style_color = paragraph.style.font.color.rgb
        if style_color is not None:
            heading_colors.add(str(style_color))
        for run in paragraph.runs:
            run_color = run.font.color.rgb
            if run_color is not None:
                heading_colors.add(str(run_color))
    assert {NAVY, BLUE} <= heading_colors


def test_live_climate_word_orders_reader_sections_and_keeps_specific_review_signals(
    tmp_path,
):
    assessment = _assessment()
    assessment["operation_context"] = {
        "document_type": "PCN",
        "instrument_type": "unknown",
        "preparation_regime": "unresolved_policy_source",
    }
    model = build_reader_model(assessment)
    model["summary_overview"] = [
        "Overview summary marker: the project evidence supports this assessment. "
        "[Verify before sharing: confirm project scope against the signed document.]"
    ]
    model["priorities"][0]["caution"] = (
        "Source caveat. [Verify before sharing: check the signed annex for the "
        "priority scope.]"
    )
    model["reader_review_notice"] = "Generic reader review caveat must stay absent."
    model["runtime_mode"] = "smoke"
    model["analysis_status"] = "partial"
    model["analysis_message"] = "Confirm coverage for remaining project locations."
    model["recommendation_status"] = "partial"
    model["recommendation_message"] = "One proposed action needs confirmation."
    from app import _word_report_meta

    report_meta = _word_report_meta(
        {
            "metadata": {
                "project_name": "Live layout test project",
                "project_id": "P123456",
                "country": "Testland",
                "document_type": "PCN",
                "instrument": "IPF",
                "date_str": "2026-10-27",
            }
        },
        title="Climate and FCV Recommendations Note",
    )
    path = tmp_path / "climate-reader.docx"

    write_reader_docx(
        model, path, snapshot_sha256=SNAPSHOT, report_meta=report_meta
    )

    document = Document(path)
    _assert_approved_word_brand(
        document,
        "Climate and FCV Recommendations Note",
        report_meta["project_line"],
    )
    blocks = list(_body_blocks(document))
    body = _all_body_text(blocks)
    advisory_index = _contains_index(blocks, BE_AWARE[0])
    overview_index = _heading_index(blocks, "Overview")
    overview_text_index = _contains_index(blocks, model["summary_overview"][0])
    findings_index = _heading_index(blocks, "Assessment findings")
    ratings_heading_index = _heading_index(
        blocks, "Climate and FCV Sensitivity and Responsiveness"
    )
    sensitivity_heading_index = _rating_heading_index(
        blocks, "sensitivity", ratings_heading_index
    )
    responsiveness_heading_index = _rating_heading_index(
        blocks, "responsiveness", ratings_heading_index
    )
    scorecard_index = _scorecard_table_index(blocks)
    status_indices = [
        _contains_index(blocks, "Assessment evidence is incomplete"),
        _contains_index(blocks, "Recommendations need review"),
        _contains_index(blocks, SMOKE_RUNTIME_WARNING),
        _contains_index(blocks, CANDIDATE_PREVIEW_WARNING),
    ]
    priorities_heading_index = _heading_index(blocks, HEADINGS[2])
    priority_intro_index = _contains_index(
        blocks, model["priority_summary"]["statement"]
    )
    priority_table_index = _priority_overview_table_index(
        blocks, model["priority_summary"]["titles"]
    )
    priority_card_index = next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Table)
        and "PRIORITY 1" in _block_text(block)
        and model["priorities"][0]["title"] in _block_text(block)
    )
    detailed_priority_index = _contains_index(
        blocks, model["priorities"][0]["minimum_action"]
    )

    assert advisory_index < scorecard_index < min(status_indices)
    assert max(status_indices) < overview_index < overview_text_index
    assert overview_text_index < findings_index < ratings_heading_index
    assert scorecard_index < overview_index
    assert ratings_heading_index < sensitivity_heading_index
    assert sensitivity_heading_index < responsiveness_heading_index
    assert responsiveness_heading_index < priorities_heading_index
    rating_section_text = "\n".join(
        _block_text(block)
        for block in blocks[ratings_heading_index:priorities_heading_index]
    )
    for dimension in (DIMENSIONS[1], DIMENSIONS[2]):
        assert dimension[1].casefold() in rating_section_text.casefold()
        assert dimension[2].casefold() in rating_section_text.casefold()
        assert assessment["judgments"][dimension[0]]["rationale"].casefold() in rating_section_text.casefold()
    assert priorities_heading_index < priority_intro_index < priority_table_index
    assert priority_table_index < priority_card_index < detailed_priority_index
    assert all(text in body for text in BE_AWARE)
    assert "Generic reader review caveat must stay absent." not in body
    assert "Assessment evidence is incomplete" in body
    assert "Recommendations need review" in body
    assert SMOKE_RUNTIME_WARNING in body
    assert CANDIDATE_PREVIEW_WARNING in body
    assert "Two sections state different financing totals." in body
    assert "Confirm the controlling total." in body
    assert "Suggested document wording is not shown because" in body
    assert f"Review snapshot: {SNAPSHOT}" in body
    assert "[Verify before sharing: confirm project scope against the signed document.]" in body
    practical_qualification = (
        "[Verify before sharing: check the signed annex for the priority scope.]"
    )
    assert any(
        practical_qualification in paragraph.text
        for paragraph in (block for block in blocks if isinstance(block, Paragraph))
    )

    qualification_run = next(
        run
        for paragraph in (
            block for block in blocks if isinstance(block, Paragraph)
        )
        for run in paragraph.runs
        if "Verify before sharing" in run.text
    )
    assert qualification_run.font.color.rgb == RGBColor.from_string(AMBER_TXT)
    practical_qualification_runs = [
        run
        for paragraph in (block for block in blocks if isinstance(block, Paragraph))
        for run in paragraph.runs
        if "check the signed annex for the priority scope" in run.text
    ]
    assert practical_qualification_runs
    assert all(
        run.font.color.rgb == RGBColor.from_string(AMBER_TXT)
        for run in practical_qualification_runs
    )


def test_legacy_climate_word_without_priorities_has_no_empty_priority_table(
    tmp_path,
):
    assessment = _assessment()
    assessment["priorities"] = []
    assessment["judgments"]["sensitivity"]["value"] = "unclear"
    assessment["judgments"]["responsiveness"]["value"] = "not_expected"
    model = build_reader_model(assessment)
    path = tmp_path / "legacy-climate-reader.docx"

    write_reader_docx(model, path)

    blocks = list(_body_blocks(Document(path)))
    body = _all_body_text(blocks)
    paragraphs = [block.text.strip() for block in blocks if isinstance(block, Paragraph)]
    assert "Strategic priorities" not in paragraphs
    assert not any(
        isinstance(block, Table)
        and len(block.columns) == 2
        and block.rows
        and block.rows[0].cells[0].text.strip() in {"#", "Rank"}
        and "priority" in block.rows[0].cells[1].text.casefold()
        for block in blocks
    )
    assert "Unclear" in body
    assert "Not expected for this project" in body


def test_climate_docx_preserves_standalone_qualification_in_amber(tmp_path):
    model = build_reader_model(_assessment())
    qualification = (
        "[Verify before sharing: confirm the documented financing total.]"
    )
    model["summary_overview"] = [qualification]
    path = tmp_path / "climate-qualification.docx"

    write_reader_docx(model, path)

    document = Document(path)
    paragraphs = [
        block for block in _body_blocks(document)
        if isinstance(block, Paragraph)
    ]
    assert any(paragraph.text == qualification for paragraph in paragraphs)
    highlighted = [
        run
        for paragraph in paragraphs
        for run in paragraph.runs
        if "Verify before sharing" in run.text
    ]
    assert highlighted
    assert all(
        run.font.color.rgb == RGBColor.from_string(AMBER_TXT)
        for run in highlighted
    )


def test_normal_report_download_opens_with_shared_advisory_and_narrative_before_ratings():
    from app import _word_report_meta, app

    payload = {
        "summary": (
            "# Narrative title that should not override metadata\n\n"
            "Narrative ordering marker: the project context is set out first."
        ),
        "metadata": {
            "project_name": "Supplied metadata project",
            "project_id": "P765432",
            "country": "Testland",
            "document_type": "PCN",
            "instrument": "IPF",
            "date_str": "2026-10-27",
        },
        "priorities": [
            {
                "title": "Priority 1 - Keep the alternate delivery route available",
                "the_gap": "The route is not yet recorded in the project document.",
                "why_it_matters": "Access can change during implementation.",
                "actions": [
                    {
                        "document_element": "Operations Manual",
                        "guidance": "Record the alternate route and responsible unit.",
                        "suggested_language": (
                            "The project will document an alternate delivery route."
                        ),
                    }
                ],
                "when": "Implementation",
                "who_acts": "Project implementation unit",
                "project_cycle": {
                    "primary_label": "At concept stage",
                    "primary_text": "Confirm the route against the approved project cycle.",
                },
            }
        ],
        "fcv_rating": "Well Embedded",
        "sensitivity_summary": "Rating rationale marker for sensitivity.",
    }
    response = app.test_client().post("/api/download-report", json=payload)

    assert response.status_code == 200, response.get_data(as_text=True)
    document = Document(BytesIO(response.data))
    report_meta = _word_report_meta(payload)
    _assert_approved_word_brand(
        document, "FCV Recommendations Note", report_meta["project_line"]
    )
    blocks = list(_body_blocks(document))
    body = _all_body_text(blocks)
    narrative_index = _contains_index(blocks, "Narrative ordering marker")
    advisory_indices = [_contains_index(blocks, paragraph) for paragraph in BE_AWARE]
    scorecard_index = _scorecard_table_index(blocks)
    first_assessment_heading_index = next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Paragraph)
        and block.style.name.startswith("Heading")
    )
    ratings_index = _heading_index(blocks, "FCV Sensitivity")
    assert max(advisory_indices) < scorecard_index < first_assessment_heading_index
    assert scorecard_index < narrative_index < ratings_index
    assert all(text in body for text in BE_AWARE)
    assert "At concept stage" in body
    assert "Supplied metadata project" in _part_text(document.sections[0].first_page_header)
    assert "Narrative title that should not override metadata" not in _part_text(
        document.sections[0].first_page_header
    )
    assert "Narrative title that should not override metadata" not in body
    assert "Generated by WBG FCV Project Screener" not in body
    assert "AI-assisted output. Analytical framework:" not in body


@pytest.mark.parametrize("verified", [False, True], ids=["normal", "climate-verified"])
def test_live_download_routes_use_the_approved_word_brand(verified):
    from app import app

    expected_project_line = None
    if verified:
        payload = _verified_download_payload()
        expected_title = "FCV Recommendations Note"
    else:
        fallback_title = "Somalia Water Access Operation"
        payload = {
            "summary": (
                f"# {fallback_title}\n\n"
                "## Overview\n\n"
                "A short narrative keeps the route local and avoids model calls."
            ),
            "priorities": [],
            "fcv_rating": "Well Embedded",
        }
        expected_title = "FCV Recommendations Note"
        expected_project_line = fallback_title

    response = app.test_client().post("/api/download-report", json=payload)

    assert response.status_code == 200, response.get_data(as_text=True)
    document = Document(BytesIO(response.data))
    _assert_approved_word_brand(document, expected_title, expected_project_line)
    if expected_project_line:
        body = _all_body_text(list(_body_blocks(document)))
        assert expected_project_line not in body


@pytest.mark.parametrize(
    ("builder", "fixture_name", "narrative_markers"),
    [
        (
            build_climate_readout,
            "climate_readout_sample.json",
            lambda data: [
                _first_sentence(data["overview"][0]),
                _first_sentence(data["findings"][0]["paras"][0]),
            ],
        ),
        (
            build_fcv_readout,
            "fcv_readout_sample.json",
            lambda data: [
                _first_sentence(data["exec"]["opening"]),
                _first_sentence(data["risk"]["How FCV risks could affect this project"][0]),
            ],
        ),
    ],
)
def test_offline_builder_scorecard_precedes_narrative_and_detail_ratings(
    builder, fixture_name, narrative_markers, tmp_path
):
    data = json.loads((FIXTURES / fixture_name).read_text(encoding="utf-8"))
    path = tmp_path / f"{builder.__name__}.docx"
    builder.render(
        data,
        {
            "title": "Reader layout test",
            "project_line": "Test project",
            "running_label": "Test",
            "classification": "Official Use Only",
            "generator": "",
            "disclaimer": "",
            "its_url": "https://fcvriskdashboardqa.worldbank.org/project-screener",
        },
        str(path),
    )

    blocks = list(_body_blocks(Document(path)))
    markers = [
        _contains_index(blocks, marker) for marker in narrative_markers(data)
    ]
    scorecard_index = next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Table)
        and len(block.columns) == 3
        and "FCV RESPONSIVENESS" in _block_text(block).upper()
    )
    advisory_index = _contains_index(blocks, BE_AWARE[0])
    first_assessment_heading_index = next(
        index
        for index, block in enumerate(blocks)
        if isinstance(block, Paragraph)
        and block.style.name.startswith("Heading")
    )
    detail_ratings_index = next(
        index
        for index, block in enumerate(blocks)
        if index > max(markers)
        and isinstance(block, Paragraph)
        and block.style.name.startswith("Heading")
        and "sensitivity" in block.text.casefold()
        and "responsiveness" in block.text.casefold()
    )
    assert advisory_index < scorecard_index < first_assessment_heading_index
    assert scorecard_index < min(markers) <= max(markers) < detail_ratings_index


def test_offline_fcv_unknown_scale6_rating_remains_unscored(tmp_path):
    data = json.loads(
        (FIXTURES / "fcv_readout_sample.json").read_text(encoding="utf-8")
    )
    data["ratings"]["FCV Sensitivity"]["value"] = "Not yet assessed"
    path = tmp_path / "fcv-unknown-rating.docx"

    build_fcv_readout.render(
        data,
        {
            "title": "Unknown rating test",
            "project_line": "",
            "running_label": "",
            "classification": "",
            "generator": "",
            "disclaimer": "",
            "its_url": "",
        },
        str(path),
    )

    blocks = list(_body_blocks(Document(path)))
    scorecard = next(
        block
        for block in blocks
        if isinstance(block, Table)
        and len(block.columns) == 3
        and "FCV SENSITIVITY" in _block_text(block).upper()
        and "NOT YET ASSESSED" in _block_text(block).upper()
    )
    assert "Not yet assessed" in _block_text(scorecard)
    meter = scorecard.cell(0, 0).tables[0]
    fills = [
        cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill"))
        for cell in meter.rows[0].cells
    ]
    assert fills == [METER_OFF] * 6
