"""No-provider checks for exact Climate export snapshot consistency."""
import json
from io import BytesIO

import pytest
from docx import Document
import app as app_module
from sector_lenses.pipeline import normalize_climate_assessment
from sector_lenses.climate_verified_render import build_reader_model, attach_provenance
from tests.test_climate_summary_completeness import _run_multiple


@pytest.fixture
def pair():
    assessment, _, _ = _run_multiple(1)
    assessment = normalize_climate_assessment(json.loads(json.dumps(assessment)))
    return json.loads(json.dumps({"climate_assessment": assessment,
            "climate_reader": attach_provenance(build_reader_model(assessment), assessment)}))


def test_snapshot_and_word_share_exact_reader(pair):
    client = app_module.app.test_client()
    response = client.post('/api/climate-report-snapshot', json=pair)
    assert response.status_code == 200
    snapshot = response.get_json()
    assert snapshot['reader'] == pair['climate_reader']
    assert snapshot['publication_approved'] is False
    assert len(snapshot['sha256']) == 64
    word = client.post('/api/download-report', json=pair)
    assert word.status_code == 200
    text = '\n'.join(p.text for p in Document(BytesIO(word.data)).paragraphs)
    assert pair['climate_reader']['priorities'][0]['minimum_action'] in text


@pytest.mark.parametrize('route', ['/api/climate-report-snapshot', '/api/download-report'])
@pytest.mark.parametrize('field', ['executive_readout', 'sources'])
def test_reader_only_changes_fail_instead_of_silent_rebuild(pair, route, field):
    pair['climate_reader'][field] = 'Changed in browser only.'
    result = app_module.app.test_client().post(route, json=pair)
    assert result.status_code == 409
    assert 'READER_SNAPSHOT_MISMATCH' in result.get_json()['reason_codes']


def test_fingerprint_binds_assessment_even_if_visible_reader_unchanged(pair):
    client = app_module.app.test_client()
    before = client.post('/api/climate-report-snapshot', json=pair).get_json()
    pair['climate_assessment']['run_id'] = 'another-run'
    pair['climate_reader'] = attach_provenance(build_reader_model(pair['climate_assessment']), pair['climate_assessment'])
    after = client.post('/api/climate-report-snapshot', json=pair).get_json()
    assert before['sha256'] != after['sha256']


@pytest.mark.parametrize('reader', [None, [], {}, 'invalid'])
def test_snapshot_requires_complete_reader(pair, reader):
    pair['climate_reader'] = reader
    result = app_module.app.test_client().post('/api/climate-report-snapshot', json=pair)
    assert result.status_code in (409, 422)


def test_canonical_correction_survives_both_exports(pair):
    from sector_lenses.climate_verified_render import render_reader_html
    correction = 'Check seasonal access before finalizing the design.'
    pair['climate_assessment']['priorities'][0]['minimum_action'] = correction
    pair['climate_reader'] = attach_provenance(build_reader_model(pair['climate_assessment']), pair['climate_assessment'])
    client = app_module.app.test_client()
    snapshot = client.post('/api/climate-report-snapshot', json=pair)
    assert snapshot.status_code == 200
    assert correction in render_reader_html(snapshot.get_json()['reader'])
    word = client.post('/api/download-report', json=pair)
    assert word.status_code == 200
    assert correction in '\n'.join(p.text for p in Document(BytesIO(word.data)).paragraphs)


@pytest.mark.parametrize('scenario', ['ok', 'reject', 'changed', 'offline', 'missing'])
def test_browser_html_waits_for_consistency_check(scenario):
    import subprocess
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function
    function = 'async ' + _extract_js_function(INDEX.read_text(encoding='utf-8'), 'downloadHTML')
    script = """
let stageOutputs={3:'ready'},stageThreePriorities=[];
let climateVerifiedAssessment={schema_version:'climate-verified-v2.1'};
let climateVerifiedReader={executive_readout:'Corrected text.'};
let alerts=[],downloads=0, rendered=0;
const alert=x=>alerts.push(x);
const document={querySelectorAll:()=>[],body:{appendChild:()=>{}},createElement:()=>({click:()=>downloads++,remove:()=>{}})};
const URL={createObjectURL:()=>'',revokeObjectURL:()=>{}};
const renderClimateVerifiedAssessment=r=>{rendered++;return r.executive_readout;};
const climatePrintDisclosureScript=()=>'';
const scenario=SCENARIO;
if(scenario==='missing')climateVerifiedReader=null;
const fetch=async()=>{
 if(scenario==='offline')throw new Error('offline');
 if(scenario==='changed')climateVerifiedReader.executive_readout='Changed while checking.';
 return {ok:!['reject','missing'].includes(scenario),json:async()=>({error:'Mismatch'})};
};
FUNCTION
(async()=>{await downloadHTML();
 const ok=scenario==='ok';
 if(downloads!==(ok?1:0)||rendered!==(ok?1:0)||alerts.length!==(ok?0:1))throw new Error(JSON.stringify({downloads,rendered,alerts}));
})().catch(e=>{console.error(e);process.exit(1)});
""".replace('SCENARIO', json.dumps(scenario)).replace('FUNCTION', function)
    result = subprocess.run(['node', '-e', script], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_runtime_reader_survives_json_round_trip(monkeypatch):
    import sector_lenses.climate_verified_runtime as runtime
    from tests.test_climate_verified_runtime import _doc_parts
    assessment, _, _ = _run_multiple(1)
    monkeypatch.setattr(runtime, 'run_verified_climate_pipeline', lambda **kwargs: assessment)
    result = runtime.run_verified_from_doc_parts(
        doc_parts=_doc_parts(), clients=None, climate_grounding={}, run_id='snapshot-test')
    payload = json.loads(json.dumps({'climate_assessment': result['assessment'],
                                    'climate_reader': result['reader']}, sort_keys=True))
    response = app_module.app.test_client().post('/api/climate-report-snapshot', json=payload)
    assert response.status_code == 200, response.get_json()
