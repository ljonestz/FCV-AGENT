"""Shared sidebar layout with native Climate rating categories."""

import json
import re
import subprocess

import pytest

from tests.test_climate_lens_frontend import INDEX, _extract_js_function, _js_escape_helper


def render_cards(responsiveness="moderate"):
    source = INDEX.read_text(encoding="utf-8")
    reader = {"judgments": [
        {"dimension": "sensitivity", "value": "moderate", "rationale": "Recognises risk."},
        {"dimension": "responsiveness", "value": responsiveness, "rationale": "Scope matters."},
    ]}
    script = (_js_escape_helper()
              + _extract_js_function(source, "climateVerifiedRatingCards")
              + "\nconsole.log(climateVerifiedRatingCards(" + json.dumps(reader) + "));")
    result = subprocess.run(["node", "-e", script], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return result.stdout


def test_verified_climate_has_two_native_rating_cards():
    html = render_cards()
    assert 'data-rating-key="climate-sensitivity"' in html
    assert 'data-rating-key="climate-responsiveness"' in html
    assert "Climate &amp; FCV sensitivity" in html
    assert "Moderate contribution" in html
    assert "Recognises risk." in html and "Scope matters." in html
    assert html.count('aria-valuemax="5"') == 2
    assert "integration rating" not in html


def test_not_expected_and_unclear_are_not_low_responsiveness_scores():
    for value, label in [("not_expected", "Not expected"), ("unclear", "Unclear")]:
        html = render_cards(value).split('data-rating-key="climate-responsiveness"')[1]
        assert label in html
        assert 'role="progressbar"' not in html
        assert "Scope matters." in html


def test_sidebar_glosses_are_short_and_full_rationale_is_collapsed():
    html = render_cards()
    glosses = re.findall(r'<p class="stage3-rating-meaning">(.*?)</p>', html)
    assert len(glosses) == 2
    assert all(len(gloss.split()) <= 20 for gloss in glosses)
    assert "Recognises risk." not in glosses
    assert html.count('<details class="stage3-rating-explanation">') == 2
    assert "<summary>Why this rating</summary>" in html


def test_responsiveness_is_additional_contribution_and_legacy_is_unscored():
    html = render_cards("emerging").split('data-rating-key="climate-responsiveness"')[1]
    assert "Emerging (legacy scale)" in html
    assert 'role="progressbar"' not in html
    assert "Contributing to responsiveness is often a sign of going above and beyond." in html
    assert "The expected contribution depends on the project&#039;s scope." in html


def test_responsiveness_contract_has_five_levels_and_unscored_exceptions():
    from sector_lenses.climate_judgments import ALLOWED
    assert ALLOWED["responsiveness"] == {
        "very_limited", "limited", "moderate", "strong", "very_strong",
        "not_expected", "unclear",
    }


def test_summary_keeps_shared_sidebar_and_verified_climate_uses_it():
    source = INDEX.read_text(encoding="utf-8")
    assert '.stage3-reading-shell[data-stage3-view="summary"]{grid-template-columns:1fr}' not in source
    assert '.stage3-reading-shell[data-stage3-view="summary"] .stage3-rating-rail' not in source
    assert 'const _showDetailedRail=isLast;' in source
    assert 'climateVerifiedRatingCards(climateVerifiedReader)' in source


def test_mobile_utility_bar_height_override_follows_desktop_rule():
    source = INDEX.read_text(encoding="utf-8")
    assert source.index('.session-bar{height:auto;') > source.index('.session-bar{display:flex;')


def test_live_detailed_avoids_duplicate_old_sensitivity_panel():
    source = INDEX.read_text(encoding="utf-8")
    assert 'renderClimateVerifiedAssessment(climateVerifiedReader,true)' in source
    assert "ratingsInSidebar?'':" in _extract_js_function(source, "renderClimateVerifiedAssessment")


def test_exports_include_responsiveness_contribution_and_scope(tmp_path):
    from docx import Document
    from sector_lenses.climate_verified_render import render_reader_html, write_reader_docx
    reader = {"judgments": [{"dimension": "responsiveness", "value": "moderate"}]}
    output = tmp_path / "report.docx"
    write_reader_docx(reader, output)
    document = Document(output)
    doc_text = "\n".join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    for text in [render_reader_html(reader), doc_text]:
        assert "Moderate contribution" in text
        assert "often a sign of going above and beyond" in text
        assert "Very limited / Limited / Moderate / Strong / Very strong" in text


@pytest.mark.parametrize(
    "status, warning",
    [
        pytest.param({"recommendation_status": "incomplete"}, "Recommendations stage incomplete", id="incomplete"),
        pytest.param(
            {"analysis_status": "partial", "analysis_message": "Evidence remains incomplete."},
            "Assessment evidence is incomplete", id="partial",
        ),
        pytest.param({"runtime_mode": "smoke"}, "Smoke test: validates workflow completion", id="smoke"),
        pytest.param(
            {"evidence_status": "preview; not approved"},
            "Candidate country evidence: preview; not approved.", id="candidate",
        ),
    ],
)
def test_genuine_status_warnings_precede_assessment(status, warning, tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from sector_lenses.climate_verified_render import render_reader_html, write_reader_docx
    reader = {**status, "executive_readout": "Assessment body."}
    output = tmp_path / "status.docx"
    write_reader_docx(reader, output)
    document = Document(output)
    text = "\n".join(Paragraph(element, document).text
                        for element in document.element.body.iter(qn("w:p")))
    for content in [text, render_reader_html(reader)]:
        normalized = content.casefold()
        warning_index = normalized.find(warning.casefold())
        body_index = normalized.find("Assessment body.".casefold())
        assert warning_index >= 0, f"missing status warning: {warning}"
        assert body_index >= 0, "missing assessment body"
        assert warning_index < body_index
