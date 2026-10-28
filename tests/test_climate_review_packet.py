"""External review packets are drafts, never app-authenticated approvals."""
import hashlib
import json
from io import BytesIO
from zipfile import ZipFile

from docx import Document
import pytest
import app as app_module
from tests.test_climate_export_snapshot import pair
from sector_lenses.climate_verified_render import render_reader_html, write_reader_docx


def packet(payload):
    return app_module.app.test_client().post('/api/climate-review-packet', json=payload)


def test_packet_binds_exact_exported_files_and_leaves_review_blank(pair):
    response = packet(pair)
    assert response.status_code == 200
    assert response.mimetype == 'application/zip'
    with ZipFile(BytesIO(response.data)) as z:
        manifest = json.loads(z.read('manifest.json'))
        record = json.loads(z.read('review_record_template.json'))
        assert manifest['publication_approved'] is False
        assert record['decision'] == 'pending'
        assert record['reviewer_name'] is None
        assert record['authenticated_workflow_record'] is None
        assert record['snapshot_sha256'] == manifest['snapshot_sha256']
        for name, digest in manifest['files'].items():
            assert hashlib.sha256(z.read(name)).hexdigest() == digest
        snapshot = json.loads(z.read('snapshot.json'))
        encoded = json.dumps(snapshot, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode()
        assert hashlib.sha256(encoded).hexdigest() == manifest['snapshot_sha256']
        doc = Document(BytesIO(z.read('report.docx')))
        word = '\n'.join(p.text for p in doc.paragraphs)
        html = z.read('report.html').decode()
        assert manifest['snapshot_sha256'] in word and manifest['snapshot_sha256'] in html
        assert 'Draft - human review required' not in word
        assert 'Draft - human review required' not in html
        assert pair['climate_reader']['priorities'][0]['minimum_action'] in word
        assert 'unflagged' in z.read('REVIEW_INSTRUCTIONS.txt').decode()


@pytest.mark.parametrize('change', ['reader', 'missing'])
def test_packet_reuses_snapshot_integrity_gate(pair, change):
    if change == 'reader':
        pair['climate_reader']['executive_readout'] = 'Unreviewed replacement.'
    else:
        pair.pop('climate_reader')
    assert packet(pair).status_code in (409, 422)


def test_changed_version_cannot_reuse_packet_identity(pair):
    first = packet(pair)
    assert first.status_code == 200
    pair['climate_assessment']['run_id'] = 'new-version'
    from sector_lenses.climate_verified_render import build_reader_model, attach_provenance
    pair['climate_reader'] = attach_provenance(build_reader_model(pair['climate_assessment']), pair['climate_assessment'])
    second = packet(pair)
    assert second.status_code == 200
    with ZipFile(BytesIO(first.data)) as a, ZipFile(BytesIO(second.data)) as b:
        assert json.loads(a.read('manifest.json'))['snapshot_sha256'] != json.loads(b.read('manifest.json'))['snapshot_sha256']


def test_client_approval_flags_do_not_change_exported_artifacts(pair):
    reader = pair['climate_reader']
    before_html = render_reader_html(reader)
    before_docx = BytesIO()
    write_reader_docx(reader, before_docx)
    before_word = Document(BytesIO(before_docx.getvalue()))
    before_content = (
        [paragraph.text for paragraph in before_word.paragraphs],
        [[cell.text for row in table.rows for cell in row.cells] for table in before_word.tables],
    )

    reader.update(publication_approved=True, human_review_complete=True, reviewer_name='Untrusted')
    after_html = render_reader_html(reader)
    after_docx = BytesIO()
    write_reader_docx(reader, after_docx)
    after_word = Document(BytesIO(after_docx.getvalue()))
    after_content = (
        [paragraph.text for paragraph in after_word.paragraphs],
        [[cell.text for row in table.rows for cell in row.cells] for table in after_word.tables],
    )

    assert 'Draft - human review required' not in after_html
    assert 'Draft - human review required' not in '\n'.join(after_content[0])
    assert after_html == before_html
    assert after_content == before_content


@pytest.mark.parametrize('scenario', ['ok', 'reject', 'changed', 'offline'])
def test_packet_download_checks_drift_and_restores_button(scenario):
    import subprocess
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function
    helper = 'async ' + _extract_js_function(INDEX.read_text(encoding='utf-8'), 'downloadClimateReviewPacket')
    script = """
let climateVerifiedAssessment={schema_version:'climate-verified-v2.1'},climateVerifiedReader={text:'Draft'};
let downloads=0,alerts=[];
const alert=x=>alerts.push(x),scenario=SCENARIO;
const document={body:{appendChild:()=>{}},createElement:()=>({click:()=>downloads++,remove:()=>{}})};
const URL={createObjectURL:()=>'',revokeObjectURL:()=>{}};
const fetch=async()=>{
 if(scenario==='offline')throw new Error('offline');
 return {ok:scenario!=='reject',json:async()=>({error:'Mismatch'}),blob:async()=>{
   if(scenario==='changed')climateVerifiedReader.text='Changed';return new Blob(['zip']);
 }};
};
HELPER
(async()=>{const btn={textContent:'Download review packet',disabled:false};await downloadClimateReviewPacket(btn);
 if(downloads!==(scenario==='ok'?1:0)||alerts.length!==(scenario==='ok'?0:1)||btn.disabled||btn.textContent!=='Download review packet')throw new Error(JSON.stringify({downloads,alerts,btn}));
})().catch(e=>{console.error(e);process.exit(1)});
""".replace('SCENARIO', json.dumps(scenario)).replace('HELPER', helper)
    result = subprocess.run(['node', '-e', script], capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr


def test_browser_views_omit_generic_draft_notice():
    from tests.test_climate_lens_frontend import INDEX, _extract_js_function
    source = INDEX.read_text(encoding='utf-8')
    for name in ['renderClimateVerifiedSummary', 'renderClimateVerifiedAssessment']:
        helper = _extract_js_function(source, name)
        assert 'Draft - human review required' not in helper
        assert 'publication_approved' not in helper


def test_posted_signoff_does_not_approve_packet(pair):
    pair.update(publication_approved=True, reviewer_name='Untrusted', decision='approved')
    response = packet(pair)
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as z:
        record = json.loads(z.read('review_record_template.json'))
        assert record['decision'] == 'pending' and record['reviewer_name'] is None
        assert json.loads(z.read('manifest.json'))['publication_approved'] is False
