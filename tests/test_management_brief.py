"""Management exports preserve the brief projection and safe presentation."""

from copy import deepcopy
from io import BytesIO
import re

from docx import Document
import pytest


def brief_fixture(count=5):
    readout = {
        "headline": "Delivery arrangements need clearer access decisions",
        "overview": (
            "The project recognizes conflict-related access risks and includes community "
            "feedback. Its main opportunity is to translate these commitments into "
            "clear delivery decisions. Better access arrangements would strengthen "
            "conflict-sensitive implementation. Directly changing conflict drivers "
            "is not an objective of this operation, so limited responsiveness should "
            "not be interpreted as poor design."
        ),
        "strengths": [
            {"title": "Community feedback", "text": "The design includes accessible feedback channels."},
            {"title": "Context awareness", "text": "The document identifies access constraints."},
        ],
    }
    priorities = [
        {
            "number": number,
            "title": f"Detailed priority {number}",
            "concise": {
                "title": f"Clarify access decision {number}",
                "why": "Unclear responsibilities can delay delivery to affected communities.",
                "how": [
                    f"Consider assigning the decision owner for activity {number}.",
                    "Further implementation detail belongs in the full assessment.",
                ],
                "suggested_wording": {"text": "DRAFTING ONLY IN DETAILED"},
            },
        }
        for number in range(1, count + 1)
    ]
    return readout, priorities


@pytest.mark.parametrize("count", [1, 2, 5])
def test_html_and_word_preserve_all_priorities_and_leading_actions(count):
    from fcv_management_brief import (
        render_management_brief_docx,
        render_management_brief_html,
    )

    readout, priorities = brief_fixture(count)
    html = render_management_brief_html(readout, priorities)
    document = Document(BytesIO(render_management_brief_docx(readout, priorities)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for output in (html, text):
        assert readout["headline"] in output
        assert "DRAFTING ONLY IN DETAILED" not in output
        assert "Further implementation detail belongs" not in output
        previous = -1
        for priority in priorities:
            card = priority["concise"]
            position = output.index(card["title"])
            assert position > previous
            previous = position
            assert card["how"][0] in output
        assert "full assessment" in output.lower()
        assert "not requirements" in output.lower()
    assert document.sections[0].page_width.mm == pytest.approx(210, abs=0.1)
    assert document.sections[0].page_height.mm == pytest.approx(297, abs=0.1)
    assert document.styles["Normal"].font.name == "Arial"


def test_html_escapes_model_text_and_needs_no_external_assets():
    from fcv_management_brief import render_management_brief_html

    readout, priorities = brief_fixture(1)
    readout["headline"] = '<img src=x onerror="alert(1)">'
    priorities[0]["concise"]["why"] = "A & B < C"
    html = render_management_brief_html(readout, priorities)
    assert "<img" not in html
    assert "&lt;img" in html
    assert "A &amp; B &lt; C" in html
    assert not re.search(r'<(?:script|link|iframe)\b', html)
    assert "@page" in html
    assert "@media print" in html


def test_no_strengths_does_not_invent_positive_findings():
    from fcv_management_brief import render_management_brief_html

    readout, priorities = brief_fixture(1)
    readout["strengths"] = []
    html = render_management_brief_html(readout, priorities)
    assert "What the project does well" not in html
    assert priorities[0]["concise"]["title"] in html


@pytest.mark.parametrize("change", ["no-priorities", "six-priorities", "missing-concise", "missing-action", "empty-headline"])
def test_incomplete_bundle_cannot_produce_a_misleading_brief(change):
    from fcv_management_brief import render_management_brief_html

    readout, priorities = brief_fixture(6 if change == "six-priorities" else 1)
    if change == "no-priorities":
        priorities = []
    elif change == "missing-concise":
        priorities[0].pop("concise")
    elif change == "missing-action":
        priorities[0]["concise"]["how"] = []
    elif change == "empty-headline":
        readout["headline"] = " "
    with pytest.raises(ValueError):
        render_management_brief_html(readout, priorities)


def test_exports_do_not_mutate_canonical_data():
    from fcv_management_brief import render_management_brief_html

    readout, priorities = brief_fixture(2)
    before = deepcopy((readout, priorities))
    render_management_brief_html(readout, priorities)
    assert (readout, priorities) == before


def test_brief_orders_gaps_before_action_only_priorities():
    from fcv_management_brief import (
        render_management_brief_docx,
        render_management_brief_html,
    )

    readout, priorities = brief_fixture(2)
    priorities[0]["concise"]["gap"] = (
        "Ownership is unclear. This can delay decisions."
    )
    priorities[1]["concise"]["why"] = "The second reason remains relevant to delivery."
    html = render_management_brief_html(readout, priorities)
    text = "\n".join(
        paragraph.text
        for paragraph in Document(BytesIO(render_management_brief_docx(readout, priorities))).paragraphs
    )

    for output in (html, text):
        assert output.index("What the project does well") < output.index("Potential gaps")
        assert output.index("Potential gaps") < output.index("Suggested priorities")
        for sentence in ("Ownership is unclear.", "This can delay decisions."):
            assert sentence in output
        assert priorities[0]["concise"]["why"] not in output
        assert "Why it matters:" not in output[output.index("Suggested priorities"):]
        assert priorities[0]["concise"]["how"][0] in output


def test_brief_bolds_robust_first_sentence_and_normalizes_prose_punctuation():
    from fcv_management_brief import render_management_brief_html

    readout, priorities = brief_fixture(1)
    readout["overview"] = (
        "The U.S. route costs 1.5m \u2014 before works. "
        "See https://example.org/a.b for context; later."
    )
    html = render_management_brief_html(readout, priorities)

    assert (
        "<strong>The U.S. route costs 1.5m - before works.</strong> "
        "See https://example.org/a.b for context; later."
    ) in html
