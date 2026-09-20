"""Render the standard FCV management projection as print HTML or editable Word.

Callers must first admit the bundle through the canonical Stage 3 validator.
The two formats share one projection and never modify the detailed assessment.
"""

from html import escape
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Mm, Pt, RGBColor

from fcv_presentation import normalize_display_text, split_first_sentence
from fcv_word_style import style_fcv_word_document


ADVISORY = (
    "AI-assisted suggestions for professional review, not requirements. "
    "Assess their relevance with the task team and relevant specialists."
)
INTERPRETATION = (
    "Sensitivity concerns how the project is designed for its FCV context. "
    "Responsiveness concerns contributions to FCV drivers. This is not an expectation "
    "for every operation."
)
DETAIL_NOTE = (
    "This brief presents the leading action for each priority. See the full "
    "assessment for all actions, supporting evidence, formal ratings and suggested wording."
)


def _text(value: Any) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError("A complete validated management brief is required.")
    return normalize_display_text(value.strip())


def _html_body(text: str, *, label: str = "") -> str:
    first, remainder = split_first_sentence(text)
    prefix = f"<strong>{escape(label)}</strong> " if label else ""
    body = f"<strong>{escape(first)}</strong>"
    if remainder:
        body += f" {escape(remainder)}"
    return prefix + body


def _projection(readout: dict, priorities: list) -> dict:
    """Select the same concise fields in both export formats, without truncation."""
    if not isinstance(readout, dict) or not isinstance(priorities, list):
        raise ValueError("A complete validated management brief is required.")
    if not 1 <= len(priorities) <= 5:
        raise ValueError("A management brief requires one to five priorities.")
    strengths_raw = readout.get("strengths")
    if not isinstance(strengths_raw, list) or len(strengths_raw) > 3:
        raise ValueError("A management brief supports up to three strengths.")
    strengths = []
    for strength in strengths_raw:
        if not isinstance(strength, dict):
            raise ValueError("Invalid strength.")
        strengths.append((_text(strength.get("title")), _text(strength.get("text"))))
    cards = []
    for priority in priorities:
        concise = priority.get("concise") if isinstance(priority, dict) else None
        if not isinstance(concise, dict):
            raise ValueError("Every priority needs an admitted concise projection.")
        actions = concise.get("how")
        if not isinstance(actions, list) or not actions:
            raise ValueError("Every priority needs a leading action.")
        gap = concise.get("gap") or concise.get("why")
        cards.append({
            "title": _text(concise.get("title")),
            "gap": _text(gap),
            "action": _text(actions[0]),
        })
    return {
        "headline": _text(readout.get("headline")),
        "overview": _text(readout.get("overview")),
        "strengths": strengths,
        "priorities": cards,
    }


def render_management_brief_html(readout: dict, priorities: list) -> str:
    """Return self-contained, escaped A4 HTML suitable for Print / Save as PDF."""
    brief = _projection(readout, priorities)
    parts = [
        '<!doctype html><html lang="en"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        '<title>FCV management brief</title>',
        '<style>',
        'body{font:11pt/1.42 Arial,sans-serif;color:#172536;max-width:760px;',
        'margin:32px auto;padding:0 20px;overflow-wrap:anywhere}',
        'h1{font-size:19pt;color:#000;margin:0 0 12px}',
        'h2{font-size:12.5pt;color:#153956;margin:16px 0 8px}',
        'h3{font-size:11pt;color:#153956;margin:0 0 6px}',
        'p{margin:0 0 10px}.section{margin-bottom:12px}',
        '.strengths{background:#eff8f0;border-left:4px solid #4f9d69;padding:10px 12px}',
        '.gaps{background:#fff4e5;border-left:4px solid #d98c34;padding:10px 12px}',
        '.strength,.gap{margin-bottom:8px}.priority{margin-bottom:14px;',
        'break-inside:avoid}.note{font-size:9pt;color:#435166}',
        '@page{size:A4;margin:16mm}@media print{body{margin:0;padding:0;',
        'max-width:none;font-size:11pt;line-height:1.32}h1{font-size:17pt}',
        'h2{margin-top:12px}.priority{margin-bottom:10px}}</style>',
        '</head><body><main><h1>FCV management brief</h1>',
        '<section class="section overall"><h2>Overall assessment</h2>',
        f'<p class="headline"><strong>{escape(brief["headline"])}</strong></p>',
        f'<p>{_html_body(brief["overview"])}</p></section>',
    ]
    if brief["strengths"]:
        parts.append('<section class="section strengths"><h2>What the project does well</h2>')
        for title, text in brief["strengths"]:
            parts.append(f'<p class="strength">{_html_body(text, label=title + ":")}</p>')
        parts.append("</section>")
    parts.append('<section class="section gaps"><h2>Potential gaps</h2>')
    for number, card in enumerate(brief["priorities"], 1):
        parts.append(
            f'<p class="gap">{_html_body(card["gap"], label=f"{number}.")}</p>'
        )
    parts.append("</section>")
    parts.append('<section class="section priorities">')
    parts.append("<h2>Suggested priorities</h2>")
    for number, card in enumerate(brief["priorities"], 1):
        parts.extend([
            f'<section class="priority"><h3>{number}. {escape(card["title"])}</h3>',
            f'<p>{_html_body(card["action"], label="Suggested action:")}</p></section>',
        ])
    parts.append("</section>")
    for text in (INTERPRETATION, DETAIL_NOTE, ADVISORY):
        parts.append(f'<p class="note">{_html_body(normalize_display_text(text))}</p>')
    parts.append("</main></body></html>")
    return "".join(parts)


def _paragraph(document, text: str, *, label: str = "", style=None):
    """Add an editable paragraph with a bold first sentence and plain remainder."""
    paragraph = document.add_paragraph(style=style)
    first, remainder = split_first_sentence(normalize_display_text(text))
    if label:
        paragraph.add_run(normalize_display_text(label) + " ").bold = True
    if first:
        paragraph.add_run(first).bold = True
    if remainder:
        paragraph.add_run(" " + remainder)
    return paragraph


def render_management_brief_docx(readout: dict, priorities: list) -> bytes:
    """Return an editable A4 brief with comfortable native Word flow."""
    brief = _projection(readout, priorities)
    document = Document()
    document.core_properties.title = "FCV management brief"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(17)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(17)
    document.add_paragraph("FCV management brief", style="Title")
    document.add_heading("Overall assessment", level=1)
    _paragraph(document, brief["headline"])
    _paragraph(document, brief["overview"])
    if brief["strengths"]:
        document.add_heading("What the project does well", level=1)
        for title, text in brief["strengths"]:
            _paragraph(document, text, label=title + ":", style="List Bullet")
    document.add_heading("Potential gaps", level=1)
    for number, card in enumerate(brief["priorities"], 1):
        gap_paragraph = _paragraph(
            document,
            card["gap"],
            label=f"Gap {number}: ",
            style="List Bullet",
        )
        gap_paragraph.paragraph_format.keep_together = True
    document.add_heading("Suggested priorities", level=1)
    for number, card in enumerate(brief["priorities"], 1):
        document.add_heading(f'{number}. {card["title"]}', level=2)
        action = _paragraph(
            document,
            card["action"],
            label="Suggested action:",
        )
        action.paragraph_format.keep_with_next = False
    for text in (INTERPRETATION, DETAIL_NOTE, ADVISORY):
        _paragraph(document, text)
    style_fcv_word_document(document, variant="brief")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
